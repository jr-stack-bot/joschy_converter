"""
Kreis-Konverter – futuristischer, runder Datei-Konverter für Windows (HUD-Design)
© JOSCHY – Lizenz: AGPL-3.0
Quellcode: https://github.com/jr-stack-bot/joschy_converter

Dieses Programm ist freie Software: Sie können es unter den Bedingungen der GNU Affero General Public
License, Version 3, wie von der Free Software Foundation veröffentlicht, weitergeben und/oder verändern.
Es wird OHNE JEDE GEWÄHRLEISTUNG bereitgestellt. Den vollständigen Lizenztext enthält die Datei LICENSE.
"""
import csv
import asyncio
import functools
import io
import json
import lzma
import math
import os
import queue
import random
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
import threading
import time
import tkinter as tk
import traceback
import zipfile
import base64
import importlib
import types
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from tkinter import filedialog

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageOps, ImageTk

import hud_grafik as hg
from hud_grafik import (  # noqa: F401  (auch als kreis_konverter.X erreichbar)
    BTN_DX, BTN_R, BTN_Y, EMB_GLYPH, FARBEN, PILL_Y, R_DISC, S_MINI, S_NORMAL, STANDARD_THEME, THEMES, TRANS,
    TRANS_RGBA, BildDaten, Grafik, P, _Lazy, auf_ebene, farbe_aus, hexfarbe, hexrgb, kreis_maske, mix, polar,
    ring_maske, setze_skalierung, theme_daten, verkleinere, verlauf)

try:
    import winsound
    SOUND_AVAILABLE = True
except Exception:
    SOUND_AVAILABLE = False

try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    DND_IMPORT = True
except Exception:
    DND_IMPORT = False

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
    HEIC_AVAILABLE = True
except Exception:
    HEIC_AVAILABLE = False

try:
    import pymupdf as fitz          # PDF -> Bild, PDF komprimieren, OCR-Vorstufe
except Exception:
    try:
        import fitz
    except Exception:
        fitz = None

# PyMuPDF ist nicht thread-sicher, es werden aber mehrere Dateien gleichzeitig umgewandelt:
# Alle Zugriffe laufen deshalb nacheinander (die Umwandlung anderer Formate bleibt parallel).
FITZ_LOCK = threading.RLock()
# Microsoft Office wird nur genutzt, wenn "office_bevorzugen": true in einstellungen.json steht.
# Standard: aus - die eingebauten Konverter liefern auf jedem Rechner dasselbe Ergebnis und können nicht
# an einem Office-Dialog (Aktivierung, Erststart, Passwort) hängen bleiben.
OFFICE_BEVORZUGEN = [False]
# Bildrate: Nach so vielen Sekunden ohne Eingabe bleibt das Bild im Leerlauf stehen (Standbild);
# danach prüft ein langsamer Takt (ms) nur noch den Wechsel des Hinweistexts.
STANDBILD_NACH = 20.0
PRUEF_TAKT = 500


def fitz_sicher(funktion):
    @functools.wraps(funktion)
    def innen(*a, **k):
        with FITZ_LOCK:
            return funktion(*a, **k)
    return innen


try:
    import pystray                  # Tray-Icon (optional)
    TRAY_AVAILABLE = True
except Exception:
    pystray = None
    TRAY_AVAILABLE = False

# Prozess-Priorität unter Windows auf "Hintergrund" senken (PC bleibt flüssig)
if os.name == "nt":
    try:
        import ctypes
        # 0x00004000 = BELOW_NORMAL_PRIORITY_CLASS
        ctypes.windll.kernel32.SetPriorityClass(ctypes.windll.kernel32.GetCurrentProcess(), 0x00004000)
    except Exception:
        pass

# ------------------------------------------------------------------ Design & Themes
WEISS, GRAU = "#ffffff", "#86a9ad"
FH, FM = "Bahnschrift", "Consolas"


def dpi_bewusst_machen():
    """Windows: Prozess DPI-bewusst machen, damit nichts unscharf hochgerechnet wird.
    Muss vor dem ersten Tk-Fenster passieren. Fehler werden ignoriert (dann bleibt alles wie bisher)."""
    if os.name != "nt":
        return
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(2)        # Windows 8.1+: pro Monitor
        return
    except Exception:
        pass
    try:
        import ctypes
        ctypes.windll.user32.SetProcessDPIAware()             # Windows Vista/7
    except Exception:
        pass


def ermittle_skalierung(root):
    """Skalierungsfaktor in 0,25-Schritten (1.0 ... 3.0). KREIS_SKALIERUNG=1.5 erzwingt einen Wert (Tests)."""
    f = None
    roh = os.environ.get("KREIS_SKALIERUNG", "").strip()
    if roh:
        try:
            f = float(roh.replace(",", "."))
            root.tk.call("tk", "scaling", f * 96.0 / 72.0)    # Schriften wie unter Windows mitskalieren
        except (ValueError, tk.TclError):
            f = None
    if f is None and os.name == "nt":
        try:
            import ctypes
            dpi = ctypes.windll.user32.GetDpiForSystem()      # Windows 10 1607+
            if dpi > 0:
                f = dpi / 96.0
        except Exception:
            pass
    if f is None:
        try:
            f = root.winfo_fpixels("1i") / 96.0
        except Exception:
            f = 1.0
    f = round(f * 4) / 4.0
    return max(1.0, min(3.0, f))


def monitor_bereich(cx, cy, root=None):
    """Arbeitsbereich (links, oben, rechts, unten) des Monitors, auf dem der Punkt liegt - sonst None."""
    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes

            class MONITORINFO(ctypes.Structure):
                _fields_ = [("cbSize", wintypes.DWORD), ("rcMonitor", wintypes.RECT),
                            ("rcWork", wintypes.RECT), ("dwFlags", wintypes.DWORD)]
            user32 = ctypes.windll.user32
            user32.MonitorFromPoint.restype = wintypes.HANDLE
            user32.MonitorFromPoint.argtypes = [wintypes.POINT, wintypes.DWORD]
            user32.GetMonitorInfoW.argtypes = [wintypes.HANDLE, ctypes.POINTER(MONITORINFO)]
            h = user32.MonitorFromPoint(wintypes.POINT(int(cx), int(cy)), 0)     # 0 = MONITOR_DEFAULTTONULL
            if not h:
                return None
            mi = MONITORINFO()
            mi.cbSize = ctypes.sizeof(MONITORINFO)
            if user32.GetMonitorInfoW(h, ctypes.byref(mi)):
                r = mi.rcWork
                return (r.left, r.top, r.right, r.bottom)
        except Exception:
            pass
    if root is None:
        return None
    sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
    return (0, 0, sw, sh) if (0 <= cx < sw and 0 <= cy < sh) else None


def fenster_position(gespeichert, groesse, root):
    """Linke obere Ecke für den Kreis: gespeicherte Mitte, falls sie auf einem Monitor liegt
    (dann ganz in dessen Arbeitsbereich geschoben), sonst mittig auf dem Hauptbildschirm."""
    sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
    mitte = ((sw - groesse) // 2, (sh - groesse) // 2)
    try:
        cx, cy = int(gespeichert["x"]), int(gespeichert["y"])
    except Exception:
        return mitte
    bereich = monitor_bereich(cx, cy, root)
    if not bereich:
        return mitte
    l, t, r, b = bereich
    x = min(max(cx - groesse // 2, l), max(l, r - groesse))
    y = min(max(cy - groesse // 2, t), max(t, b - groesse))
    return x, y


KATEGORIEN = {
    "Bild": ["png", "jpg", "webp", "bmp", "gif", "tiff", "ico", "svg", "heic", "resize 50%", "resize 1920px", "resize 1280px", "resize 800px"],
    "Office": ["docx", "xlsx", "pptx"],
    "PDF": ["pdf", "pdf (komp)", "pdf teilen", "pdf zusammenführen", "pdf zu bild", "txt (ocr)"],
    "Daten": ["csv", "json", "xml", "txt", "md"],
    "E-Book": ["epub"],
    "Archiv": ["zip", "zip (alle)", "7z", "7z (alle)", "entpacken"],
    "QR": ["qrcode"],
    "Audio": ["mp3", "wav", "flac", "ogg", "aac", "m4a"],
    "Video": ["mp4", "mkv", "webm", "avi", "mov"],
}
ALLE_FORMATE = [f for liste in KATEGORIEN.values() for f in liste]
EMB_KEYS = set(KATEGORIEN) | {"ok", "err"}
BILD_FORMATE = {"png": "PNG", "jpg": "JPEG", "jpeg": "JPEG", "webp": "WEBP", "bmp": "BMP",
                "gif": "GIF", "tiff": "TIFF", "tif": "TIFF", "ico": "ICO", "pdf": "PDF",
                "pdf (komp)": "PDF", "svg": "SVG", "heic": "HEIF"}
BILD_EINGABE = {"png", "jpg", "jpeg", "webp", "bmp", "gif", "tiff", "tif", "heic", "ico"}
AUDIO_EINGABE = {"mp3", "wav", "flac", "ogg", "aac", "m4a", "wma", "opus"}
MEDIEN_EINGABE = {"mp3", "wav", "flac", "ogg", "aac", "m4a", "wma", "opus",
                  "mp4", "mkv", "webm", "avi", "mov", "wmv", "flv", "m4v"}
OFFICE_EINGABE = {"docx", "doc", "xlsx", "xls", "pptx", "ppt", "rtf", "odt", "pdf"}
DATEN_EINGABE = {"csv", "json", "xml", "txt", "md", "log"}
EBOOK_EINGABE = {"epub"}
ARCHIV_EINGABE = {"zip", "7z", "rar"}                     # entpacken: zip (eingebaut), 7z/rar (über 7-Zip)
PACK_AKTIONEN = {"zip", "zip (alle)", "7z", "7z (alle)"}   # nehmen jede Datei und ganze Ordner
ARCHIV_AKTIONEN = PACK_AKTIONEN | {"entpacken"}
SPEZIAL_ZIELE = {"pdf teilen", "pdf zusammenführen", "pdf zu bild", "txt (ocr)", "qrcode"} | ARCHIV_AKTIONEN
AUDIO_AUSGABE = set(KATEGORIEN["Audio"])
MEDIEN_AUSGABE = AUDIO_AUSGABE | set(KATEGORIEN["Video"])
ALLE_EINGABE = (BILD_EINGABE | MEDIEN_EINGABE | OFFICE_EINGABE | DATEN_EINGABE
                | EBOOK_EINGABE | ARCHIV_EINGABE)
GLEICH = {("jpg", "jpeg"), ("jpeg", "jpg"), ("tif", "tiff"), ("tiff", "tif")}

# Optionale Erweiterungen (siehe lade_erweiterungen): werden beim Start eingetragen, sonst bleiben sie leer
ERWEITERUNGEN = []                 # geladene Erweiterungen (Beschreibungen als dict)
ERW_FORMATE = {}                   # Format -> {"quellen", "endung", "umwandeln", "falsche_quelle"}
# Bei diesen Aktionen wird nie automatisch erkannt und umgeleitet (Sammel-/Archiv-Aktionen)
ERKENNUNG_AUSNAHMEN = {"pdf zusammenführen"} | ARCHIV_AKTIONEN

PFAD_LOCK = threading.Lock()
OFFICE_LOCK = threading.Lock()       # Office-Programme immer nur einzeln ansprechen
MEDIEN_SEM = threading.Semaphore(2)  # höchstens 2 ffmpeg-Prozesse gleichzeitig
BILD_SEM = threading.Semaphore(4)    # begrenzt den Speicherbedarf bei großen Bildern
PROZESSE = set()                     # laufende ffmpeg-Prozesse (zum Abbrechen)


class Abgebrochen(Exception):
    pass


def basis_ordner():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def daten_ordner():
    """Beschreibbarer Ordner für Log und Einstellungen (auch bei Installation unter Program Files)."""
    wurzel = os.environ.get("LOCALAPPDATA") or os.path.join(os.path.expanduser("~"), ".local", "share")
    pfad = os.path.join(wurzel, "KreisKonverter")
    try:
        os.makedirs(pfad, exist_ok=True)
        return pfad
    except Exception:
        return basis_ordner()


def logge(text):
    try:
        log_pfad = os.path.join(daten_ordner(), "kreis_fehler.log")
        try:
            if os.path.getsize(log_pfad) > 512 * 1024:
                os.replace(log_pfad, log_pfad + ".alt")
        except OSError:
            pass
        t_str = time.strftime("%Y-%m-%d %H:%M:%S")
        with open(log_pfad, "a", encoding="utf-8") as f:
            f.write(f"{t_str} {text}\n")
    except Exception:
        pass


def lade_einstellungen():
    try:
        with open(os.path.join(daten_ordner(), "einstellungen.json"), "r", encoding="utf-8") as f:
            d = json.load(f)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def speichere_einstellungen(d):
    try:
        pfad = os.path.join(daten_ordner(), "einstellungen.json")
        with open(pfad + ".tmp", "w", encoding="utf-8") as f:
            json.dump(d, f, indent=2, ensure_ascii=False)
        os.replace(pfad + ".tmp", pfad)
    except Exception:
        pass


def kat_von_format(f):
    return next(k for k, v in KATEGORIEN.items() if f in v)


def sichtbare_kategorien():
    """Kategorien für Menü/Pfeiltasten (inkl. geladener Erweiterungen)."""
    return list(KATEGORIEN)


def sichtbare_formate():
    return [f for k in sichtbare_kategorien() for f in KATEGORIEN[k]]


FFMPEG_LOCK = threading.Lock()
_FFMPEG_PFAD = [None]


def ressourcen_ordner():
    """Ordner mit den in die exe eingebetteten Dateien (PyInstaller entpackt sie nach _MEIPASS)."""
    return getattr(sys, "_MEIPASS", None) or os.path.dirname(os.path.abspath(__file__))


def _entpacke_ffmpeg(xz):
    """Entpackt das eingebettete ffmpeg einmalig nach %LOCALAPPDATA%\\KreisKonverter\\ffmpeg (Cache)."""
    ziel_dir = os.path.join(daten_ordner(), "ffmpeg")
    os.makedirs(ziel_dir, exist_ok=True)
    name = f"ffmpeg-{os.path.getsize(xz)}.exe"
    ziel = os.path.join(ziel_dir, name)
    if os.path.isfile(ziel) and os.path.getsize(ziel) > 1_000_000:
        return ziel
    tmp = f"{ziel}.{os.getpid()}.tmp"
    try:
        with lzma.open(xz, "rb") as quelle, open(tmp, "wb") as f_out:
            shutil.copyfileobj(quelle, f_out, 1024 * 1024)
        try:
            os.replace(tmp, ziel)
        except OSError:
            if not (os.path.isfile(ziel) and os.path.getsize(ziel) > 1_000_000):
                raise            # sonst hat gerade eine andere Instanz dieselbe Datei fertiggestellt
    finally:
        raeume_auf(tmp)
    try:
        os.chmod(ziel, 0o755)
    except OSError:
        pass
    for alt_datei in os.listdir(ziel_dir):      # alte Versionen aufräumen
        if alt_datei.startswith("ffmpeg-") and alt_datei.endswith(".exe") and alt_datei != name:
            try:
                os.remove(os.path.join(ziel_dir, alt_datei))
            except OSError:
                pass
    return ziel


def _suche_ffmpeg():
    lokal = os.path.join(basis_ordner(), "ffmpeg.exe")          # 1. eigene ffmpeg.exe neben der exe hat Vorrang
    if os.path.isfile(lokal):
        return lokal
    for ordner in (ressourcen_ordner(), basis_ordner()):         # 2. in die exe eingebettet
        xz = os.path.join(ordner, "ffmpeg.xz")
        if os.path.isfile(xz):
            try:
                return _entpacke_ffmpeg(xz)
            except Exception:
                logge("ffmpeg entpacken fehlgeschlagen:\n" + traceback.format_exc())
    try:                                                         # 3. Python-Paket (beim Start aus dem Quelltext)
        import imageio_ffmpeg
        pfad = imageio_ffmpeg.get_ffmpeg_exe()
        if pfad and os.path.isfile(pfad):
            return pfad
    except Exception:
        pass
    return shutil.which("ffmpeg")                                # 4. systemweit installiert


def finde_ffmpeg(neu_suchen=False):
    with FFMPEG_LOCK:
        if neu_suchen or not (_FFMPEG_PFAD[0] and os.path.isfile(_FFMPEG_PFAD[0])):
            _FFMPEG_PFAD[0] = _suche_ffmpeg()
        return _FFMPEG_PFAD[0]


def freier_pfad(pfad):
    """Liefert einen freien Dateinamen und legt ihn sofort als leere Datei an (Reservierung),
    damit parallele Threads sich nicht gegenseitig überschreiben."""
    name, endung = os.path.splitext(pfad)
    kandidat, i = pfad, 0
    with PFAD_LOCK:
        while True:
            try:
                with open(kandidat, "x"):
                    return kandidat
            except (FileExistsError, PermissionError):
                if not os.path.exists(kandidat):
                    raise
                i += 1
                kandidat = f"{name} ({i}){endung}"


def freier_ordner(pfad):
    """Legt einen noch nicht vorhandenen Ordner an (name, name (1), ...)."""
    kandidat, i = pfad, 0
    with PFAD_LOCK:
        while True:
            try:
                os.makedirs(kandidat)
                return kandidat
            except FileExistsError:
                i += 1
                kandidat = f"{pfad} ({i})"


def raeume_auf(pfad):
    """Entfernt eine Hilfsdatei bzw. eine leere Platzhalter-Reservierung."""
    try:
        if pfad and os.path.isfile(pfad):
            os.remove(pfad)
    except OSError:
        pass


def ext_von(pfad):
    return os.path.splitext(pfad)[1].lower().lstrip(".")


def kuerze(text, n):
    return text if len(text) <= n else text[: n - 1] + "…"


def umbrechen(text, breite, max_zeilen):
    zeilen = textwrap.wrap(str(text), breite, break_long_words=True) or [""]
    if len(zeilen) > max_zeilen:
        zeilen = zeilen[:max_zeilen]
        zeilen[-1] = kuerze(zeilen[-1] + "…", breite)
    return zeilen


def gesperrt(text):
    return " ".join(text.upper())


def natur_key(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", os.path.basename(s))]


def verstaendlich(ex):
    """Übersetzt technische Fehlermeldungen der Bibliotheken in verständliche Sätze."""
    if type(ex) is RuntimeError:                  # eigene, bereits verständliche Meldungen
        return str(ex)
    name, text = type(ex).__name__, str(ex)
    if name == "UnidentifiedImageError":
        return "Die Datei ist beschädigt oder kein gültiges Bild"
    if name.startswith("DecompressionBomb"):
        return "Das Bild ist zu groß, um es sicher zu öffnen"
    if isinstance(ex, MemoryError):
        return "Nicht genug Arbeitsspeicher für diese Datei"
    if isinstance(ex, PermissionError):
        return "Kein Zugriff: Die Datei ist geöffnet, gesperrt oder schreibgeschützt"
    if isinstance(ex, FileNotFoundError):
        return "Datei nicht gefunden (wurde sie verschoben oder gelöscht?)"
    if isinstance(ex, OSError) and getattr(ex, "errno", None) == 28:
        return "Der Datenträger ist voll"
    if name == "XLRDError":
        return "Die Excel-Datei ist beschädigt oder passwortgeschützt"
    if name == "JSONDecodeError":
        return "Die Datei enthält kein gültiges JSON"
    if name == "Error" and type(ex).__module__ == "_csv":
        return "Die Datei ist keine gültige CSV-Tabelle"
    if name in ("InvalidFileException", "BadZipFile", "PackageNotFoundError") or isinstance(ex, (zipfile.BadZipFile, ET.ParseError, KeyError)):
        return "Die Datei ist beschädigt oder kein gültiges Dokument dieses Typs"
    if name in ("PdfReadError", "EmptyFileError", "PdfStreamError", "FileDataError", "PyPdfError") or "pdf" in name.lower():
        return "Die PDF-Datei ist beschädigt oder leer"
    return text or name


def standard_ausgabe_ordner():
    """Wohin Dateien ohne gewählten Zielordner kommen (z. B. QR-Codes): Desktop, sonst Benutzerordner."""
    home = os.path.expanduser("~")
    desktop = os.path.join(home, "Desktop")
    return desktop if os.path.isdir(desktop) else home


def groesse_von(pfad):
    try:
        if os.path.isfile(pfad):
            return os.path.getsize(pfad)
        gesamt = 0
        for root, _, files in os.walk(pfad):
            for f in files:
                try:
                    gesamt += os.path.getsize(os.path.join(root, f))
                except OSError:
                    pass
        return gesamt
    except OSError:
        return 0


def ausgabe_ext(ziel_ext):
    if ziel_ext in ERW_FORMATE:
        return ERW_FORMATE[ziel_ext]["endung"]
    return {"pdf (komp)": "pdf"}.get(ziel_ext, ziel_ext)


def spiele_sound(art):
    if not SOUND_AVAILABLE:
        return
    def _run():
        try:
            if art == "drop":
                winsound.Beep(800, 60)
            elif art == "ok":
                winsound.Beep(1200, 80)
                winsound.Beep(1600, 100)
            elif art == "err":
                winsound.Beep(400, 150)
        except Exception:
            pass
    threading.Thread(target=_run, daemon=True).start()


# ------------------------------------------------------------------ JAI (Sprachansagen)
class JAI:
    """Kurze Sprachansagen aus jai/<id>.mp3 (in der exe aus _MEIPASS, sonst neben dem Skript).
    Wiedergabe über Windows-MCI (winmm, per ctypes) in einem eigenen Thread – keine Zusatzpakete.
    Eine neue Ansage unterbricht die laufende. Fehlt der Ordner oder ein Clip oder klappt die
    Wiedergabe nicht, passiert einfach nichts."""
    ALIAS = "kreis_jai"

    def __init__(self, mci=None):
        self.an = False
        self._mci = mci                 # Testen: eigene Funktion statt winmm
        self._auftraege = queue.Queue()
        self._thread = None
        self._offen = False

    @staticmethod
    def ordner():
        return os.path.join(ressourcen_ordner(), "jai")

    def clip(self, cid):
        pfad = os.path.join(self.ordner(), f"{cid}.mp3")
        return pfad if os.path.isfile(pfad) else None

    def _sende(self, befehl):
        if self._mci is None:
            import ctypes
            winmm = ctypes.windll.winmm
            fn = winmm.mciSendStringW

            def mci(b):
                return fn(b, None, 0, None)
            self._mci = mci
        return self._mci(befehl)

    def _lauf(self):
        while True:
            pfad = self._auftraege.get()
            try:
                while True:              # nur die neueste Ansage zählt
                    pfad = self._auftraege.get_nowait()
            except queue.Empty:
                pass
            try:
                if self._offen:
                    self._sende(f"close {self.ALIAS}")
                    self._offen = False
                if self._sende(f'open "{pfad}" type mpegvideo alias {self.ALIAS}') == 0:
                    self._offen = True
                    self._sende(f"play {self.ALIAS} from 0")
            except Exception:
                pass                     # winmm nicht verfügbar o. Ä.: still überspringen

    def sage(self, cid, immer=False):
        """True, wenn eine Ansage gestartet wurde."""
        if not (self.an or immer):
            return False
        try:
            pfad = self.clip(cid)
            if not pfad or (self._mci is None and not hasattr(sys, "getwindowsversion")):
                return False
            if self._thread is None:
                self._thread = threading.Thread(target=self._lauf, daemon=True, name="JAI")
                self._thread.start()
            self._auftraege.put(pfad)
            return True
        except Exception:
            return False


# ------------------------------------------------------------------ Bilder
def _gedreht(img):
    """EXIF-Drehung immer anwenden (sonst liegen Handyfotos nach dem Konvertieren schief)."""
    try:
        return ImageOps.exif_transpose(img)
    except Exception:
        return img


def _meta_args(img, fmt, strip_exif):
    """Metadaten (EXIF, Farbprofil) für das Speichern – nur wenn nicht entfernt werden soll."""
    args = {}
    if strip_exif or fmt not in ("JPEG", "WEBP", "PNG", "HEIF"):
        return args
    try:
        ex = img.getexif()
        if ex:
            ex[0x0112] = 1  # Drehung wurde bereits angewendet
            args["exif"] = ex.tobytes()
    except Exception:
        pass
    icc = img.info.get("icc_profile")
    if icc:
        args["icc_profile"] = icc
    return args


def konvertiere_bild_verkleinern(quelle, zielordner, ziel_ext, qualitaet, strip_exif=False):
    with Image.open(quelle) as img0:
        img = _gedreht(img0)
        w, h = img.size

        if ziel_ext == "resize 50%":
            neu_w = max(1, int(w * 0.5))
            neu_h = max(1, int(h * 0.5))
        else:
            max_dim = int(ziel_ext.replace("resize ", "").replace("px", ""))
            if max(w, h) > max_dim:
                ratio = max_dim / float(max(w, h))
                neu_w = max(1, int(w * ratio))
                neu_h = max(1, int(h * ratio))
            else:
                neu_w, neu_h = w, h

        quell_ext = ext_von(quelle)
        if quell_ext in ("jpg", "jpeg", "heic"):
            fmt, ext_ausgabe = "JPEG", "jpg"
        elif quell_ext == "png":
            fmt, ext_ausgabe = "PNG", "png"
        elif quell_ext == "webp":
            fmt, ext_ausgabe = "WEBP", "webp"
        else:
            fmt, ext_ausgabe = "JPEG", "jpg"

        meta = _meta_args(img, fmt, strip_exif)
        with img.resize((neu_w, neu_h), Image.LANCZOS) as img_res:
            if fmt == "JPEG" and img_res.mode in ("RGBA", "LA", "P"):
                with img_res.convert("RGBA") as rgba:
                    with Image.new("RGB", rgba.size, (255, 255, 255)) as hg:
                        hg.paste(rgba, mask=rgba.split()[-1])
                        img_res = hg
            elif fmt == "JPEG" and img_res.mode != "RGB":
                img_res = img_res.convert("RGB")

            basis = os.path.splitext(os.path.basename(quelle))[0]
            ziel = freier_pfad(os.path.join(zielordner, f"{basis}_{neu_w}x{neu_h}.{ext_ausgabe}"))
            if fmt in ("JPEG", "WEBP"):
                img_res.save(ziel, format=fmt, quality=int(qualitaet), **meta)
            else:
                img_res.save(ziel, format=fmt, **meta)
            return ziel


def konvertiere_bild(quelle, ziel, ziel_ext, qualitaet, strip_exif=False):
    if ext_von(quelle) == "heic" and not HEIC_AVAILABLE:
        raise RuntimeError("HEIC-Unterstützung ist in dieser Programmversion nicht enthalten")
    with Image.open(quelle) as img0:
        img = _gedreht(img0)
        if ziel_ext == "svg":
            with io.BytesIO() as puffer:
                img.convert("RGBA").save(puffer, format="PNG")
                b64 = base64.b64encode(puffer.getvalue()).decode("ascii")
            w, h = img.size
            svg_inhalt = (f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
                          f'width="{w}" height="{h}" viewBox="0 0 {w} {h}">'
                          f'<image width="{w}" height="{h}" xlink:href="data:image/png;base64,{b64}"/></svg>')
            with open(ziel, "w", encoding="utf-8") as f_out:
                f_out.write(svg_inhalt)
            return

        fmt = BILD_FORMATE.get(ziel_ext, "JPEG")
        if fmt == "HEIF" and not HEIC_AVAILABLE:
            raise RuntimeError("HEIC-Unterstützung ist in dieser Programmversion nicht enthalten")
        meta = _meta_args(img, fmt, strip_exif)

        if fmt in ("JPEG", "PDF", "BMP") and img.mode in ("RGBA", "LA", "P"):
            with img.convert("RGBA") as rgba:
                with Image.new("RGB", rgba.size, (255, 255, 255)) as hg:
                    hg.paste(rgba, mask=rgba.split()[-1])
                    img = hg
        elif img.mode not in ("RGB", "RGBA", "L", "P"):
            img = img.convert("RGB")

        if fmt == "ICO":
            img.save(ziel, format=fmt, sizes=[(256, 256), (128, 128), (64, 64), (32, 32), (16, 16)])
        elif fmt in ("JPEG", "WEBP", "HEIF"):
            img.save(ziel, format=fmt, quality=int(qualitaet), **meta)
        elif ziel_ext == "pdf (komp)":
            w, h = img.size
            max_dim = 1920
            if max(w, h) > max_dim:
                ratio = max_dim / float(max(w, h))
                with img.resize((int(w * ratio), int(h * ratio)), Image.LANCZOS) as img_resized:
                    img_resized.save(ziel, format="PDF", quality=min(int(qualitaet), 75), optimize=True)
            else:
                img.save(ziel, format="PDF", quality=min(int(qualitaet), 75), optimize=True)
        elif fmt == "PDF":
            img.save(ziel, format="PDF", quality=int(qualitaet), optimize=True)
        else:
            img.save(ziel, format=fmt, **meta)


def bild_als_png_temp(quelle):
    """Normalisiert ein beliebiges Bild (auch HEIC/WEBP) zu einer temporären PNG für Office-COM."""
    pfad = os.path.join(tempfile.mkdtemp(prefix="kreis_"), "bild.png")
    with Image.open(quelle) as img:
        img = _gedreht(img)
        if img.mode not in ("RGB", "RGBA", "L"):
            img = img.convert("RGBA")
        img.save(pfad, format="PNG")
    return pfad


# ------------------------------------------------------------------ PDF
def _pdf_pruefen(doc):
    if getattr(doc, "needs_pass", False):
        raise RuntimeError("Das PDF ist passwortgeschützt")


def _pypdf_oeffnen(f_in, name=""):
    import pypdf
    reader = pypdf.PdfReader(f_in)
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise RuntimeError(f"{name or 'Das PDF'} ist passwortgeschützt")
    return reader


def konvertiere_pdf_teilen(quelle, zielordner):
    import pypdf
    with open(quelle, "rb") as f_in:
        reader = _pypdf_oeffnen(f_in)
        basis = os.path.splitext(os.path.basename(quelle))[0]
        target_dir = os.path.join(zielordner, f"{basis}_Seiten")
        os.makedirs(target_dir, exist_ok=True)
        stellen = len(str(len(reader.pages)))
        for idx, seite in enumerate(reader.pages, 1):
            writer = pypdf.PdfWriter()
            writer.add_page(seite)
            ziel_pfad = freier_pfad(os.path.join(target_dir, f"{basis}_Seite_{idx:0{stellen}d}.pdf"))
            with open(ziel_pfad, "wb") as f_out:
                writer.write(f_out)
    return target_dir


@fitz_sicher
def konvertiere_pdf_zu_bild(quelle, zielordner, dpi=150):
    if fitz is None:
        raise RuntimeError("Für 'PDF zu Bild' fehlt das PDF-Modul (PyMuPDF)")
    basis = os.path.splitext(os.path.basename(quelle))[0]
    target_dir = os.path.join(zielordner, f"{basis}_Bilder")
    os.makedirs(target_dir, exist_ok=True)
    with fitz.open(quelle) as doc:
        _pdf_pruefen(doc)
        stellen = len(str(len(doc)))
        for idx, seite in enumerate(doc, 1):
            zoom = dpi / 72.0
            laengste = max(seite.rect.width, seite.rect.height) * zoom
            if laengste > 6000:                       # Poster-Seiten nicht in Gigabyte-Bilder verwandeln
                zoom *= 6000 / laengste
            pix = seite.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            pix.save(freier_pfad(os.path.join(target_dir, f"{basis}_Seite_{idx:0{stellen}d}.png")))
    return target_dir


def konvertiere_pdf_zusammenfuehren(dateien, zielordner):
    import pypdf
    from contextlib import ExitStack
    writer = pypdf.PdfWriter()
    os.makedirs(zielordner, exist_ok=True)
    with ExitStack() as stack:            # Quelldateien bleiben bis zum Schreiben geöffnet
        for d in sorted(dateien, key=natur_key):
            if ext_von(d) != "pdf":
                continue
            reader = _pypdf_oeffnen(stack.enter_context(open(d, "rb")), os.path.basename(d))
            for page in reader.pages:
                writer.add_page(page)
        if len(writer.pages) == 0:
            raise RuntimeError("Die PDF-Dateien enthalten keine Seiten")
        ziel_pfad = freier_pfad(os.path.join(zielordner, "Zusammengefuehrt.pdf"))
        try:
            with open(ziel_pfad, "wb") as f_out:
                writer.write(f_out)
        except BaseException:
            raeume_auf(ziel_pfad)
            raise
    return ziel_pfad


@fitz_sicher
def komprimiere_pdf(quelle, ziel, qualitaet=85):
    q = int(qualitaet)
    if fitz is not None:
        with fitz.open(quelle) as doc:
            _pdf_pruefen(doc)
            try:    # Bilder in zu hoher Auflösung herunterrechnen (bringt bei Scans/Fotos am meisten)
                dpi = 150 if q >= 95 else (120 if q >= 85 else 100)
                doc.rewrite_images(dpi_threshold=dpi + 30, dpi_target=dpi, quality=max(40, min(85, q - 15)))
            except Exception:
                pass
            doc.save(ziel, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    else:
        import pypdf
        with open(quelle, "rb") as f_in:
            reader = _pypdf_oeffnen(f_in)
            writer = pypdf.PdfWriter()
            for seite in reader.pages:
                writer.add_page(seite)
            for seite in writer.pages:
                try:
                    seite.compress_content_streams()
                except Exception:
                    pass
            with open(ziel, "wb") as f_out:
                writer.write(f_out)
    if os.path.getsize(ziel) >= os.path.getsize(quelle):    # Komprimieren darf nie größer machen
        shutil.copyfile(quelle, ziel)


# ------------------------------------------------------------------ Texterkennung (OCR)
# Nutzt die in Windows 10/11 eingebaute Texterkennung (über die winrt-Pakete, die in der exe stecken).
# Es muss nichts installiert werden; benötigt wird nur ein Windows-OCR-Sprachpaket (bei deutschem Windows vorhanden).
def _ocr_engine():
    from winrt.windows.media.ocr import OcrEngine
    from winrt.windows.globalization import Language
    for tag in ("de-DE", "de", "en-US", "en"):
        try:
            sprache = Language(tag)
            if OcrEngine.is_language_supported(sprache):
                engine = OcrEngine.try_create_from_language(sprache)
                if engine is not None:
                    return engine
        except Exception:
            continue
    engine = OcrEngine.try_create_from_user_profile_languages()
    if engine is None:
        raise RuntimeError("Windows-Texterkennung nicht verfügbar: Bitte in den Windows-Einstellungen unter "
                           "Zeit & Sprache > Sprache das Sprachpaket mit Texterkennung (OCR) aktivieren")
    return engine


async def _ocr_async(engine, rgba, breite, hoehe):
    from winrt.windows.storage.streams import DataWriter
    from winrt.windows.graphics.imaging import SoftwareBitmap, BitmapPixelFormat
    writer = DataWriter()
    writer.write_bytes(rgba)
    bitmap = SoftwareBitmap.create_copy_from_buffer(writer.detach_buffer(), BitmapPixelFormat.RGBA8, breite, hoehe)
    ergebnis = await engine.recognize_async(bitmap)
    return "\n".join(zeile.text for zeile in ergebnis.lines)


def ocr_bild(img):
    """Liest den Text aus einem PIL-Bild."""
    if os.name != "nt":
        raise RuntimeError("Texterkennung (OCR) ist nur unter Windows 10/11 verfügbar")
    try:
        import winrt.windows.media.ocr  # noqa: F401
    except Exception:
        raise RuntimeError("Texterkennung ist in dieser Installation nicht enthalten")
    erg = {}

    def lauf():
        try:
            from winrt.windows.media.ocr import OcrEngine
            engine = _ocr_engine()
            grenze = int(OcrEngine.max_image_dimension)
            bild = img.convert("RGBA")
            if bild.getchannel("A").getextrema()[0] < 255:       # Transparenz -> weißer Hintergrund
                hg = Image.new("RGBA", bild.size, (255, 255, 255, 255))
                hg.alpha_composite(bild)
                bild = hg
            w, h = bild.size
            faktor = 2.0 if max(w, h) < 1000 else 1.0           # kleine Bilder vergrößern: bessere Erkennung
            if max(w, h) * faktor > grenze:
                faktor = grenze / float(max(w, h))
            if faktor != 1.0:
                bild = bild.resize((max(1, int(w * faktor)), max(1, int(h * faktor))), Image.LANCZOS)
            erg["text"] = asyncio.run(_ocr_async(engine, bild.tobytes(), bild.width, bild.height))
        except BaseException as ex:
            erg["fehler"] = ex

    t = threading.Thread(target=lauf, daemon=True)    # eigener Thread: eigene COM-Umgebung für WinRT
    t.start()
    t.join()
    if "fehler" in erg:
        ex = erg["fehler"]
        if isinstance(ex, RuntimeError):
            raise ex
        raise RuntimeError(f"Texterkennung fehlgeschlagen: {ex}")
    return erg.get("text", "").strip()


@fitz_sicher
def _pdf_seitentexte(quelle):
    """Text jeder Seite (leer bei Seiten ohne Textebene, z. B. Scans)."""
    if fitz is not None:
        with fitz.open(quelle) as doc:
            _pdf_pruefen(doc)
            return [seite.get_text("text").strip() for seite in doc]
    with open(quelle, "rb") as f_in:
        return [(seite.extract_text() or "").strip() for seite in _pypdf_oeffnen(f_in).pages]


def konvertiere_pdf_ocr(quelle, ziel, abbruch=None, melde=None):
    """Text aus PDF holen. Seiten ohne Textebene (Scans) werden per Windows-Texterkennung gelesen."""
    seiten = _pdf_seitentexte(quelle)
    leer = [i for i, t in enumerate(seiten) if len(t) < 3]
    if leer:
        if fitz is None:
            raise RuntimeError("Gescanntes PDF: Die Texterkennung braucht das PDF-Modul (PyMuPDF)")
        for n, i in enumerate(leer, 1):
            if abbruch is not None and abbruch.is_set():
                raise Abgebrochen()
            with FITZ_LOCK, fitz.open(quelle) as doc:
                seite = doc[i]
                zoom = 300 / 72.0
                laengste = max(seite.rect.width, seite.rect.height) * zoom
                if laengste > 6000:
                    zoom *= 6000 / laengste
                png = seite.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).tobytes("png")
            with Image.open(io.BytesIO(png)) as bild:
                seiten[i] = ocr_bild(bild)
            if melde:
                melde(0.98 * n / len(leer))
    text = [f"--- Seite {i} ---\n{t}" for i, t in enumerate(seiten, 1) if t]
    if not text:
        raise RuntimeError("Kein Text gefunden")
    with open(ziel, "w", encoding="utf-8") as f_out:
        f_out.write("\n\n".join(text))


def konvertiere_bild_ocr(quelle, ziel):
    with Image.open(quelle) as img0:
        text = ocr_bild(_gedreht(img0))
    if not text:
        raise RuntimeError("Kein Text im Bild erkannt")
    with open(ziel, "w", encoding="utf-8") as f_out:
        f_out.write(text)


def qr_bild(text):
    import qrcode
    try:
        return qrcode.make(text)
    except Exception as ex:
        if type(ex).__name__ == "DataOverflowError":
            raise RuntimeError("Der Text ist zu lang für einen QR-Code (max. ca. 2900 Zeichen)")
        raise


def konvertiere_qrcode(quelle, ziel):
    if ext_von(quelle) not in DATEN_EINGABE:
        raise RuntimeError("QR-Code braucht eine Textdatei (.txt, .md, .csv ...) mit dem Inhalt")
    text = _lies_text(quelle).strip()
    if not text:
        text = os.path.basename(quelle)
    qr_bild(text).save(ziel, format="PNG")


from xml.sax.saxutils import escape as xml_esc


# ------------------------------------------------------------------ Auftrag
class Auftrag:
    """Einstellungen und Rückmeldekanal für eine einzelne Konvertierung."""
    def __init__(self, ziel_ext, qual, strip_exif, ffmpeg, abbruch, melde=None):
        self.ziel_ext = ziel_ext
        self.qual = qual
        self.strip_exif = strip_exif
        self.ffmpeg = ffmpeg
        self.abbruch = abbruch
        self.melde = melde or (lambda frac: None)
        self.warnungen = []          # Zweifelsfälle: Datei wird erstellt, im Kreis erscheint eine Warnung


# ------------------------------------------------------------------ Daten laden
def _lies_text(pfad):
    with open(pfad, "rb") as f:
        roh = f.read()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return roh.decode(enc)
        except UnicodeDecodeError:
            continue
    return roh.decode("utf-8", errors="replace")


class _Semikolon(csv.excel):
    delimiter = ";"


def _tabelle_aus_csv(text):
    try:
        dialekt = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
    except csv.Error:
        dialekt = _Semikolon
    return [zeile for zeile in csv.reader(io.StringIO(text), dialekt)]


def _json_zu_tabelle(obj):
    """Liste von Objekten / Listen -> Tabelle, sonst None."""
    if not isinstance(obj, list) or not obj:
        return None
    if all(isinstance(e, dict) for e in obj):
        kopf = []
        for e in obj:
            for k in e:
                if k not in kopf:
                    kopf.append(k)
        zeilen = [kopf]
        for e in obj:
            zeilen.append([("" if e.get(k) is None else
                            (json.dumps(e[k], ensure_ascii=False) if isinstance(e[k], (dict, list)) else e[k]))
                           for k in kopf])
        return zeilen
    if all(isinstance(e, list) for e in obj):
        return [["" if c is None else c for c in e] for e in obj]
    if all(not isinstance(e, (dict, list)) for e in obj):
        return [[e] for e in obj]
    return None


def _xml_zu_dict(el):
    kinder = list(el)
    d = {}
    if el.attrib:
        d["@attribute"] = dict(el.attrib)
    text = (el.text or "").strip()
    if not kinder:
        if d:
            if text:
                d["#text"] = text
            return d
        return text
    for k in kinder:
        v = _xml_zu_dict(k)
        if k.tag in d:
            if not isinstance(d[k.tag], list):
                d[k.tag] = [d[k.tag]]
            d[k.tag].append(v)
        else:
            d[k.tag] = v
    return d


def _xml_laden(pfad):
    wurzel = ET.parse(pfad).getroot()
    kinder = list(wurzel)
    if kinder and len({k.tag for k in kinder}) == 1:
        flach = all(not list(k) or all(not list(g) for g in k) for k in kinder)
        if flach:
            kopf = []
            for k in kinder:
                for key in list(k.attrib) + [g.tag for g in k]:
                    if key not in kopf:
                        kopf.append(key)
            if kopf:
                zeilen = [kopf]
                for k in kinder:
                    werte = dict(k.attrib)
                    for g in k:
                        werte[g.tag] = (g.text or "").strip()
                    zeilen.append([werte.get(key, "") for key in kopf])
                return ("tabelle", zeilen)
    return ("baum", {wurzel.tag: _xml_zu_dict(wurzel)})


def _xlsx_zeilen(pfad):
    import openpyxl
    wb = openpyxl.load_workbook(pfad, read_only=True, data_only=True)
    try:
        ws = wb.active
        zeilen = [["" if c is None else c for c in z] for z in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    while zeilen and not any(str(c).strip() for c in zeilen[-1]):
        zeilen.pop()
    return zeilen


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _docx_text(pfad):
    with zipfile.ZipFile(pfad) as zf:
        wurzel = ET.fromstring(zf.read("word/document.xml"))
    zeilen = []
    for p in wurzel.iter(W_NS + "p"):
        teile = []
        for el in p.iter():
            if el.tag == W_NS + "t":
                teile.append(el.text or "")
            elif el.tag == W_NS + "tab":
                teile.append("\t")
            elif el.tag == W_NS + "br":
                teile.append("\n")
        zeile = "".join(teile).strip()
        if zeile:
            zeilen.append(zeile)
    return zeilen


def _pptx_text(pfad):
    zeilen = []
    with zipfile.ZipFile(pfad) as zf:
        folien = sorted((n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        for nr, name in enumerate(folien, 1):
            zeilen.append(f"[Folie {nr}]")
            wurzel = ET.fromstring(zf.read(name))
            for p in wurzel.iter(A_NS + "p"):
                zeile = "".join(t.text or "" for t in p.iter(A_NS + "t")).strip()
                if zeile:
                    zeilen.append(zeile)
    return zeilen


@fitz_sicher
def _pdf_zeilen(quelle):
    """Text eines PDFs zeilenweise (pypdf zuerst, bei Misserfolg PyMuPDF)."""
    zeilen = []
    try:
        with open(quelle, "rb") as f_in:
            reader = _pypdf_oeffnen(f_in)
            for seite in reader.pages:
                t = seite.extract_text()
                if t:
                    zeilen.extend(z.strip() for z in t.splitlines() if z.strip())
    except RuntimeError:
        raise
    except Exception:
        zeilen = []
    if not zeilen and fitz is not None:
        with fitz.open(quelle) as doc:
            _pdf_pruefen(doc)
            for seite in doc:
                zeilen.extend(z.strip() for z in seite.get_text("text").splitlines() if z.strip())
    return zeilen


def _xls_zeilen(pfad):
    """Altes Excel-Format (.xls) ohne Excel lesen."""
    import xlrd
    buch = xlrd.open_workbook(pfad, on_demand=True)
    try:
        blatt = next((buch.sheet_by_index(i) for i in range(buch.nsheets) if buch.sheet_by_index(i).nrows), None)
        zeilen = []
        if blatt is not None:
            for r in range(blatt.nrows):
                zeile = []
                for c in range(blatt.ncols):
                    zelle = blatt.cell(r, c)
                    t, v = zelle.ctype, zelle.value
                    if t == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate_as_datetime(v, buch.datemode)
                        except Exception:
                            pass
                    elif t == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                        v = int(v)
                    elif t == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    elif t == xlrd.XL_CELL_ERROR:
                        v = "#FEHLER"
                    elif t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        v = ""
                    zeile.append(v)
                zeilen.append(zeile)
    finally:
        buch.release_resources()
    while zeilen and not any(str(c).strip() for c in zeilen[-1]):
        zeilen.pop()
    return zeilen


def _doc_text(pfad):
    """Text aus einem alten Word-Dokument (.doc, Word 97-2003) ohne Word lesen."""
    import olefile
    import struct
    ole = olefile.OleFileIO(pfad)
    try:
        wd = ole.openstream("WordDocument").read()
        if len(wd) < 0x1AA or struct.unpack_from("<H", wd, 0)[0] != 0xA5EC:
            raise RuntimeError("Kein gültiges Word-97-Dokument")
        flags = struct.unpack_from("<H", wd, 0x0A)[0]
        if flags & 0x0100:
            raise RuntimeError("Das Word-Dokument ist passwortgeschützt")
        tabelle = ole.openstream("1Table" if flags & 0x0200 else "0Table").read()
    finally:
        ole.close()
    ccp_text = struct.unpack_from("<i", wd, 0x4C)[0]
    fc_clx, lcb_clx = struct.unpack_from("<II", wd, 0x1A2)
    clx = tabelle[fc_clx:fc_clx + lcb_clx]
    i = 0
    while i < len(clx) and clx[i] == 1:
        i += 3 + struct.unpack_from("<H", clx, i + 1)[0]
    if i >= len(clx) or clx[i] != 2:
        raise RuntimeError("Word-Dokument hat eine unbekannte Struktur")
    lcb = struct.unpack_from("<I", clx, i + 1)[0]
    plc = clx[i + 5:i + 5 + lcb]
    n = (lcb - 4) // 12
    cps = struct.unpack_from(f"<{n + 1}I", plc, 0)
    teile = []
    for k in range(n):
        fc = struct.unpack_from("<I", plc, 4 * (n + 1) + 8 * k + 2)[0]
        anzahl = cps[k + 1] - cps[k]
        if fc & 0x40000000:
            start = (fc & 0x3FFFFFFF) // 2
            teile.append(wd[start:start + anzahl].decode("cp1252", "replace"))
        else:
            teile.append(wd[fc:fc + 2 * anzahl].decode("utf-16-le", "replace"))
    text = "".join(teile)[:max(0, ccp_text)]
    for _ in range(5):                              # Feldfunktionen (Seitenzahl, Verweise ...) entfernen
        neu = re.sub(r"\x13[^\x13\x14\x15]*\x14", "", text)
        neu = re.sub(r"\x13[^\x13\x15]*\x15", "", neu)
        if neu == text:
            break
        text = neu
    text = (text.replace("\x07\x07", "\n").replace("\r", "\n").replace("\x0b", "\n").replace("\x0c", "\n")
            .replace("\x07", "\t").replace("\x1e", "-").replace("\x1f", ""))
    text = _UNGUELTIG.sub("", text)
    return [z.strip() for z in text.split("\n") if z.strip()]


def _ppt_folien(pfad):
    """Text aus einer alten PowerPoint-Datei (.ppt) ohne PowerPoint lesen -> Liste von Folien (je Liste von Zeilen)."""
    import olefile
    import struct
    ole = olefile.OleFileIO(pfad)
    try:
        if not ole.exists("PowerPoint Document"):
            raise RuntimeError("Kein gültiges PowerPoint-97-Dokument")
        daten = ole.openstream("PowerPoint Document").read()
    finally:
        ole.close()

    def atom_text(typ, start, laenge):
        roh = daten[start:start + laenge]
        t = roh.decode("utf-16-le", "replace") if typ == 0x0FA0 else roh.decode("cp1252", "replace")
        return [z.strip() for z in _UNGUELTIG.sub("", t.replace("\r", "\n").replace("\x0b", "\n")).split("\n") if z.strip()]

    persist, folien_b = [], []      # Text aus SlideListWithText bzw. aus den Folien selbst

    def lauf(start, ende, ziel_liste, in_slidelist):
        pos = start
        while pos + 8 <= ende:
            ver_inst, typ, laenge = struct.unpack_from("<HHI", daten, pos)
            ab = pos + 8
            if ab + laenge > ende:
                break
            if ver_inst & 0x0F == 0x0F:
                if typ == 0x03EE:                      # Slide
                    folien_b.append([])
                    lauf(ab, ab + laenge, folien_b[-1], False)
                elif typ == 0x0FF0:                    # SlideListWithText (Instanz 0 = Folien)
                    lauf(ab, ab + laenge, ziel_liste, (ver_inst >> 4) == 0)
                else:
                    lauf(ab, ab + laenge, ziel_liste, in_slidelist)
            elif typ == 0x03F3 and in_slidelist:       # SlidePersistAtom -> nächste Folie
                persist.append([])
            elif typ in (0x0FA0, 0x0FA8):
                if in_slidelist and persist:
                    persist[-1].extend(atom_text(typ, ab, laenge))
                elif ziel_liste is not persist and ziel_liste is not None:
                    ziel_liste.extend(atom_text(typ, ab, laenge))
            pos = ab + laenge

    lauf(0, len(daten), None, False)
    folien = []
    for i in range(max(len(persist), len(folien_b))):
        a = persist[i] if i < len(persist) else []
        b = folien_b[i] if i < len(folien_b) else []
        folien.append(a + [t for t in b if t not in a])
    if not any(folien):
        raise RuntimeError("Kein Text in der PowerPoint-Datei gefunden")
    return folien


def _ppt_text(pfad):
    zeilen = []
    for nr, folie in enumerate(_ppt_folien(pfad), 1):
        zeilen.append(f"[Folie {nr}]")
        zeilen.extend(folie)
    return zeilen


def _rtf_text(pfad):
    from striprtf.striprtf import rtf_to_text
    with open(pfad, "rb") as f:
        roh = f.read().decode("cp1252", errors="replace")
    zeilen = []
    for z in rtf_to_text(roh, errors="ignore").splitlines():
        z = re.sub("[\uf000-\uf8ff]", "•", z)                  # Aufzählungszeichen aus Symbol-Schriften
        if z.endswith("|") and z.count("|") > 1:                 # Tabellenzeile
            z = z.rstrip("|").replace("|", "\t")
        if z.strip():
            zeilen.append(z.strip())
    return zeilen


ODF_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _odt_text(pfad):
    with zipfile.ZipFile(pfad) as zf:
        wurzel = ET.fromstring(zf.read("content.xml"))

    def inhalt(el):
        teile = [el.text or ""]
        for k in el:
            if k.tag == ODF_TEXT + "s":
                teile.append(" " * int(k.get(ODF_TEXT + "c", "1")))
            elif k.tag == ODF_TEXT + "tab":
                teile.append("\t")
            elif k.tag == ODF_TEXT + "line-break":
                teile.append("\n")
            else:
                teile.append(inhalt(k))
            teile.append(k.tail or "")
        return "".join(teile)

    zeilen = []
    for el in wurzel.iter():
        if el.tag in (ODF_TEXT + "p", ODF_TEXT + "h"):
            for z in inhalt(el).split("\n"):
                if z.strip():
                    zeilen.append(z.strip())
    return zeilen


def _epub_text(pfad):
    import posixpath
    from html.parser import HTMLParser
    from urllib.parse import unquote

    class _Text(HTMLParser):
        BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "blockquote"}

        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.teile, self.skip = [], 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style"):
                self.skip += 1
            elif tag in self.BLOCK:
                self.teile.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style"):
                self.skip = max(0, self.skip - 1)
            elif tag in self.BLOCK:
                self.teile.append("\n")

        def handle_data(self, data):
            if not self.skip:
                self.teile.append(data)

    zeilen = []
    with zipfile.ZipFile(pfad) as zf:
        container = ET.fromstring(zf.read("META-INF/container.xml"))
        opf_pfad = container.find(".//{*}rootfile").get("full-path")
        opf = ET.fromstring(zf.read(opf_pfad))
        basis = posixpath.dirname(opf_pfad)
        manifest = {e.get("id"): e.get("href") for e in opf.iter() if e.tag.split("}")[-1] == "item"}
        for ref in (e.get("idref") for e in opf.iter() if e.tag.split("}")[-1] == "itemref"):
            href = manifest.get(ref)
            if not href:
                continue
            name = posixpath.normpath(posixpath.join(basis, unquote(href.split("#")[0])))
            try:
                roh = zf.read(name)
            except KeyError:
                continue
            parser = _Text()
            parser.feed(roh.decode("utf-8", errors="replace"))
            parser.close()
            zeilen.extend(z.strip() for z in "".join(parser.teile).splitlines() if z.strip())
    return zeilen


def lade_daten(quelle):
    """Liest eine Datei in eine gemeinsame Form: ('tabelle', zeilen) | ('text', zeilen) | ('baum', objekt)."""
    ext = ext_von(quelle)
    if ext == "csv":
        return ("tabelle", _tabelle_aus_csv(_lies_text(quelle)))
    if ext == "json":
        obj = json.loads(_lies_text(quelle))
        t = _json_zu_tabelle(obj)
        return ("tabelle", t) if t else ("baum", obj)
    if ext == "xml":
        return _xml_laden(quelle)
    if ext in ("txt", "md", "log"):
        return ("text", [z.strip() for z in _lies_text(quelle).splitlines() if z.strip()])
    if ext == "xlsx":
        return ("tabelle", _xlsx_zeilen(quelle))
    if ext == "xls":
        try:
            return ("tabelle", _xls_zeilen(quelle))
        except Exception as ex_xlrd:
            tmp_dir = tempfile.mkdtemp(prefix="kreis_")
            try:
                tmp = os.path.join(tmp_dir, "tabelle.xlsx")
                try:
                    _excel_konvertieren(quelle, tmp, "xlsx")
                except Exception:
                    raise RuntimeError(f"Die .xls-Datei konnte nicht gelesen werden ({ex_xlrd})")
                return ("tabelle", _xlsx_zeilen(tmp))
            finally:
                shutil.rmtree(tmp_dir, ignore_errors=True)
    if ext == "docx":
        return ("text", _docx_text(quelle))
    if ext == "pptx":
        return ("text", _pptx_text(quelle))
    if ext == "ppt":
        try:
            return ("text", _ppt_text(quelle))
        except Exception as ex_ole:
            tmp_dir = tempfile.mkdtemp(prefix="kreis_")
            try:
                tmp = os.path.join(tmp_dir, "folien.pptx")
                try:
                    _powerpoint_konvertieren(quelle, tmp, "pptx")
                except Exception:
                    raise RuntimeError(f"Die .ppt-Datei konnte nicht gelesen werden ({ex_ole})")
                return ("text", _pptx_text(tmp))
            finally:
                shutil.rmtree(tmp_dir, ignore_errors=True)
    if ext in ("doc", "rtf", "odt"):
        leser = {"doc": _doc_text, "rtf": _rtf_text, "odt": _odt_text}[ext]
        try:
            return ("text", leser(quelle))
        except Exception as ex_intern:
            try:
                return ("text", [z.strip() for z in _word_zu_text(quelle).splitlines() if z.strip()])
            except Exception:
                raise RuntimeError(f"Die .{ext}-Datei konnte nicht gelesen werden ({ex_intern})")
    if ext == "epub":
        return ("text", _epub_text(quelle))
    if ext == "pdf":
        zeilen = _pdf_zeilen(quelle)
        if not zeilen:
            raise RuntimeError("Kein Text im PDF gefunden (gescannt?) – 'txt (ocr)' verwenden")
        return ("text", zeilen)
    raise RuntimeError(f".{ext} kann nicht gelesen werden")


# ------------------------------------------------------------------ Daten schreiben
_UNGUELTIG = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]")


def _sauber(text):
    return _UNGUELTIG.sub("", str(text))


def _xml_name(s, i):
    s = re.sub(r"[^\w.-]", "_", str(s).strip())
    if not s:
        return f"spalte{i}"
    if not (s[0].isalpha() or s[0] == "_"):
        s = "_" + s
    return s


def _hat_kopf(zeilen):
    if len(zeilen) < 2:
        return False
    kopf = [str(c).strip() for c in zeilen[0]]
    return all(kopf) and len(set(kopf)) == len(kopf) and all(isinstance(c, str) for c in zeilen[0])


def _als_text(wert):
    return "" if wert is None else str(wert)


def _tabelle_zu_zeilentext(zeilen):
    return ["; ".join(_als_text(c) for c in z) for z in zeilen]


def _zahl(v):
    """Text, der wie eine Zahl aussieht, als Zahl speichern (Excel rechnet damit), führende Nullen bleiben Text."""
    if isinstance(v, str) and re.fullmatch(r"-?(0|[1-9]\d{0,14})([.,]\d+)?", v.strip()):
        try:
            return float(v.strip().replace(",", ".")) if re.search(r"[.,]", v) else int(v.strip())
        except ValueError:
            return v
    return v


def _baum_zu_xml(tag, val, parent=None):
    el = ET.SubElement(parent, tag) if parent is not None else ET.Element(tag)
    if isinstance(val, dict):
        for k, v in val.items():
            if k == "@attribute" and isinstance(v, dict):
                el.attrib.update({_xml_name(a, 0): _sauber(b) for a, b in v.items()})
            elif k == "#text":
                el.text = _sauber(v)
            elif isinstance(v, list):
                for item in v:
                    _baum_zu_xml(_xml_name(k, 0), item, el)
            else:
                _baum_zu_xml(_xml_name(k, 0), v, el)
    elif isinstance(val, list):
        for item in val:
            _baum_zu_xml("eintrag", item, el)
    else:
        el.text = _sauber(_als_text(val))
    return el


def _xml_speichern(el, ziel):
    if hasattr(ET, "indent"):
        ET.indent(el)
    ET.ElementTree(el).write(ziel, encoding="utf-8", xml_declaration=True)


def _als_zeilen(daten):
    typ, inhalt = daten
    if typ == "tabelle":
        return _tabelle_zu_zeilentext(inhalt)
    if typ == "text":
        return list(inhalt)
    return json.dumps(inhalt, indent=2, ensure_ascii=False, default=str).splitlines()


def schreibe_epub(zeilen, ziel, titel):
    import uuid
    uid = "urn:uuid:" + str(uuid.uuid4())
    t = xml_esc(_sauber(titel))
    absaetze = "".join(f"<p>{xml_esc(_sauber(z))}</p>" for z in zeilen) or "<p></p>"
    container = ('<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                 '</rootfiles></container>')
    opf = (f'<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" '
           f'unique-identifier="BookId" version="2.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           f'<dc:title>{t}</dc:title><dc:creator>Kreis-Konverter</dc:creator><dc:language>de</dc:language>'
           f'<dc:identifier id="BookId">{uid}</dc:identifier></metadata><manifest>'
           f'<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>'
           f'<item id="page" href="page.xhtml" media-type="application/xhtml+xml"/></manifest>'
           f'<spine toc="ncx"><itemref idref="page"/></spine></package>')
    ncx = (f'<?xml version="1.0" encoding="utf-8"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">'
           f'<head><meta name="dtb:uid" content="{uid}"/></head><docTitle><text>{t}</text></docTitle>'
           f'<navMap><navPoint id="n1" playOrder="1"><navLabel><text>{t}</text></navLabel>'
           f'<content src="page.xhtml"/></navPoint></navMap></ncx>')
    seite = (f'<?xml version="1.0" encoding="utf-8"?><!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.1//EN" '
             f'"http://www.w3.org/TR/xhtml11/DTD/xhtml11.dtd"><html xmlns="http://www.w3.org/1999/xhtml">'
             f'<head><title>{t}</title></head><body><h1>{t}</h1>{absaetze}</body></html>')
    with zipfile.ZipFile(ziel, "w") as zf:
        # EPUB verlangt: 'mimetype' als erster Eintrag und unkomprimiert
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        for name, inhalt in (("META-INF/container.xml", container), ("OEBPS/content.opf", opf),
                             ("OEBPS/toc.ncx", ncx), ("OEBPS/page.xhtml", seite)):
            zf.writestr(name, inhalt, compress_type=zipfile.ZIP_DEFLATED)


def _schreibe_docx_einfach(zeilen, ziel):
    """Schreibt ein schlichtes, gültiges .docx – ohne Word (reiner Text, ein Absatz pro Zeile)."""
    def p(t):
        return '<w:p><w:r><w:t xml:space="preserve">' + xml_esc(_sauber(t)) + '</w:t></w:r></w:p>'
    body = "".join(p(z) for z in zeilen) or p("")
    doc = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
           + body + '<w:sectPr/></w:body></w:document>')
    typen = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
             'wordprocessingml.document.main+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            'officeDocument" Target="word/document.xml"/></Relationships>')
    with zipfile.ZipFile(ziel, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", typen)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", doc)


def schreibe_xlsx(daten, ziel):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    typ, inhalt = daten
    if typ == "tabelle":
        zeilen = inhalt
    elif typ == "text":
        zeilen = [["Inhalt"]] + [[z] for z in inhalt]
    else:
        raise RuntimeError("Verschachtelte Struktur lässt sich nicht als Tabelle (XLSX) speichern")
    wb = openpyxl.Workbook()
    ws = wb.active
    for z in zeilen:
        ws.append([_zahl(_sauber(c)) if isinstance(c, str) else c for c in z])
    kopf = typ == "text" or _hat_kopf(zeilen)
    if kopf and ws.max_row >= 1:
        for zelle in ws[1]:
            zelle.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for spalte in ws.columns:
        breite = max((len(str(c.value)) for c in spalte if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(spalte[0].column)].width = min(max(breite + 2, 8), 60)
    wb.save(ziel)


def schreibe_daten(daten, ziel, ziel_ext, name):
    typ, inhalt = daten
    if ziel_ext == "csv":
        if typ == "baum":
            raise RuntimeError("Verschachtelte Struktur lässt sich nicht als Tabelle (CSV) speichern")
        with open(ziel, "w", encoding="utf-8-sig", newline="") as f:   # BOM: Excel zeigt Umlaute richtig an
            w = csv.writer(f, delimiter=";")
            if typ == "tabelle":
                w.writerows(inhalt)
            else:
                w.writerow(["Index", "Inhalt"])
                for i, z in enumerate(inhalt, 1):
                    w.writerow([i, z])
    elif ziel_ext == "json":
        if typ == "tabelle":
            if _hat_kopf(inhalt):
                kopf = [str(c).strip() for c in inhalt[0]]
                obj = [dict(zip(kopf, z)) for z in inhalt[1:]]
            else:
                obj = inhalt
        elif typ == "text":
            obj = {"datei": name, "zeilen": inhalt}
        else:
            obj = inhalt
        with open(ziel, "w", encoding="utf-8") as f:
            json.dump(obj, f, indent=2, ensure_ascii=False, default=str)
    elif ziel_ext == "xml":
        if typ == "tabelle":
            wurzel = ET.Element("tabelle", {"name": _sauber(name)})
            if _hat_kopf(inhalt):
                namen = [_xml_name(c, i) for i, c in enumerate(inhalt[0])]
                daten_zeilen = inhalt[1:]
            else:
                namen = None
                daten_zeilen = inhalt
            for z in daten_zeilen:
                ze = ET.SubElement(wurzel, "zeile")
                for i, c in enumerate(z):
                    ET.SubElement(ze, namen[i] if namen and i < len(namen) else f"spalte{i + 1}").text = _sauber(_als_text(c))
        elif typ == "text":
            wurzel = ET.Element("dokument", {"name": _sauber(name)})
            for z in inhalt:
                ET.SubElement(wurzel, "zeile").text = _sauber(z)
        else:
            if isinstance(inhalt, dict) and len(inhalt) == 1:
                tag, val = next(iter(inhalt.items()))
            else:
                tag, val = "wurzel", inhalt
            wurzel = _baum_zu_xml(_xml_name(tag, 0), val)
        _xml_speichern(wurzel, ziel)
    elif ziel_ext == "md":
        with open(ziel, "w", encoding="utf-8") as f:
            if typ == "tabelle" and inhalt:
                breite = max(len(z) for z in inhalt)
                def zeile(z):
                    z = list(z) + [""] * (breite - len(z))
                    return "| " + " | ".join(_als_text(c).replace("|", "\\|").replace("\n", " ") for c in z) + " |"
                f.write(zeile(inhalt[0]) + "\n| " + " | ".join(["---"] * breite) + " |\n")
                for z in inhalt[1:]:
                    f.write(zeile(z) + "\n")
            elif typ == "baum":
                f.write("```json\n" + json.dumps(inhalt, indent=2, ensure_ascii=False, default=str) + "\n```\n")
            else:
                f.write("# " + name + "\n\n" + "\n\n".join(_als_zeilen((typ, inhalt))) + "\n")
    elif ziel_ext == "txt":
        with open(ziel, "w", encoding="utf-8") as f:
            f.write("\n".join(_als_zeilen((typ, inhalt))))
    elif ziel_ext == "epub":
        schreibe_epub(_als_zeilen((typ, inhalt)), ziel, name)
    elif ziel_ext == "docx":
        schreibe_docx_daten((typ, inhalt), ziel)
    elif ziel_ext == "pdf":
        daten_zu_pdf((typ, inhalt), ziel, name)
    elif ziel_ext == "xlsx":
        schreibe_xlsx((typ, inhalt), ziel)
    else:
        raise RuntimeError(f"Zielformat .{ziel_ext} wird hier nicht unterstützt")


def konvertiere_daten(quelle, ziel, ziel_ext):
    schreibe_daten(lade_daten(quelle), ziel, ziel_ext, os.path.basename(quelle))


# ------------------------------------------------------------------ Eingebaute Dokument-Konverter (ohne Office)
PDF_CSS = """
body { font-family: sans-serif; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } h3 { font-size: 13pt; } h4, h5, h6 { font-size: 11pt; }
p { margin: 0 0 5pt 0; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 8pt 0; }
td, th { border: 0.6pt solid #777777; padding: 3pt; font-size: 9pt; }
th { background-color: #dde6f3; }
"""
SEITENUMBRUCH = "<!--SEITE-->"


def _html_esc(text):
    return xml_esc(_sauber(text))


@fitz_sicher
def _html_zu_pdf(html, ziel, querformat=False, archiv=None, rand=50, schrift=11, css=PDF_CSS):
    """Rendert HTML (mit SEITENUMBRUCH-Markern) in ein PDF – ohne Word/Office."""
    if fitz is None:
        raise RuntimeError("PDF-Erzeugung braucht das PDF-Modul (PyMuPDF)")
    seite = fitz.paper_rect("a4-l" if querformat else "a4")
    feld = seite + (rand, rand, -rand, -rand)
    writer = fitz.DocumentWriter(ziel)
    seiten = 0
    try:
        for teil in html.split(SEITENUMBRUCH):
            if not teil.strip():
                continue
            story = fitz.Story(html="<html><body>" + teil + "</body></html>", user_css=css, em=schrift, archive=archiv)
            weiter = 1
            while weiter:
                dev = writer.begin_page(seite)
                weiter, _ = story.place(feld)
                story.draw(dev)
                writer.end_page()
                seiten += 1
                if seiten > 3000:
                    raise RuntimeError("Das Dokument ist zu groß für die PDF-Erzeugung")
        if seiten == 0:
            writer.begin_page(seite)
            writer.end_page()
    finally:
        writer.close()


def _zeilen_html(zeilen):
    out = []
    for z in zeilen:
        z = _sauber(z).expandtabs(4)
        if not z.strip():
            out.append("<p>&#160;</p>")
            continue
        lead = len(z) - len(z.lstrip(" "))
        out.append("<p>" + "&#160;" * lead + xml_esc(z.lstrip(" ")) + "</p>")
    return "".join(out)


def _wert_text(v):
    import datetime
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%d.%m.%Y" if (v.hour, v.minute, v.second) == (0, 0, 0) else "%d.%m.%Y %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        return ("%.10g" % v).replace(".", ",")
    return str(v)


def _tabelle_html(zeilen, kopf=False):
    if not zeilen:
        return ""
    spalten = max(len(z) for z in zeilen)
    out = ["<table>"]
    for i, z in enumerate(zeilen):
        z = list(z) + [""] * (spalten - len(z))
        tag = "th" if (kopf and i == 0) else "td"
        zellen = []
        for c in z:
            rechts = isinstance(c, (int, float)) and not isinstance(c, bool)
            stil = ' style="text-align:right"' if rechts else ""
            zellen.append(f"<{tag}{stil}>{_html_esc(_wert_text(c)) or '&#160;'}</{tag}>")
        out.append("<tr>" + "".join(zellen) + "</tr>")
    out.append("</table>")
    return "".join(out)


def daten_zu_pdf(daten, ziel, name=""):
    typ, inhalt = daten
    if typ == "tabelle":
        breit = max((len(z) for z in inhalt), default=0) > 6
        _html_zu_pdf(_tabelle_html(inhalt, _hat_kopf(inhalt)), ziel, querformat=breit)
    elif typ == "text":
        _html_zu_pdf(_zeilen_html(inhalt), ziel)
    else:
        _html_zu_pdf(_zeilen_html(json.dumps(inhalt, indent=2, ensure_ascii=False, default=str).splitlines()), ziel)


def _tabellen_blaetter(quelle):
    """[(Titel, Zeilen)] aller Blätter einer .xlsx/.xls."""
    blaetter = []
    if ext_von(quelle) == "xls":
        import xlrd
        buch = xlrd.open_workbook(quelle, on_demand=True)
        try:
            for i in range(buch.nsheets):
                b = buch.sheet_by_index(i)
                zeilen = []
                for r in range(b.nrows):
                    zeile = []
                    for c in range(b.ncols):
                        t, v = b.cell_type(r, c), b.cell_value(r, c)
                        if t == xlrd.XL_CELL_DATE:
                            try:
                                v = xlrd.xldate_as_datetime(v, buch.datemode)
                            except Exception:
                                pass
                        elif t == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                            v = int(v)
                        zeile.append("" if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else v)
                    zeilen.append(zeile)
                blaetter.append((b.name, zeilen))
        finally:
            buch.release_resources()
    else:
        import openpyxl
        wb = openpyxl.load_workbook(quelle, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                blaetter.append((ws.title, [["" if c is None else c for c in z] for z in ws.iter_rows(values_only=True)]))
        finally:
            wb.close()
    ergebnis = []
    for titel, zeilen in blaetter:
        while zeilen and not any(str(c).strip() for c in zeilen[-1]):
            zeilen.pop()
        if zeilen:
            ergebnis.append((titel, zeilen))
    return ergebnis


def tabelle_zu_pdf(quelle, ziel):
    blaetter = _tabellen_blaetter(quelle)
    if not blaetter:
        raise RuntimeError("Die Tabelle enthält keine Daten")
    breit = any(max(len(z) for z in zeilen) > 6 for _, zeilen in blaetter)
    teile = [f"<h2>{_html_esc(t)}</h2>" + _tabelle_html(z) for t, z in blaetter]
    _html_zu_pdf(SEITENUMBRUCH.join(teile), ziel, querformat=breit)


# ---- Word (.docx) -> HTML -> PDF
def _docx_html(pfad):
    import docx
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    d = docx.Document(pfad)
    archiv = fitz.Archive()
    zaehler = {"bild": 0, "nummer": 0}
    W_P, W_TBL, W_R = qn("w:p"), qn("w:tbl"), qn("w:r")

    def bild_html(elem):
        blips = elem.findall(".//" + qn("a:blip"))
        if not blips:
            return ""
        rid = blips[0].get(qn("r:embed"))
        try:
            blob = d.part.related_parts[rid].blob
            with Image.open(io.BytesIO(blob)) as im:
                if im.format not in ("JPEG", "PNG"):
                    puffer = io.BytesIO()
                    im.convert("RGBA").save(puffer, "PNG")
                    blob, endung = puffer.getvalue(), "png"
                else:
                    endung = "jpg" if im.format == "JPEG" else "png"
                breite_px, hoehe_px = im.size
        except Exception:
            return ""
        ext = elem.find(".//" + qn("wp:extent"))
        breite = int(ext.get("cx")) / 12700.0 if ext is not None else breite_px * 0.75
        breite = max(10.0, min(breite, 480.0))
        zaehler["bild"] += 1
        name = f"bild{zaehler['bild']}.{endung}"
        archiv.add(blob, name)
        return f'<img src="{name}" style="width:{breite:.0f}pt">'

    def eigene_runs(p):
        for r in p.iter(W_R):
            a = r.getparent()
            while a is not None and a.tag != W_P:
                a = a.getparent()
            if a is p:
                yield r

    def absatz(p):
        para = Paragraph(p, d.part)
        try:
            stil = (para.style.name or "") if para.style is not None else ""
        except Exception:
            stil = ""
        teile = []
        if p.find(".//" + qn("w:pageBreakBefore")) is not None:
            teile.append(SEITENUMBRUCH)
        for r in eigene_runs(p):
            inner = []
            for k in r:
                if k.tag == qn("w:t"):
                    inner.append(xml_esc(_sauber(k.text or "")))
                elif k.tag == qn("w:tab"):
                    inner.append("&#160;" * 4)
                elif k.tag == qn("w:br"):
                    if k.get(qn("w:type")) == "page":
                        inner.append("</span>" + SEITENUMBRUCH + "<span>")
                    else:
                        inner.append("<br/>")
                elif k.tag == qn("w:drawing"):
                    inner.append(bild_html(k))
            text = "".join(inner)
            if not text:
                continue
            css = []
            rpr = r.find(qn("w:rPr"))
            def an(name):
                e = rpr.find(qn(name)) if rpr is not None else None
                return e is not None and e.get(qn("w:val")) not in ("0", "false", "none")
            sz = rpr.find(qn("w:sz")) if rpr is not None else None
            if sz is not None and sz.get(qn("w:val")):
                css.append(f"font-size:{int(sz.get(qn('w:val'))) / 2:g}pt")
            col = rpr.find(qn("w:color")) if rpr is not None else None
            if col is not None and (col.get(qn("w:val")) or "").lower() not in ("", "auto"):
                css.append("color:#" + col.get(qn("w:val")))
            if an("w:b"):
                css.append("font-weight:bold")
            if an("w:i"):
                css.append("font-style:italic")
            if an("w:u"):
                css.append("text-decoration:underline")
            teile.append(f'<span style="{";".join(css)}">{text}</span>' if css else f"<span>{text}</span>")
        inhalt = "".join(teile)
        ppr = p.find(qn("w:pPr"))
        liste = ppr is not None and ppr.find(qn("w:numPr")) is not None
        liste = liste or stil.startswith("List")
        ausrichtung = ""
        try:
            ausr = {1: "center", 2: "right", 3: "justify"}.get(int(para.alignment) if para.alignment is not None else 0)
            if ausr:
                ausrichtung = f"text-align:{ausr};"
        except Exception:
            pass
        if not liste:
            zaehler["nummer"] = 0
        if not inhalt.replace("<span>", "").replace("</span>", "").strip():
            return inhalt + "<p>&#160;</p>" if SEITENUMBRUCH not in inhalt else inhalt
        if stil.startswith("Heading") or stil in ("Title", "Subtitle"):
            ebene = 1 if stil in ("Title",) else (2 if stil == "Subtitle" else min(6, int(stil.split()[-1]) if stil.split()[-1].isdigit() else 1))
            return f'<h{ebene} style="{ausrichtung}">{inhalt}</h{ebene}>'
        if liste:
            if "Number" in stil:
                zaehler["nummer"] += 1
                marke = f"{zaehler['nummer']}."
            else:
                marke = "•"
            return f'<p style="margin-left:18pt;{ausrichtung}">{marke}&#160;{inhalt}</p>'
        return f'<p style="{ausrichtung}">{inhalt}</p>'

    def tabelle(tbl):
        zeilen = []
        for tr in tbl.findall(qn("w:tr")):
            zellen = []
            for tc in tr.findall(qn("w:tc")):
                tcpr = tc.find(qn("w:tcPr"))
                span = tcpr.find(qn("w:gridSpan")) if tcpr is not None else None
                colspan = f' colspan="{span.get(qn("w:val"))}"' if span is not None else ""
                inhalt = "".join(absatz(p) for p in tc.findall(qn("w:p"))) or "&#160;"
                zellen.append(f"<td{colspan}>{inhalt}</td>")
            zeilen.append("<tr>" + "".join(zellen) + "</tr>")
        return "<table>" + "".join(zeilen) + "</table>"

    teile = []
    for kind in d.element.body.iterchildren():
        if kind.tag == W_P:
            teile.append(absatz(kind))
        elif kind.tag == W_TBL:
            teile.append(tabelle(kind))
    return "".join(teile), archiv


@fitz_sicher
def docx_zu_pdf(quelle, ziel):
    html, archiv = _docx_html(quelle)
    _html_zu_pdf(html or "<p>&#160;</p>", ziel, archiv=archiv)


def text_dokument_zu_pdf(quelle, ziel):
    """doc / rtf / odt / epub / ppt: reiner Text (ohne Formatierung) als PDF."""
    daten = lade_daten(quelle)
    daten_zu_pdf(daten, ziel, os.path.basename(quelle))


# ---- PowerPoint -> PDF (vereinfachte Darstellung: Texte, Bilder, Tabellen, einfache Formen)
def _farbe_rgb(farbe):
    try:
        rgb = farbe.rgb
        return tuple(int(str(rgb)[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
    except Exception:
        return None


def _hell(rgb):
    return 0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]


def _folien_hintergrund(folie):
    from pptx.enum.dml import MSO_FILL
    for quelle in (folie, folie.slide_layout, folie.slide_layout.slide_master):
        try:
            fuell = quelle.background.fill
            if fuell.type == MSO_FILL.SOLID:
                rgb = _farbe_rgb(fuell.fore_color)
                if rgb:
                    return rgb
        except Exception:
            continue
    return (1.0, 1.0, 1.0)


@fitz_sicher
def _pptx_pdf_eigen(quelle, ziel):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.enum.dml import MSO_FILL
    from pptx.oxml.ns import qn
    if fitz is None:
        raise RuntimeError("PDF-Erzeugung braucht das PDF-Modul (PyMuPDF)")
    prs = Presentation(quelle)
    k = 72.0 / 914400.0
    doc = fitz.open()

    def text_html(shape, hg):
        gross = {PP_PLACEHOLDER.TITLE: 30, PP_PLACEHOLDER.CENTER_TITLE: 36, PP_PLACEHOLDER.SUBTITLE: 22}
        try:
            ph = shape.placeholder_format.type if shape.is_placeholder else None
        except Exception:
            ph = None
        standard = gross.get(ph, 18)
        mitte = ph in (PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE)
        body = ph in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
        teile = []
        for para in shape.text_frame.paragraphs:
            runs = []
            for r in para.runs:
                t = xml_esc(_sauber(r.text))
                if not t:
                    continue
                css = []
                if r.font.size:
                    css.append(f"font-size:{r.font.size.pt:g}pt")
                rgb = _farbe_rgb(r.font.color) if r.font.color and r.font.color.type is not None else None
                if rgb is None or abs(_hell(rgb) - _hell(hg)) < 0.35:
                    rgb = (0, 0, 0) if _hell(hg) > 0.5 else (1, 1, 1)
                css.append("color:#%02x%02x%02x" % tuple(int(c * 255) for c in rgb))
                if r.font.bold:
                    css.append("font-weight:bold")
                if r.font.italic:
                    css.append("font-style:italic")
                if r.font.underline:
                    css.append("text-decoration:underline")
                runs.append(f'<span style="{";".join(css)}">{t}</span>')
            if not runs:
                teile.append("<p>&#160;</p>")
                continue
            ppr = para._p.find(qn("a:pPr"))
            punkt = ppr is not None and ppr.find(qn("a:buChar")) is not None
            if body and not (ppr is not None and ppr.find(qn("a:buNone")) is not None):
                punkt = True
            ausr = "center" if mitte else {1: "left", 2: "center", 3: "right", 4: "justify"}.get(
                int(para.alignment) if para.alignment is not None else 0, "left")
            einzug = 16 * (para.level or 0) + (14 if punkt else 0)
            vorne = "•&#160;" if punkt else ""
            fg = (0, 0, 0) if _hell(hg) > 0.5 else (1, 1, 1)
            teile.append(f'<p style="text-align:{ausr};margin:0 0 4pt {einzug}pt;'
                         f'color:#%02x%02x%02x;font-size:{standard}pt">{vorne}{"".join(runs)}</p>' % tuple(int(c * 255) for c in fg))
        return "<body style='font-family:sans-serif'>" + "".join(teile) + "</body>"

    def zeichne(seite, shapes, hg, t):
        ax, bx, ay, by = t
        for sh in shapes:
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    xf = sh._element.find(".//" + qn("a:xfrm"))
                    off, ext = xf.find(qn("a:off")), xf.find(qn("a:ext"))
                    choff, chext = xf.find(qn("a:chOff")), xf.find(qn("a:chExt"))
                    sx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
                    sy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
                    zeichne(seite, sh.shapes, hg, (
                        ax * sx, ax * (int(off.get("x")) - int(choff.get("x")) * sx) + bx,
                        ay * sy, ay * (int(off.get("y")) - int(choff.get("y")) * sy) + by))
                    continue
                if sh.left is None or sh.width is None:
                    continue
                x0 = (sh.left * ax + bx) * k
                y0 = (sh.top * ay + by) * k
                rect = fitz.Rect(x0, y0, x0 + sh.width * ax * k, y0 + sh.height * ay * k)
                if rect.is_empty or rect.is_infinite:
                    continue
                if getattr(sh, "has_table", False) and sh.has_table:
                    tbl = sh.table
                    y = rect.y0
                    for zeile in tbl.rows:
                        x = rect.x0
                        h = zeile.height * ay * k
                        for c, spalte in enumerate(tbl.columns):
                            b = spalte.width * ax * k
                            zr = fitz.Rect(x, y, x + b, y + h)
                            seite.draw_rect(zr, color=(0.5, 0.5, 0.5), width=0.5)
                            zelle = tbl.cell(list(tbl.rows).index(zeile), c)
                            if zelle.text.strip():
                                seite.insert_htmlbox(zr + (3, 2, -3, -2), "<body style='font-family:sans-serif'><p style='font-size:11pt;margin:0'>"
                                                     + xml_esc(_sauber(zelle.text)).replace("\n", "<br/>") + "</p></body>", scale_low=0)
                            x += b
                        y += h
                    continue
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    blob = sh.image.blob
                    if sh.crop_left or sh.crop_right or sh.crop_top or sh.crop_bottom:
                        with Image.open(io.BytesIO(blob)) as im:
                            w, h = im.size
                            im2 = im.crop((int(w * sh.crop_left), int(h * sh.crop_top),
                                           int(w * (1 - sh.crop_right)), int(h * (1 - sh.crop_bottom))))
                            puffer = io.BytesIO()
                            im2.convert("RGBA").save(puffer, "PNG")
                            blob = puffer.getvalue()
                    seite.insert_image(rect, stream=blob)
                    continue
                if sh.shape_type in (MSO_SHAPE_TYPE.LINE,) or sh.__class__.__name__ == "Connector":
                    seite.draw_line((rect.x0, rect.y0), (rect.x1, rect.y1), color=(0.4, 0.4, 0.4), width=1)
                    continue
                hg_shape = hg
                try:
                    if sh.fill.type == MSO_FILL.SOLID:
                        rgb = _farbe_rgb(sh.fill.fore_color)
                        if rgb:
                            seite.draw_rect(rect, color=None, fill=rgb)
                            hg_shape = rgb
                except Exception:
                    pass
                if sh.has_text_frame and sh.text_frame.text.strip():
                    seite.insert_htmlbox(rect + (4, 3, -4, -3), text_html(sh, hg_shape), scale_low=0)
            except Exception:
                logge("PPTX-Shape übersprungen:\n" + traceback.format_exc())

    for folie in prs.slides:
        seite = doc.new_page(width=prs.slide_width * k, height=prs.slide_height * k)
        hg = _folien_hintergrund(folie)
        seite.draw_rect(seite.rect, color=None, fill=hg)
        zeichne(seite, folie.shapes, hg, (1, 0, 1, 0))
    if doc.page_count == 0:
        raise RuntimeError("Die Präsentation enthält keine Folien")
    doc.save(ziel, garbage=3, deflate=True)
    doc.close()


def ppt_zu_pdf(quelle, ziel):
    """Alte .ppt ohne PowerPoint: eine Seite je Folie mit dem Text."""
    teile = []
    for nr, zeilen in enumerate(_ppt_folien(quelle), 1):
        kopf = f"<h2>{_html_esc(zeilen[0])}</h2>" if zeilen else f"<h2>Folie {nr}</h2>"
        teile.append(kopf + _zeilen_html(zeilen[1:]))
    _html_zu_pdf(SEITENUMBRUCH.join(teile), ziel, querformat=True)


# ---- Schreiben: DOCX (python-docx), PDF -> DOCX, Bild -> DOCX/PPTX
def schreibe_docx_daten(daten, ziel):
    try:
        import docx
        from docx.shared import Cm
    except ImportError:
        return _schreibe_docx_einfach(_als_zeilen(daten), ziel)
    typ, inhalt = daten
    d = docx.Document()
    sek = d.sections[0]
    sek.page_width, sek.page_height = Cm(21.0), Cm(29.7)
    sek.left_margin = sek.right_margin = sek.top_margin = sek.bottom_margin = Cm(2.0)
    if typ == "tabelle" and inhalt:
        spalten = max(len(z) for z in inhalt)
        t = d.add_table(rows=0, cols=spalten)
        t.style = "Table Grid"
        kopf = _hat_kopf(inhalt)
        for i, z in enumerate(inhalt):
            zellen = t.add_row().cells
            for j in range(spalten):
                zellen[j].text = _sauber(_wert_text(z[j])) if j < len(z) else ""
                if kopf and i == 0:
                    for r in zellen[j].paragraphs[0].runs:
                        r.bold = True
    else:
        for z in _als_zeilen(daten) or [""]:
            d.add_paragraph(_sauber(z))
    d.save(ziel)


@fitz_sicher
def pdf_zu_docx(quelle, ziel):
    if fitz is None:
        raise RuntimeError("PDF → DOCX braucht das PDF-Modul (PyMuPDF)")
    import docx
    from docx.shared import Pt, Cm
    d = docx.Document()
    sek = d.sections[0]
    sek.page_width, sek.page_height = Cm(21.0), Cm(29.7)
    sek.left_margin = sek.right_margin = sek.top_margin = sek.bottom_margin = Cm(2.0)
    hat_text = False
    with fitz.open(quelle) as pdf:
        _pdf_pruefen(pdf)
        for nr, seite in enumerate(pdf):
            if nr:
                d.add_page_break()
            for block in seite.get_text("dict", sort=True)["blocks"]:
                if block.get("type") != 0:
                    continue
                para = d.add_paragraph()
                for zi, zeile in enumerate(block["lines"]):
                    if zi:
                        para.add_run(" ")
                    for sp in zeile["spans"]:
                        t = _sauber(sp["text"])
                        if not t.strip() and not t:
                            continue
                        r = para.add_run(t)
                        r.font.size = Pt(max(6, min(40, round(sp["size"] * 2) / 2)))
                        r.bold = bool(sp["flags"] & 16)
                        r.italic = bool(sp["flags"] & 2)
                        hat_text = hat_text or bool(t.strip())
    if not hat_text:
        raise RuntimeError("Kein Text im PDF gefunden (gescannt?) – erst 'txt (ocr)' verwenden")
    d.save(ziel)


def bild_zu_docx(quelle, ziel):
    import docx
    from docx.shared import Cm
    png = bild_als_png_temp(quelle)
    try:
        with Image.open(png) as im:
            w, h = im.size
        breite = min(17.0, 24.0 * w / h)               # in cm: passt auf die A4-Seite
        d = docx.Document()
        sek = d.sections[0]
        sek.page_width, sek.page_height = Cm(21.0), Cm(29.7)
        sek.left_margin = sek.right_margin = sek.top_margin = sek.bottom_margin = Cm(2.0)
        d.add_picture(png, width=Cm(breite))
        d.save(ziel)
    finally:
        shutil.rmtree(os.path.dirname(png), ignore_errors=True)


def bild_zu_pptx(quelle, ziel):
    from pptx import Presentation
    from pptx.util import Inches
    png = bild_als_png_temp(quelle)
    try:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        folie = prs.slides.add_slide(prs.slide_layouts[6])
        with Image.open(png) as im:
            w, h = im.size
        faktor = min(prs.slide_width / w, prs.slide_height / h)
        bw, bh = int(w * faktor), int(h * faktor)
        folie.shapes.add_picture(png, (prs.slide_width - bw) // 2, (prs.slide_height - bh) // 2, bw, bh)
        prs.save(ziel)
    finally:
        shutil.rmtree(os.path.dirname(png), ignore_errors=True)


def ppt_zu_pptx(quelle, ziel):
    """Alte .ppt ohne PowerPoint: neue .pptx mit dem Text (Titel + Inhalt je Folie, ohne Layout/Bilder)."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for nr, zeilen in enumerate(_ppt_folien(quelle), 1):
        folie = prs.slides.add_slide(prs.slide_layouts[1])
        folie.shapes.title.text = _sauber(zeilen[0]) if zeilen else f"Folie {nr}"
        folie.placeholders[1].text_frame.text = "\n".join(_sauber(z) for z in zeilen[1:])
    prs.save(ziel)


def xls_zu_xlsx(quelle, ziel):
    """Altes .xls ohne Excel in .xlsx umwandeln (Werte, keine Formatierung)."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for titel, zeilen in _tabellen_blaetter(quelle) or [("Tabelle1", [[""]])]:
        ws = wb.create_sheet(_sauber(titel)[:31] or "Tabelle")
        for z in zeilen:
            ws.append([_sauber(c) if isinstance(c, str) else c for c in z])
    wb.save(ziel)



# ------------------------------------------------------------------ Office (COM, nur Windows + installiertes Office)
@contextmanager
def com_umgebung():
    if not OFFICE_BEVORZUGEN[0]:
        raise RuntimeError("Die Office-Anbindung ist ausgeschaltet")
    try:
        import pythoncom
        import win32com.client  # noqa: F401
    except ImportError:
        raise RuntimeError("Office-Funktion braucht Windows, 'pywin32' und Microsoft Office")
    pythoncom.CoInitialize()
    try:
        yield
    finally:
        pythoncom.CoUninitialize()


MSO_AUTOMATION_SECURITY_FORCE_DISABLE = 3     # Makros in geöffneten Dateien niemals ausführen


def _starte(programm, name):
    import win32com.client
    try:
        app = win32com.client.DispatchEx(programm)
    except Exception as ex:
        raise RuntimeError(f"{name} ist nicht verfügbar: {ex}")
    # Sicherheit: sofort nach dem Start Makros abschalten (msoAutomationSecurityForceDisable).
    # Klappt das nicht, läuft die Umwandlung trotzdem weiter (nur ins Log schreiben).
    try:
        app.AutomationSecurity = MSO_AUTOMATION_SECURITY_FORCE_DISABLE
    except Exception as ex:
        logge(f"{name}: AutomationSecurity konnte nicht gesetzt werden: {ex}")
    return app


def _beende(app):
    if app is not None:
        try:
            app.Quit()
        except Exception:
            pass


def _word_konvertieren(quelle, ziel, ziel_ext):
    with OFFICE_LOCK, com_umgebung():
        word = None
        try:
            word = _starte("Word.Application", "Word")
            word.Visible = False
            word.DisplayAlerts = 0
            doc = word.Documents.Open(os.path.abspath(quelle), ConfirmConversions=False,
                                      ReadOnly=True, AddToRecentFiles=False)
            try:
                if ziel_ext == "pdf":
                    doc.ExportAsFixedFormat(OutputFileName=os.path.abspath(ziel), ExportFormat=17)
                elif ziel_ext == "docx":
                    doc.SaveAs2(os.path.abspath(ziel), FileFormat=16)
                else:
                    raise RuntimeError(f"Word kann nicht nach .{ziel_ext} umwandeln")
            finally:
                doc.Close(False)
        finally:
            _beende(word)


def _word_zu_text(quelle):
    """Beliebiges Word-lesbares Dokument (doc, rtf, odt, pdf) als Klartext."""
    tmp_dir = tempfile.mkdtemp(prefix="kreis_")
    try:
        tmp = os.path.join(tmp_dir, "text.txt")
        with OFFICE_LOCK, com_umgebung():
            word = None
            try:
                word = _starte("Word.Application", "Word")
                word.Visible = False
                word.DisplayAlerts = 0
                doc = word.Documents.Open(os.path.abspath(quelle), ConfirmConversions=False,
                                          ReadOnly=True, AddToRecentFiles=False)
                try:
                    doc.SaveAs2(tmp, FileFormat=7)   # Unicode-Text
                finally:
                    doc.Close(False)
            finally:
                _beende(word)
        with open(tmp, "rb") as f:
            roh = f.read()
        for enc in ("utf-16", "utf-8-sig", "cp1252"):
            try:
                return roh.decode(enc)
            except UnicodeDecodeError:
                continue
        return roh.decode("utf-8", errors="replace")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _excel_konvertieren(quelle, ziel, ziel_ext):
    with OFFICE_LOCK, com_umgebung():
        excel = None
        try:
            excel = _starte("Excel.Application", "Excel")
            excel.Visible = False
            excel.DisplayAlerts = False
            try:
                excel.AskToUpdateLinks = False
            except Exception as ex:
                logge(f"Excel: AskToUpdateLinks konnte nicht gesetzt werden: {ex}")
            # UpdateLinks=0: externe Verknüpfungen nicht aktualisieren (keine Netzwerkzugriffe, keine Rückfragen)
            try:
                wb = excel.Workbooks.Open(os.path.abspath(quelle), UpdateLinks=0, ReadOnly=True,
                                          IgnoreReadOnlyRecommended=True, AddToMru=False, Notify=False)
            except TypeError:                # ältere Anbindung ohne benannte Argumente: gleiche Werte per Position
                wb = excel.Workbooks.Open(os.path.abspath(quelle), 0, True)
            try:
                if ziel_ext == "pdf":
                    wb.ExportAsFixedFormat(0, os.path.abspath(ziel))
                elif ziel_ext == "xlsx":
                    wb.SaveAs(os.path.abspath(ziel), FileFormat=51)
                else:
                    raise RuntimeError(f"Excel kann nicht nach .{ziel_ext} umwandeln")
            finally:
                wb.Close(False)
        finally:
            _beende(excel)


def _powerpoint_konvertieren(quelle, ziel, ziel_ext):
    with OFFICE_LOCK, com_umgebung():
        ppt = None
        try:
            ppt = _starte("PowerPoint.Application", "PowerPoint")
            try:
                ppt.DisplayAlerts = 1        # ppAlertsNone
            except Exception as ex:
                logge(f"PowerPoint: DisplayAlerts konnte nicht gesetzt werden: {ex}")
            pres = ppt.Presentations.Open(os.path.abspath(quelle), True, False, False)
            try:
                if ziel_ext == "pdf":
                    pres.SaveAs(os.path.abspath(ziel), 32)
                elif ziel_ext == "pptx":
                    pres.SaveAs(os.path.abspath(ziel), 24)
                else:
                    raise RuntimeError(f"PowerPoint kann nicht nach .{ziel_ext} umwandeln")
            finally:
                pres.Close()
        finally:
            _beende(ppt)


# ------------------------------------------------------------------ Audio / Video
def stoppe_prozesse():
    for p in list(PROZESSE):
        try:
            p.terminate()
        except Exception:
            pass


def konvertiere_medien(auf, quelle, ziel, ziel_ext):
    ffmpeg = finde_ffmpeg()
    if not ffmpeg:
        raise RuntimeError("Der eingebaute Medien-Konverter (ffmpeg) konnte nicht entpackt werden")
    q = int(auf.qual)
    gerade = "scale=trunc(iw/2)*2:trunc(ih/2)*2"      # libx264/VP9 brauchen gerade Bildmaße
    crf = max(18, min(30, round(18 + (100 - q) * 0.5)))
    rate = "320k" if q >= 95 else ("192k" if q >= 85 else "128k")
    cmd = [ffmpeg, "-y", "-hide_banner", "-nostdin", "-i", quelle]
    if ziel_ext in AUDIO_AUSGABE:
        cmd += ["-vn"]
        if ziel_ext == "mp3":
            cmd += ["-b:a", rate]
        elif ziel_ext in ("aac", "m4a"):
            cmd += ["-c:a", "aac", "-b:a", rate]
        elif ziel_ext == "ogg":
            cmd += ["-c:a", "libvorbis", "-q:a", str(max(3, min(9, round(4 + (q - 75) * 0.16))))]
    elif ziel_ext == "gif":
        cmd += ["-an", "-vf", "fps=12,scale=480:-1:flags=lanczos,split[s0][s1];[s0]palettegen[p];[s1][p]paletteuse"]
    elif ziel_ext == "mp4":
        cmd += ["-vf", gerade, "-c:v", "libx264", "-crf", str(crf), "-preset", "medium", "-pix_fmt", "yuv420p",
                "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart"]
    elif ziel_ext in ("mkv", "mov"):
        cmd += ["-vf", gerade, "-c:v", "libx264", "-crf", str(crf), "-preset", "medium", "-pix_fmt", "yuv420p",
                "-c:a", "aac", "-b:a", "160k"]
    elif ziel_ext == "webm":
        cmd += ["-vf", gerade, "-c:v", "libvpx-vp9", "-crf", str(min(40, crf + 8)), "-b:v", "0", "-row-mt", "1",
                "-pix_fmt", "yuv420p", "-c:a", "libopus"]
    elif ziel_ext == "avi":
        cmd += ["-vf", gerade, "-c:v", "mpeg4", "-q:v", "4", "-c:a", "libmp3lame", "-b:a", rate]
    cmd += ["-progress", "pipe:1", "-nostats", ziel]
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0

    with MEDIEN_SEM:
        if auf.abbruch.is_set():
            raise Abgebrochen()
        try:
            proc = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                                    encoding="utf-8", errors="replace", creationflags=flags)
        except OSError:                      # z. B. vom Virenscanner entfernt: einmal neu entpacken
            neu = finde_ffmpeg(neu_suchen=True)
            if not neu:
                raise RuntimeError("ffmpeg konnte nicht gestartet werden")
            cmd[0] = neu
            proc = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                                    encoding="utf-8", errors="replace", creationflags=flags)
        PROZESSE.add(proc)
        dauer, letzte = None, []
        try:
            for zeile in proc.stdout:
                zeile = zeile.strip()
                if auf.abbruch.is_set():
                    proc.terminate()
                    break
                if dauer is None:
                    m = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", zeile)
                    if m:
                        dauer = int(m.group(1)) * 3600 + int(m.group(2)) * 60 + float(m.group(3))
                if zeile.startswith(("out_time_us=", "out_time_ms=")):
                    try:
                        us = int(zeile.split("=", 1)[1])
                        if dauer:
                            auf.melde(max(0.0, min(0.99, us / 1e6 / dauer)))
                    except ValueError:
                        pass
                elif zeile and "=" not in zeile[:20]:
                    letzte.append(zeile)
                    del letzte[:-15]
            rc = proc.wait()
        finally:
            PROZESSE.discard(proc)
            if proc.poll() is None:
                proc.kill()
    if auf.abbruch.is_set():
        raise Abgebrochen()
    if rc != 0:
        meldung = next((z for z in reversed(letzte) if "Conversion failed" not in z), "")
        raise RuntimeError(meldung or "ffmpeg: unbekannter Fehler")


# ------------------------------------------------------------------ Archive
def zip_erstellen(items, ziel_zip, ausschluss, abbruch, melde=None):
    """Packt Dateien UND ganze Ordner (mit Struktur) in ein ZIP."""
    ausschluss = {os.path.abspath(a) for a in ausschluss}
    eintraege = []
    for it in items:
        if os.path.isdir(it):
            basis = os.path.basename(it.rstrip("/\\"))
            gefunden = False
            for root, dirs, files in os.walk(it):
                rel = os.path.relpath(root, it)
                arc_dir = basis if rel == "." else os.path.join(basis, rel)
                if not files and not dirs:
                    eintraege.append((root, arc_dir + "/"))
                for f in files:
                    voll = os.path.join(root, f)
                    if os.path.abspath(voll) in ausschluss:
                        continue
                    eintraege.append((voll, os.path.join(arc_dir, f)))
                    gefunden = True
            if not gefunden and not eintraege:
                eintraege.append((it, basis + "/"))
        else:
            eintraege.append((it, os.path.basename(it)))
    gesamt = max(1, len(eintraege))
    with zipfile.ZipFile(ziel_zip, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for i, (voll, arc) in enumerate(eintraege, 1):
            if abbruch.is_set():
                raise Abgebrochen()
            zf.write(voll, arc)
            if melde:
                melde(i / gesamt, os.path.basename(voll))


def zip_entpacken(quelle, zielordner):
    name = os.path.splitext(os.path.basename(quelle))[0]
    ziel = freier_ordner(os.path.join(zielordner, name))
    try:
        with zipfile.ZipFile(quelle) as zf:
            zf.extractall(ziel)
    except Exception:
        shutil.rmtree(ziel, ignore_errors=True)
        raise
    return ziel

# ------------------------------------------------------------------ 7-Zip (7z erstellen, 7z/rar entpacken)
# Verwendet wird das offizielle 7z.exe + 7z.dll (nur diese Kombination kann auch RAR entpacken).
# exe-Betrieb: beide Dateien stecken in der exe (Ordner "7zip") und werden einmalig nach
# %LOCALAPPDATA%\KreisKonverter\7zip kopiert. start.bat-Betrieb: Ordner "tools\7zip" neben dem Programm.
SIEBENZIP_LOCK = threading.Lock()
_SIEBENZIP_PFAD = [None]
ARCHIV_SEM = threading.Semaphore(2)          # 7-Zip arbeitet selbst mehrkernig
SIEBENZIP_DATEIEN = ("7z.exe", "7z.dll")
SIEBENZIP_FEHLT = ("7-Zip fehlt (7z.exe + 7z.dll). Bitte die exe neu mit build_exe.bat bauen "
                   "bzw. beide Dateien in den Ordner tools\\7zip legen")
_KEIN_PASSWORT = "-pKreisKonverter-ohne-Passwort"   # verhindert die Passwort-Abfrage (würde sonst hängen)


def _siebenzip_ordner_ok(ordner):
    return all(os.path.isfile(os.path.join(ordner, f)) for f in SIEBENZIP_DATEIEN)


def _kopiere_7zip(quelle_dir):
    """Kopiert 7z.exe/7z.dll (+ Lizenz) aus der exe einmalig nach %LOCALAPPDATA%\\KreisKonverter\\7zip."""
    import filecmp
    ziel_dir = os.path.join(daten_ordner(), "7zip")
    os.makedirs(ziel_dir, exist_ok=True)
    for name in SIEBENZIP_DATEIEN + ("License.txt",):
        q = os.path.join(quelle_dir, name)
        z = os.path.join(ziel_dir, name)
        if not os.path.isfile(q):
            continue
        if (os.path.isfile(z) and os.path.getsize(z) == os.path.getsize(q)
                and filecmp.cmp(q, z, shallow=False)):
            continue
        tmp = f"{z}.{os.getpid()}.tmp"
        try:
            shutil.copyfile(q, tmp)
            try:
                os.replace(tmp, z)
            except OSError:
                # z. B. läuft gerade eine andere Instanz mit genau dieser Datei: gleiche Größe genügt dann
                if not (os.path.isfile(z) and os.path.getsize(z) == os.path.getsize(q)):
                    raise
        finally:
            raeume_auf(tmp)
    return os.path.join(ziel_dir, "7z.exe")


def _suche_7zip():
    if os.name == "nt":
        if getattr(sys, "frozen", False):                            # 1. in die exe eingebettet
            eingebettet = os.path.join(ressourcen_ordner(), "7zip")
            if _siebenzip_ordner_ok(eingebettet):
                try:
                    return _kopiere_7zip(eingebettet)
                except Exception:
                    logge("7-Zip kopieren fehlgeschlagen:\n" + traceback.format_exc())
        for ordner in (os.path.join(basis_ordner(), "tools", "7zip"),   # 2. start.bat-Betrieb
                       os.path.join(basis_ordner(), "7zip")):
            if _siebenzip_ordner_ok(ordner):
                return os.path.join(ordner, "7z.exe")
        for wurzel in (os.environ.get("ProgramW6432"), os.environ.get("ProgramFiles")):  # 3. installiertes 7-Zip
            if wurzel and _siebenzip_ordner_ok(os.path.join(wurzel, "7-Zip")):
                return os.path.join(wurzel, "7-Zip", "7z.exe")
        return None
    for name in ("7z", "7zz"):                                       # nur Entwicklung/Tests (Linux, macOS)
        pfad = shutil.which(name)
        if pfad:
            return os.path.abspath(pfad)
    return None


def finde_7zip(neu_suchen=False):
    with SIEBENZIP_LOCK:
        if neu_suchen or not (_SIEBENZIP_PFAD[0] and os.path.isfile(_SIEBENZIP_PFAD[0])):
            _SIEBENZIP_PFAD[0] = _suche_7zip()
        return _SIEBENZIP_PFAD[0]


def _7z_fehler(zeilen, rc):
    """Übersetzt 7-Zip-Meldungen in verständliche Sätze."""
    text = "\n".join(zeilen)
    klein = text.lower()
    if "wrong password" in klein or "encrypted" in klein and "password" in klein:
        return "Das Archiv ist passwortgeschützt – das wird nicht unterstützt"
    if "missing volume" in klein or "unavailable data" in klein or "need next volume" in klein:
        return "Ein Teil des mehrteiligen Archivs fehlt (alle Teile in denselben Ordner legen)"
    if "unexpected end" in klein or "headers error" in klein or "data error" in klein or "crc failed" in klein:
        return "Das Archiv ist beschädigt oder unvollständig"
    if "can not open the file as archive" in klein or "cannot open the file as archive" in klein \
            or "is not archive" in klein:
        return "Die Datei ist beschädigt oder kein gültiges Archiv"
    if "duplicate filename" in klein:
        return "Zwei Dateien mit gleichem Namen – bitte getrennt packen"
    if "not enough space" in klein or "disk full" in klein or "no space" in klein:
        return "Der Datenträger ist voll"
    if "access is denied" in klein or "zugriff verweigert" in klein or "permission denied" in klein:
        return "Kein Zugriff: Datei oder Ordner ist gesperrt oder schreibgeschützt"
    if rc == 8:
        return "Nicht genug Arbeitsspeicher für dieses Archiv"
    fehler = [z for z in zeilen if z.upper().startswith("ERROR") or "error" in z.lower()]
    letzte = (fehler or zeilen or [""])[-1].strip()
    return f"7-Zip: {letzte}" if letzte else f"7-Zip ist mit Fehlercode {rc} beendet worden"


def _7z_lauf(argumente, abbruch, melde=None):
    """Startet 7z.exe (ohne Konsole, ohne Shell) und liefert (Rückgabecode, Ausgabezeilen).
    Der Prozess steht in PROZESSE -> Abbrechen beendet ihn sofort."""
    exe = finde_7zip()
    if not exe:
        raise RuntimeError(SIEBENZIP_FEHLT)
    cmd = [exe] + list(argumente)
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    with ARCHIV_SEM:
        if abbruch.is_set():
            raise Abgebrochen()
        try:
            proc = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                                    stderr=subprocess.STDOUT, creationflags=flags, shell=False)
        except OSError:                      # z. B. vom Virenscanner entfernt: einmal neu suchen/kopieren
            neu = finde_7zip(neu_suchen=True)
            if not neu:
                raise RuntimeError(SIEBENZIP_FEHLT)
            cmd[0] = neu
            try:
                proc = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                                        stderr=subprocess.STDOUT, creationflags=flags, shell=False)
            except OSError as ex:
                raise RuntimeError(f"7-Zip konnte nicht gestartet werden: {ex}")
        PROZESSE.add(proc)
        zeilen, rest = [], ""
        try:
            while True:
                stueck = proc.stdout.read1(4096) if hasattr(proc.stdout, "read1") else proc.stdout.read(4096)
                if not stueck:
                    break
                if abbruch.is_set():
                    proc.terminate()
                    break
                rest += stueck.decode("utf-8", errors="replace")
                teile = re.split(r"[\r\n\b]+", rest)
                rest = teile.pop()
                for t in teile:
                    t = t.strip()
                    if not t:
                        continue
                    m = re.match(r"(\d{1,3})%", t)
                    if m:
                        if melde:
                            melde(max(0.0, min(0.99, int(m.group(1)) / 100)))
                    else:
                        zeilen.append(t)
                        del zeilen[:-40]
            if rest.strip() and not re.match(r"\d{1,3}%", rest.strip()):
                zeilen.append(rest.strip())
            rc = proc.wait()
        finally:
            PROZESSE.discard(proc)
            if proc.poll() is None:
                proc.kill()
    if abbruch.is_set():
        raise Abgebrochen()
    return rc, zeilen


_7Z_STANDARD = ["-y", "-bb0", "-bsp1", "-sccUTF-8"]


def _7z_pfade_pruefen(quelle, abbruch):
    """Sicherheit: Archive mit absoluten Pfaden oder '..' werden nicht entpackt (Schutz vor 'Zip-Slip')."""
    rc, zeilen = _7z_lauf(["l", "-slt", _KEIN_PASSWORT] + _7Z_STANDARD + ["--", quelle], abbruch)
    if rc >= 2:
        raise RuntimeError(_7z_fehler(zeilen, rc))
    if any(z.startswith("ERROR") for z in zeilen):
        raise RuntimeError(_7z_fehler(zeilen, rc))
    eintraege = False                    # erst nach der Linie "----------" folgen die Archiv-Einträge
    for z in zeilen:
        if z.startswith("----------"):
            eintraege = True
            continue
        if not (eintraege and z.startswith("Path = ")):
            continue
        pfad = z[7:].strip().replace("\\", "/")
        if pfad.startswith("/") or re.match(r"[A-Za-z]:", pfad) or ".." in pfad.split("/"):
            raise RuntimeError("Das Archiv enthält unsichere Pfade und wird aus Sicherheitsgründen nicht entpackt")


def _archiv_name(quelle):
    name = os.path.splitext(os.path.basename(quelle))[0]
    return re.sub(r"\.part0*1$", "", name, flags=re.IGNORECASE) or name   # name.part1.rar -> name


def siebenzip_entpacken(quelle, zielordner, abbruch, melde=None):
    """Entpackt .7z, .rar (auch mehrteilig ab .part1.rar) und Sonder-ZIPs in einen neuen Ordner."""
    quelle = os.path.abspath(quelle)
    _7z_pfade_pruefen(quelle, abbruch)
    ziel = freier_ordner(os.path.join(zielordner, _archiv_name(quelle)))
    try:
        rc, zeilen = _7z_lauf(["x", "-o" + os.path.abspath(ziel), "-aoa", _KEIN_PASSWORT] + _7Z_STANDARD
                              + ["--", quelle], abbruch, melde)
        if rc >= 2:
            raise RuntimeError(_7z_fehler(zeilen, rc))
        if rc == 1:
            logge(f"7-Zip-Warnung beim Entpacken von {quelle}: {' | '.join(zeilen[-5:])}")
    except BaseException:
        shutil.rmtree(ziel, ignore_errors=True)
        raise
    return ziel


def siebenzip_erstellen(items, ziel_7z, abbruch, melde=None, ausschluss=()):
    """Packt Dateien und ganze Ordner (mit Struktur) in ein .7z-Archiv.
    Die Pfade gehen über eine Listendatei an 7-Zip (keine Längenbegrenzung der Befehlszeile)."""
    ziel_7z = os.path.abspath(ziel_7z)
    liste_dir = tempfile.mkdtemp(prefix="kreis_7z_")
    try:
        liste = os.path.join(liste_dir, "dateien.txt")
        with open(liste, "w", encoding="utf-8") as f:
            for it in items:
                f.write(os.path.abspath(it).rstrip("/\\") + "\n")
        # eigene (Hilfs-)Ausgabedateien nie mit einpacken, falls der Zielordner in einem gepackten Ordner liegt
        weg = ["-xr!" + os.path.basename(a) for a in dict.fromkeys([ziel_7z, *ausschluss])]
        # kein "--" hier: danach würde 7-Zip "@liste" als Dateinamen lesen (Archivpfad ist absolut, also sicher)
        rc, zeilen = _7z_lauf(["a", "-t7z", "-mx=5", "-mmt=on", "-ssw", "-scsUTF-8"] + weg + _7Z_STANDARD
                              + [ziel_7z, "@" + liste], abbruch, melde)
        if rc >= 2:
            raise RuntimeError(_7z_fehler(zeilen, rc))
        if rc == 1:      # Warnung = mindestens eine Datei fehlte/war nicht lesbar -> kein unvollständiges Archiv
            logge(f"7-Zip-Warnung beim Packen nach {ziel_7z}: {' | '.join(zeilen[-8:])}")
            raise RuntimeError("Einige Dateien konnten nicht gelesen werden (geöffnet, gesperrt oder gelöscht?)")
    finally:
        shutil.rmtree(liste_dir, ignore_errors=True)


def archiv_entpacken(quelle, zielordner, abbruch, melde=None):
    """ZIP mit dem eingebauten Python-Entpacker; 7z/rar (und ZIPs, die Python nicht kann) über 7-Zip."""
    if ext_von(quelle) == "zip":
        try:
            return zip_entpacken(quelle, zielordner)
        except (NotImplementedError, RuntimeError) as ex:      # z. B. Deflate64 oder verschlüsselt
            if not finde_7zip():
                if "encrypted" in str(ex).lower() or "password" in str(ex).lower():
                    raise RuntimeError("Das Archiv ist passwortgeschützt – das wird nicht unterstützt")
                raise RuntimeError("Dieses ZIP-Format wird ohne 7-Zip nicht unterstützt")
            logge(f"ZIP über 7-Zip entpackt ({ex})")
    if ext_von(quelle) not in ARCHIV_EINGABE:
        raise RuntimeError("Entpacken geht nur bei .zip, .7z und .rar")
    return siebenzip_entpacken(quelle, zielordner, abbruch, melde)


def ist_folgeteil(pfad):
    """name.part2.rar, name.part3.rar ... werden beim Entpacken von .part1.rar automatisch mitgenommen."""
    m = re.search(r"\.part(\d+)\.rar$", os.path.basename(pfad), re.IGNORECASE)
    return bool(m) and int(m.group(1)) > 1



# ------------------------------------------------------------------ Dispatcher: was ist überhaupt möglich?
DATEN_ZIELE = {"csv", "json", "xml", "txt", "md", "epub"}
TEXT_QUELLEN = DATEN_EINGABE | {"docx", "doc", "rtf", "odt", "xlsx", "xls", "pptx", "ppt", "pdf", "epub"}
WORD_QUELLEN = {"docx", "doc", "rtf", "odt"}


def _com_dann_intern(com_funktion, intern_funktion):
    """Ist Microsoft Office installiert, liefert es die originalgetreueste Umwandlung.
    Sonst (oder bei jedem Fehler) übernimmt der in die exe eingebaute Konverter."""
    def lauf(s, t, a):
        if OFFICE_BEVORZUGEN[0]:
            try:
                com_funktion(s, t)
                return
            except Abgebrochen:
                raise
            except Exception as ex:
                logge(f"Office nicht nutzbar, eingebauter Konverter wird verwendet: {ex}")
            raeume_auf(t)
        intern_funktion(s, t, a)
    return lauf


# ------------------------------------------------------------------ Erweiterungs-Schnittstelle
# Optional: Liegt das Modul ERWEITERUNGEN_MODUL neben kreis_konverter.py (bzw. ist es beim Bauen gezielt in
# die exe eingebunden), lädt es weitere Module. Fehlt es, bleibt das Programm genau so, wie es ist.
# Der Name steht absichtlich nur in einer Variable: PyInstaller bindet das Modul so nie von selbst ein.
# Die fertige exe lädt nur, was beim Bauen eingebunden wurde – nie Dateien, die neben der exe liegen.
#
# Das Modul ERWEITERUNGEN_MODUL enthält:  MODULE = ["name1", ...]   (Namen der Zusatzmodule)
# Jedes Zusatzmodul enthält eine Funktion  erweiterung(api) -> dict  mit:
#   "kategorie":  Name der neuen Kategorie (erscheint in Formatliste und bei den Pfeiltasten)
#   "nach":       (optional) Kategorie, hinter der sie eingereiht wird – sonst ganz hinten
#   "farben":     (Farbe1, Farbe2) als "#rrggbb" (Kennfarbe der Kategorie)
#   "symbol":     (optional) "pfeil", "check" oder "warn" für die Mitte des Kreises
#   "formate":    {format: {"quellen": {"pdf", ...}, "endung": "xlsx",
#                           "umwandeln": funktion(quelle, ziel, auftrag),
#                           "falsche_quelle": "Meldung bei unpassender Datei"}}
#   "erkennen":   (optional) funktion(pfad) -> format oder None. Ein Treffer hat Vorrang vor dem gewählten
#                 Format (das gewählte Format bleibt eingestellt; nicht bei ERKENNUNG_AUSNAHMEN)
#   "erkannt_text": (optional) kurzer Hinweis im Kreis, wenn erkannt wurde
#   "lizenz_hinweis": (optional) Zusatztext im Reiter "Kreis-Konverter" des Lizenzen-Fensters
# Fehler meldet umwandeln() als RuntimeError (klare Meldung), Zweifelsfälle als Text in auftrag.warnungen
# (Datei wird erstellt, im Kreis erscheint die Warnung). Abbrechen: auftrag.abbruch / api.Abgebrochen.
ERWEITERUNGEN_MODUL = "erweiterungen"


def _erweiterung_eintragen(erw):
    kat = str(erw["kategorie"])
    formate = dict(erw["formate"])
    if not formate or kat in KATEGORIEN or any(f in ALLE_FORMATE or f in ERW_FORMATE for f in formate):
        raise ValueError(f"Kategorie oder Format bereits vorhanden: {kat}")
    for f, info in formate.items():
        if not callable(info.get("umwandeln")):
            raise ValueError(f"{f}: umwandeln fehlt")
        ERW_FORMATE[f] = {"quellen": set(info.get("quellen") or ()), "endung": str(info.get("endung") or f),
                          "umwandeln": info["umwandeln"],
                          "falsche_quelle": str(info.get("falsche_quelle") or f"{kat}: Datei passt nicht")}
    FARBEN[kat] = tuple(erw.get("farben") or FARBEN["hover"])
    if erw.get("symbol") in ("pfeil", "check", "warn"):
        EMB_GLYPH[kat] = erw["symbol"]
    neu, nach = [], erw.get("nach")
    for k, v in KATEGORIEN.items():
        neu.append((k, v))
        if k == nach:
            neu.append((kat, list(formate)))
    if nach not in KATEGORIEN:
        neu.append((kat, list(formate)))
    KATEGORIEN.clear()
    KATEGORIEN.update(neu)
    ALLE_FORMATE[:] = [f for liste in KATEGORIEN.values() for f in liste]
    EMB_KEYS.add(kat)
    ERWEITERUNGEN.append(erw)


def lade_erweiterungen():
    """Lädt die optionalen Zusatzmodule. Fehlt das Modul, passiert nichts (keine Meldung, keine Spur)."""
    if not getattr(sys, "frozen", False):
        ordner = os.path.dirname(os.path.abspath(__file__))      # Start aus dem Quelltext: Programmordner
        if ordner not in sys.path:
            sys.path.insert(0, ordner)
    try:
        liste = importlib.import_module(ERWEITERUNGEN_MODUL)
    except ImportError as ex:
        if getattr(ex, "name", None) != ERWEITERUNGEN_MODUL:
            logge(f"Erweiterungen: {ex}")
        return
    except Exception:
        logge("Erweiterungen:\n" + traceback.format_exc())
        return
    api = types.SimpleNamespace(Abgebrochen=Abgebrochen, ext_von=ext_von, kuerze=kuerze, sauber=_sauber,
                                logge=logge)
    for name in list(getattr(liste, "MODULE", []) or []):
        try:
            _erweiterung_eintragen(importlib.import_module(str(name)).erweiterung(api))
        except Exception:
            logge(f"Erweiterung {name} nicht geladen:\n" + traceback.format_exc())


def erkanntes_format(pfad, ziel_ext):
    """Erkennt eine Erweiterung die Datei, gilt deren Format statt des gewählten: (format, erweiterung).
    Sonst None."""
    if ziel_ext in ERKENNUNG_AUSNAHMEN:
        return None
    for erw in ERWEITERUNGEN:
        erkennen = erw.get("erkennen")
        if not callable(erkennen) or ziel_ext in erw["formate"]:
            continue
        try:
            f = erkennen(pfad)
        except Exception:
            logge(f"Erkennung {erw['kategorie']}: {traceback.format_exc()}")
            continue
        if f in erw["formate"]:
            return f, erw
    return None


def waehle_konverter(q, z):
    """Liefert (funktion(quelle, ziel_pfad, auftrag), art) oder wirft RuntimeError mit einer klaren Begründung.
    Nur was hier registriert ist, wird ausgeführt – es gibt keine stillen Scheinerfolge mehr."""
    if q == z or (q, z) in GLEICH:
        raise RuntimeError("Quell- und Zielformat sind gleich")
    nein = RuntimeError(f".{q} → .{z} ist nicht möglich")

    if z in ERW_FORMATE:
        info = ERW_FORMATE[z]
        if q in info["quellen"]:
            return (lambda s, t, a: info["umwandeln"](s, t, a)), "normal"
        raise RuntimeError(info["falsche_quelle"])

    if q in MEDIEN_EINGABE:
        if z == "gif" and q in AUDIO_EINGABE:
            raise RuntimeError("Aus einer Audiodatei lässt sich kein GIF erzeugen")
        if z in MEDIEN_AUSGABE or z == "gif":
            return (lambda s, t, a: konvertiere_medien(a, s, t, z)), "medien"
        raise nein

    if q in BILD_EINGABE:
        if z in BILD_FORMATE:
            def bild(s, t, a):
                with BILD_SEM:
                    konvertiere_bild(s, t, z, a.qual, a.strip_exif)
            return bild, "bild"
        if z == "docx":
            return (lambda s, t, a: bild_zu_docx(s, t)), "normal"
        if z == "pptx":
            return (lambda s, t, a: bild_zu_pptx(s, t)), "normal"
        raise nein

    if q == "pdf":
        if z == "pdf (komp)":
            return (lambda s, t, a: komprimiere_pdf(s, t, a.qual)), "normal"
        if z == "docx":
            return (lambda s, t, a: pdf_zu_docx(s, t)), "normal"
        if z in DATEN_ZIELE:
            return (lambda s, t, a: konvertiere_daten(s, t, z)), "normal"
        if z in BILD_FORMATE:
            raise RuntimeError("Für PDF → Bild bitte 'pdf zu bild' wählen")
        raise nein

    if q in WORD_QUELLEN:
        if z == "pdf":
            intern = (lambda s, t, a: docx_zu_pdf(s, t)) if q == "docx" else (lambda s, t, a: text_dokument_zu_pdf(s, t))
            return _com_dann_intern(lambda s, t: _word_konvertieren(s, t, "pdf"), intern), "office"
        if z == "docx":
            return _com_dann_intern(lambda s, t: _word_konvertieren(s, t, "docx"),
                                    lambda s, t, a: konvertiere_daten(s, t, "docx")), "office"
        if z in DATEN_ZIELE | {"xlsx"}:
            return (lambda s, t, a: konvertiere_daten(s, t, z)), "normal"
        raise nein

    if q in ("xlsx", "xls"):
        if z == "pdf":
            return _com_dann_intern(lambda s, t: _excel_konvertieren(s, t, "pdf"),
                                    lambda s, t, a: tabelle_zu_pdf(s, t)), "office"
        if z == "xlsx" and q == "xls":
            return _com_dann_intern(lambda s, t: _excel_konvertieren(s, t, "xlsx"),
                                    lambda s, t, a: xls_zu_xlsx(s, t)), "office"
        if z in DATEN_ZIELE | {"docx"}:
            return (lambda s, t, a: konvertiere_daten(s, t, z)), "normal"
        raise nein

    if q in ("pptx", "ppt"):
        if z == "pdf":
            intern = (lambda s, t, a: _pptx_pdf_eigen(s, t)) if q == "pptx" else (lambda s, t, a: ppt_zu_pdf(s, t))
            return _com_dann_intern(lambda s, t: _powerpoint_konvertieren(s, t, "pdf"), intern), "office"
        if z == "pptx" and q == "ppt":
            return _com_dann_intern(lambda s, t: _powerpoint_konvertieren(s, t, "pptx"),
                                    lambda s, t, a: ppt_zu_pptx(s, t)), "office"
        if z in DATEN_ZIELE | {"docx"}:
            return (lambda s, t, a: konvertiere_daten(s, t, z)), "normal"
        raise nein

    if q in DATEN_EINGABE or q == "epub":
        if z in DATEN_ZIELE | {"docx", "xlsx", "pdf"}:
            if z == "xlsx" and q == "epub":
                raise nein
            return (lambda s, t, a: konvertiere_daten(s, t, z)), "normal"
        raise nein

    raise nein


def ziel_passt(q, z):
    """Passt eine Datei mit Endung q zum Ziel z? (gleiche Regeln wie bei der eigentlichen Umwandlung)"""
    if z.startswith("resize "):
        return q in BILD_EINGABE
    if z in ("pdf teilen", "pdf zu bild", "pdf zusammenführen"):
        return q == "pdf"
    if z == "txt (ocr)":
        return q == "pdf" or q in BILD_EINGABE
    if z == "qrcode":
        return q in DATEN_EINGABE
    if z in PACK_AKTIONEN:
        return True
    if z == "entpacken":
        return q in ARCHIV_EINGABE
    try:
        waehle_konverter(q, z)
        return True
    except RuntimeError:
        return False


def passende_formate(q, umsetzen=None):
    """Alle sichtbaren Formate, die für Endung q möglich sind.
    umsetzen(format) -> tatsächliches Ziel (z. B. Web-Preset: Bildformat -> 'resize 1200px').
    Reihenfolge: zuerst die "eigentlichen" Umwandlungen, Archiv und QR-Code zuletzt."""
    umsetzen = umsetzen or (lambda f: f)
    formate = [f for f in sichtbare_formate() if ziel_passt(q, umsetzen(f))]
    return sorted(formate, key=lambda f: kat_von_format(f) in ("Archiv", "QR") and f != "entpacken")


def lizenz_ordner():
    """Ordner 'lizenzen' (in der exe eingebettet bzw. neben kreis_konverter.py)."""
    for ordner in (os.path.join(ressourcen_ordner(), "lizenzen"), os.path.join(basis_ordner(), "lizenzen")):
        if os.path.isdir(ordner):
            return ordner
    return os.path.join(ressourcen_ordner(), "lizenzen")


QUELLCODE_URL = "https://github.com/jr-stack-bot/joschy_converter"
EIGENE_LIZENZ_NAME = "Kreis-Konverter"
EIGENE_LIZENZ_TEXT = f"""KREIS-KONVERTER
© JOSCHY – Lizenz: AGPL-3.0 (GNU Affero General Public License, Version 3)

Dieses Programm ist freie Software: Sie dürfen es unter den Bedingungen der AGPL-3.0
weitergeben und/oder verändern. Es wird OHNE JEDE GEWÄHRLEISTUNG bereitgestellt.

Quellcode (Angebot nach AGPL-3.0): Der vollständige Quellcode dieses Programms steht
kostenlos zur Verfügung unter
    {QUELLCODE_URL}

Enthaltene Fremdsoftware und deren Lizenzen: siehe Reiter "Übersicht" und die weiteren Reiter.
"""


def eigene_lizenz_datei():
    """Die Datei LICENSE (AGPL-3.0) – in die exe eingebettet bzw. neben kreis_konverter.py."""
    for pfad in (os.path.join(ressourcen_ordner(), "LICENSE"), os.path.join(basis_ordner(), "LICENSE"),
                 os.path.join(lizenz_ordner(), "AGPL-3.0.txt")):
        if os.path.isfile(pfad):
            return pfad
    return None


def eigener_lizenz_text():
    hinweise = [str(e["lizenz_hinweis"]) for e in ERWEITERUNGEN if e.get("lizenz_hinweis")]
    text = EIGENE_LIZENZ_TEXT + ("\n" + "\n".join(hinweise) + "\n" if hinweise else "")
    pfad = eigene_lizenz_datei()
    if pfad:
        try:
            with open(pfad, "r", encoding="utf-8-sig", errors="replace") as f:
                text += "\n" + "=" * 78 + "\n\n" + f.read()
        except OSError:
            pass
    return text


def lizenz_dateien():
    """[(Anzeigename, Pfad)] – die eigene Lizenz zuerst (Pfad None: Text aus eigener_lizenz_text),
    danach die Übersicht LIZENZEN.txt und die Lizenztexte der Fremdsoftware."""
    eigene = [(EIGENE_LIZENZ_NAME, None)]
    ordner = lizenz_ordner()
    try:
        namen = sorted(f for f in os.listdir(ordner) if f.lower().endswith(".txt"))
    except OSError:
        return eigene
    namen.sort(key=lambda f: f.upper() != "LIZENZEN.TXT")
    return eigene + [("Übersicht" if f.upper() == "LIZENZEN.TXT" else os.path.splitext(f)[0].replace("_", " "),
                      os.path.join(ordner, f)) for f in namen]


lade_erweiterungen()            # optional – ohne Zusatzmodul ändert sich nichts


class KreisApp:
    EFFEKT_STUFEN = ("voll", "sparsam", "aus")     # Menü "Effekte"; "aus" = Bild wie vor den neuen Ebenen

    @staticmethod
    def _effekte_aus_cfg(cfg):
        """Effekt-Stufe aus den Einstellungen; ohne Eintrag oder bei ungültigem Wert "voll"."""
        wert = cfg.get("effekte") if isinstance(cfg, dict) else None
        return wert if isinstance(wert, str) and wert in KreisApp.EFFEKT_STUFEN else "voll"

    def __init__(self):
        self.dnd = False
        dpi_bewusst_machen()               # vor dem ersten Fenster (scharfe Darstellung bei 125/150/200 %)
        try:
            if not DND_IMPORT:
                raise RuntimeError
            self.root = TkinterDnD.Tk()
            self.dnd = True
        except Exception:
            self.root = tk.Tk()
        r = self.root
        r.title("Kreis-Konverter (by Joschy)")
        r.report_callback_exception = self._tk_fehler
        r.overrideredirect(True)
        r.configure(bg=TRANS)
        try:
            r.attributes("-transparentcolor", TRANS)
        except tk.TclError:
            pass
        r.attributes("-topmost", True)
        
        cfg = lade_einstellungen()
        OFFICE_BEVORZUGEN[0] = bool(cfg.get("office_bevorzugen", False))
        theme0 = cfg.get("theme") if cfg.get("theme") in THEMES else STANDARD_THEME
        if cfg.get("design") != 2:          # einmalig auf das neue Reaktor-Design umstellen
            theme0 = STANDARD_THEME
        self.theme_var = tk.StringVar(value=theme0)
        setze_skalierung(ermittle_skalierung(r))
        mini0 = cfg.get("mini") is True
        self.size = P(S_MINI if mini0 else S_NORMAL)
        x, y = fenster_position(cfg.get("fenster"), self.size, r)      # zuletzt benutzte Stelle
        self._pos = (x, y)                 # linke obere Ecke in Bildschirmpixeln (zum Speichern)
        r.geometry(f"{self.size}x{self.size}+{x}+{y}")

        self.c = tk.Canvas(r, width=self.size, height=self.size, bg=TRANS, highlightthickness=0, bd=0)
        self.c.pack()
        self.g = Grafik(akt_theme=self.theme_var.get())
        self.bild_id = self.c.create_image(0, 0, anchor="nw")
        self.photo = None

        self.ffmpeg = None          # wird erst beim ersten Audio-/Video-Auftrag entpackt (schneller Start)
        self.q = queue.Queue()
        self.format = cfg.get("format") if cfg.get("format") in sichtbare_formate() else "png"
        ordner0 = cfg.get("ordner")
        self.ordner = ordner0 if isinstance(ordner0, str) and ordner0 and os.path.isdir(ordner0) else ""
        try:
            qual0 = max(1, min(100, int(cfg.get("qual", 92))))
        except (TypeError, ValueError):
            qual0 = 92
        self.qual = tk.IntVar(value=qual0)
        self.open_var = tk.BooleanVar(value=bool(cfg.get("open", True)))
        self.top_var = tk.BooleanVar(value=bool(cfg.get("top", True)))
        self.sound_var = tk.BooleanVar(value=bool(cfg.get("sound", True)))
        self.jai = JAI()
        self.jai_var = tk.BooleanVar(value=cfg.get("jai") is True)     # Standard: aus
        self.jai.an = self.jai_var.get()
        self._jai_ziel = ""
        self.exif_var = tk.BooleanVar(value=bool(cfg.get("exif", False)))
        self.effekte = self._effekte_aus_cfg(cfg)
        self._effekt_knoepfe = None    # (Knöpfe je Stufe, Stil), solange das Hauptmenü offen ist
        self.watch_var = tk.BooleanVar(value=False)
        self.preset_web_var = tk.BooleanVar(value=bool(cfg.get("preset_web", False)))
        self.tray_min_var = tk.BooleanVar(value=bool(cfg.get("tray_min", False)) and TRAY_AVAILABLE)
        r.attributes("-topmost", self.top_var.get())
        wo = cfg.get("watch_ordner")
        self.watch_ordner = wo if isinstance(wo, str) else ""
        self._watch_gesehen, self._watch_groessen = set(), {}
        self.abbruch = threading.Event()
        self.wartend = []
        self.wahl = None               # Modus FORMAT WÄHLEN: {"dateien", "formate", "i", "nr", "gesamt"}
        self.wahl_queue = []           # weitere Gruppen [(dateien, formate)], die danach zur Auswahl kommen
        self._wahl_nr = 0
        self._wahl_start_t = 0.0
        self.erzeugt = set()
        self._tray = None

        self._zustand = "bereit"         # über die Property zustand: jeder Wechsel weckt den Kreis
        self.hover = False
        self.mini_mode = mini0
        self._tick_id = None           # after-ID der Bildschleife (es läuft immer genau eine)
        self._standbild = False        # True: im Leerlauf wird nicht gezeichnet, nur der Prüf-Takt läuft
        self._letzte_eingabe = self._start_t = time.time()
        self.render_zaehler = 0        # Anzahl gezeichneter Bilder (für Tests der Drosselung)
        self.welle_radius = 40.0
        self.fortschritt = 0.0
        self.prog_anzeige = 0.0
        self.info = ("", "", "")
        self.ergebnis = (0, 0, "")
        self.warnungen = []            # Warnungen der letzten Umwandlung (Datei erstellt, aber bitte prüfen)
        self.erkannt_text = ""         # Hinweis einer Erweiterung, wenn sie eine Datei erkannt hat
        self.erparnis_text = ""
        self.reset_id = None
        self.historie = []
        self.w_ticks = self.w_radar = self.w_hex = 0.0
        self.w_ringe = [0.0, 0.0, 0.0]          # Zierringe A, B, C (hud_grafik.RINGE)
        self.t0 = self.t_letzt = self.t_fertig = time.time()
        self.funken = []
        self._press = None
        self._last_watch_files = set()

        self.c.bind("<ButtonPress-1>", self._press_ev)
        self.c.bind("<B1-Motion>", self._motion_ev)
        self.c.bind("<ButtonRelease-1>", self._release_ev)
        self.c.bind("<Double-Button-1>", self._double_click_ev)
        self.c.bind("<Motion>", self._hover_cursor)
        self.c.bind("<Enter>", self._wecken)
        self.c.bind("<Button-3>", lambda e: [self._wecken(), self._modernes_haupt_menu(e.x_root, e.y_root)])
        self.hinweis_phase = 0
        r.bind("<Key>", self._taste)
        r.after(250, self._fokus)
        
        if self.dnd:
            self.c.drop_target_register(DND_FILES)
            self.c.dnd_bind("<<Drop>>", self._on_drop)
            self.c.dnd_bind("<<DropEnter>>", self._on_enter)
            self.c.dnd_bind("<<DropLeave>>", self._on_leave)
            try:
                r.drop_target_register(DND_FILES)
                r.dnd_bind("<<Drop>>", self._on_drop)
            except Exception:
                pass

        self.zeichne_texte()
        r.after(60, self._poll)
        self._tick_id = r.after(20, self._tick)
        r.after(1000, self._watch_loop)
        if self.jai.an:
            r.after(600, lambda: self._ansage("willkommen"))

    def _tk_fehler(self, exc, val, tb):
        import traceback
        logge("".join(traceback.format_exception(exc, val, tb)))
        try:
            if self.reset_id:
                self.root.after_cancel(self.reset_id)
                self.reset_id = None
            self.ergebnis = (0, 1, f"{exc.__name__}: {val}")
            self._zeige_fertig()
        except Exception:
            pass

    def _wechsel_theme(self, neues_theme):
        self.theme_var.set(neues_theme)
        self.g.rebuild_theme(neues_theme)
        self.zeichne_texte()
        self._speichern()

    def _toggle_mini(self):
        alt = self.size
        self.mini_mode = not self.mini_mode
        self.size = P(S_MINI if self.mini_mode else S_NORMAL)
        x, y = self._pos
        # linke obere Ecke bleibt; der Kreis wird aber ganz auf den Monitor geschoben, auf dem er war
        bereich = monitor_bereich(x + alt // 2, y + alt // 2, self.root)
        if bereich:
            l, t, r, b = bereich
            x = min(max(x, l), max(l, r - self.size))
            y = min(max(y, t), max(t, b - self.size))
        self._pos = (x, y)
        self.root.geometry(f"{self.size}x{self.size}+{x}+{y}")
        self.c.configure(width=self.size, height=self.size)
        self.zeichne_texte()
        self._speichern()

    def _double_click_ev(self, e):
        self._wecken()
        if self.zustand == "wahl" or time.time() - self._wahl_start_t < 0.6:
            return                         # Klicks im Auswahlmodus nicht als Doppelklick deuten
        if self.tray_min_var.get() and TRAY_AVAILABLE:
            self._tray_verstecken()
            return
        if self.mini_mode:
            self._toggle_mini()
        else:
            tag = self._tag_bei(e.x, e.y)
            if tag not in ("btn_menu", "btn_close", "prev", "next", "fmt", "datei_auswahl"):
                self._toggle_mini()

    def _anzeige_format(self):
        """Im Modus FORMAT WÄHLEN das gerade angebotene Format, sonst das eingestellte."""
        if self.zustand == "wahl" and self.wahl:
            return self.wahl["formate"][self.wahl["i"]]
        return self.format

    def _schluessel(self):
        if self.zustand == "wahl":
            return "wahl"
        if self.zustand == "bereit":
            return "hover" if self.hover else kat_von_format(self.format)
        if self.zustand == "laeuft":
            return "run"
        return "ok" if self.ergebnis[0] else "err"

    @property
    def zustand(self):
        return self._zustand

    @zustand.setter
    def zustand(self, neu):
        alt = getattr(self, "_zustand", None)
        self._zustand = neu
        if alt is not None and alt != neu:
            self._wecken()

    def _wecken(self, _e=None):
        """Jede Eingabe: Leerlauf-Uhr zurücksetzen; aus dem Standbild sofort (genau einmal) weiterzeichnen."""
        self._letzte_eingabe = time.time()
        if self._standbild:
            if self._tick_id is not None:
                try:
                    self.root.after_cancel(self._tick_id)
                except Exception:
                    pass
            self._standbild = False
            self._tick_id = self.root.after(0, self._tick)

    def _standbild_faellig(self, jetzt):
        return (self.zustand == "bereit" and not self.hover and not self.wahl_queue
                and jetzt - self._letzte_eingabe >= STANDBILD_NACH)

    def _hinweis_wechsel(self, jetzt):
        phase = int(jetzt // 4) % 2
        if phase != self.hinweis_phase:
            self.hinweis_phase = phase
            if self.zustand == "bereit" and not self.hover and not self.mini_mode:
                self.zeichne_texte()

    def _tick(self):
        try:
            if self.root.state() in ("withdrawn", "iconic"):   # minimiert / im Tray: nichts rendern
                self.t_letzt = time.time()
                self._tick_id = self.root.after(500, self._tick)
                return
        except tk.TclError:
            pass
        jetzt = time.time()
        if self._standbild_faellig(jetzt):          # Standbild: letztes Bild bleibt stehen
            self._standbild = True
            self.t_letzt = jetzt
            self._hinweis_wechsel(jetzt)
            self._tick_id = self.root.after(PRUEF_TAKT, self._tick)
            return
        self._standbild = False
        self.render_zaehler += 1
        dt = min(0.1, jetzt - self.t_letzt)
        self.t_letzt = jetzt
        self._hinweis_wechsel(jetzt)
        if self.zustand == "bereit" and self.hover:
            v_t, v_r = 45, 230
            self.welle_radius += 120 * dt
            if self.welle_radius > 150:
                self.welle_radius = 40.0
        else:
            v_t, v_r = {"bereit": (10, 38), "laeuft": (60, 150), "fertig": (25, 60), "wahl": (22, 95)}[self.zustand]
            self.welle_radius = 40.0

        self.w_ticks = (self.w_ticks + v_t * dt) % 360
        self.w_radar = (self.w_radar + v_r * dt) % 360
        self.w_hex = (self.w_hex + 22 * dt) % 360
        self.prog_anzeige += (self.fortschritt - self.prog_anzeige) * min(1.0, dt * 8)
        stufe = self.effekte
        if self.mini_mode and stufe == "voll":
            stufe = "sparsam"                       # Mini-Modus: höchstens "sparsam"
        faktor = 2.5 if self.zustand == "laeuft" else (1.5 if (self.hover or self.zustand == "wahl") else 1.0)
        for i, name in enumerate("ABC"):
            self.w_ringe[i] = (self.w_ringe[i] + hg.RINGE[name][5] * faktor * dt) % 360

        key = self._schluessel()
        anzeige = self._anzeige_format()
        kat = kat_von_format(anzeige)
        try:
            formate = self.wahl["formate"] if (self.zustand == "wahl" and self.wahl) else sichtbare_formate()
            bogen_i, bogen_n = formate.index(anzeige), len(formate)
        except Exception:
            bogen_i, bogen_n = None, 0
        d = BildDaten(zustand=self.zustand, hover=self.hover, mini=self.mini_mode, key=key, kat=kat,
                      ekey=key if key in EMB_KEYS else kat, bogen_i=bogen_i, bogen_n=bogen_n,
                      w_ticks=self.w_ticks, w_radar=self.w_radar, w_hex=self.w_hex,
                      welle_radius=self.welle_radius, prog=self.prog_anzeige,
                      t_fertig=jetzt - self.t_fertig, t_lauf=jetzt - self.t0,
                      stufe=stufe, w_ringe=tuple(self.w_ringe))
        out_img = hg.zeichne(self.g, d)

        if self.photo is None or self.photo.width() != out_img.width:
            self.photo = ImageTk.PhotoImage(out_img)
            self.c.itemconfigure(self.bild_id, image=self.photo)
        else:
            self.photo.paste(out_img)       # vorhandenes Bild wiederverwenden statt jedes Mal neu anlegen
        self.c.tag_lower(self.bild_id)
        if stufe != "aus" and self.g.ringe is None:
            self.g.ringe_starten()                  # nach dem ersten Bild im Hintergrund (höchstens einmal je Grafik)
        if not self.mini_mode:
            self._funken_schritt()
        self._tick_id = self.root.after(self._tick_pause(), self._tick)

    def _tick_pause(self):
        """Millisekunden bis zum nächsten Bild: während der Arbeit flüssig (30 fps), im Leerlauf 12 fps,
        im Mini-Modus sparsam."""
        if self.mini_mode:
            return 33 if (self.zustand == "laeuft" or self.hover) else 125
        return 33 if (self.hover or self.zustand != "bereit") else 83

    def _funken_start(self):
        t_data = theme_daten(self.theme_var.get())
        f_list = t_data["FUNKEN"]
        M = 170
        for _ in range(80):
            ang = random.uniform(0, math.tau)
            v = random.uniform(5, 16)
            leben = random.randint(34, 70)
            item = self.c.create_line(0, 0, 0, 0, fill=random.choice(f_list), width=max(2, P(2)), tags="funke")
            self.funken.append({"id": item, "x": M, "y": 112.0, "vx": math.cos(ang) * v,
                                "vy": math.sin(ang) * v, "leben": leben, "max": leben})

    def _funken_schritt(self):
        if not self.funken or self.mini_mode:
            return
        rest = []
        for p in self.funken:
            p["vx"] *= 0.94
            p["vy"] = p["vy"] * 0.94 + 0.10
            p["x"] += p["vx"]
            p["y"] += p["vy"]
            p["leben"] -= 1
            if p["leben"] <= 0:
                self.c.delete(p["id"])
                continue
            k = 2.2 * p["leben"] / p["max"]
            self.c.coords(p["id"], p["x"] * hg.SKALA, p["y"] * hg.SKALA,
                          (p["x"] - p["vx"] * k) * hg.SKALA, (p["y"] - p["vy"] * k) * hg.SKALA)
            rest.append(p)
        self.funken = rest
        self.c.tag_raise("funke")

    def zeichne_texte(self):
        c = self.c
        c.delete("txt")
        if self.mini_mode:
            c.create_text(S_MINI/2, S_MINI/2, text=self.format.upper(), fill=WEISS, font=(FH, 11, "bold"), tags="txt")
            if hg.SKALA != 1.0:
                c.scale("txt", 0, 0, hg.SKALA, hg.SKALA)
            return

        M = 170
        key = self._schluessel()
        akzent = hexfarbe(self.g.farben[key][0])
        t_data = theme_daten(self.theme_var.get())
        c_cyan, c_mag = t_data["CYAN"], t_data["MAGENTA"]

        def t(y, text, farbe, schrift, glow=False):
            if glow:
                c.create_text(M - 1, y, text=text, fill=c_cyan, font=schrift, tags="txt")
                c.create_text(M + 1, y, text=text, fill=c_mag, font=schrift, tags="txt")
            c.create_text(M, y, text=text, fill=farbe, font=schrift, tags="txt")

        if self.zustand == "bereit":
            t(62, "●  " + gesperrt("empfang" if self.hover else "bereit"), akzent, (FM, 9, "bold"))
            if self.hover:
                t(180, gesperrt("loslassen"), WEISS, (FH, 15, "bold"), glow=True)
                t(201, "AKTION ▸ " + self.format.upper(), akzent, (FM, 9, "bold"))
            else:
                t(180, gesperrt("dateien ablegen" if self.dnd else "hier klicken"), WEISS,
                  (FH, 13, "bold"), glow=True)
                if self.hinweis_phase == 1:
                    t(200, "←  →  FORMAT/MODUS  ·  ENTER ÖFFNEN", akzent, (FM, 9, "bold"))
                else:
                    t(200, "ZIEHEN ODER KLICKEN" if self.dnd else "DATEIEN AUSWÄHLEN", akzent, (FM, 9, "bold"))
            t(PILL_Y, self.format.upper(), WEISS, (FH, 15, "bold"), glow=True)
            
            if self.ordner:
                ziel_str = os.path.basename(self.ordner.rstrip("/\\")) or self.ordner
                t_ziel = f"ZIEL ▸ {kuerze(ziel_str.upper(), 18)}"
            elif self.format.startswith("resize "):
                t_ziel = "ZIEL ▸ ORDNER: Verkleinert"
            else:
                t_ziel = "ZIEL ▸ Quellordner"
                
            t(278, t_ziel, GRAU, (FM, 9))
            t(295, "© BY JOSCHY", akzent, (FM, 9, "bold"))
            
            if self.preset_web_var.get():
                c.create_text(M, 42, text="[ 🔥 WEB-PRESET AKTIV ]", fill="#00ff9d", font=(FM, 8, "bold"), tags="txt")
            elif self.watch_var.get():
                c.create_text(M, 42, text="[ ÜBERWACHUNG AKTIV ]", fill="#00ff9d", font=(FM, 8, "bold"), tags="txt")
            elif self.ordner:
                c.create_text(M, 42, text="[ ZIEL-MODUS AKTIV ]", fill="#00ff9d", font=(FM, 8, "bold"), tags="txt")

        elif self.zustand == "wahl" and self.wahl:
            w = self.wahl
            fm = w["formate"][w["i"]]
            t(62, "●  " + gesperrt("auswahl"), akzent, (FM, 9, "bold"))
            t(180, gesperrt("format wählen"), WEISS, (FH, 13, "bold"), glow=True)
            t(200, "← → WÄHLEN  ·  ENTER START", akzent, (FM, 9, "bold"))
            t(PILL_Y, fm.upper(), WEISS, (FH, 15, "bold"), glow=True)
            zeile = f"{w['i'] + 1} / {len(w['formate'])}"
            if w["gesamt"] > 1:
                zeile += f"  ·  GRUPPE {w['nr']}/{w['gesamt']}"
            t(268, zeile, GRAU, (FM, 8, "bold"))
            t(283, "✕ ABBRECHEN (ESC)", GRAU, (FM, 9, "bold"))
            t(298, "© BY JOSCHY", akzent, (FM, 9, "bold"))
            n = len(w["dateien"])
            endungen = " ".join(sorted({"." + (ext_von(d) or "?") for d in w["dateien"]}))
            kopf = f"{n} DATEI{'EN' if n != 1 else ''} {endungen.upper()} ≠ {self.format.upper()}"
            c.create_text(M, 42, text=f"[ {kuerze(kopf, 30)} ]", fill="#ffb000", font=(FM, 8, "bold"), tags="txt")

        elif self.zustand == "laeuft":
            t(62, "●  " + gesperrt("aktiv"), akzent, (FM, 9, "bold"))
            t(128, self.info[0].replace(" ", ""), WEISS, (FH, 38, "bold"), glow=True)
            t(180, gesperrt("verarbeite"), WEISS, (FH, 12, "bold"))
            t(201, kuerze(self.info[1], 28), GRAU, (FM, 9))
            t(228, self.info[2], akzent, (FM, 11, "bold"))
            if self.wartend:
                t(250, f"+{len(self.wartend)} IN WARTESCHLANGE", akzent, (FM, 8))
            t(274, "✕ ABBRECHEN (ESC)", GRAU, (FM, 9, "bold"))
            t(295, "© BY JOSCHY", akzent, (FM, 9, "bold"))
            if self.erkannt_text:
                c.create_text(M, 42, text=f"[ {self.erkannt_text} ]", fill="#ffb000", font=(FM, 8, "bold"), tags="txt")

        else:
            ok, fehler, fehltext = self.ergebnis
            t(62, "●  " + gesperrt("fertig" if ok else "fehler"), akzent, (FM, 9, "bold"))
            t(174, gesperrt(f"{ok} fertig" if ok else "abbruch"), akzent, (FH, 17, "bold"), glow=True)
            
            if self.erparnis_text and ok:
                t(200, self.erparnis_text, "#00ff9d", (FM, 9, "bold"))
                
            warnungen = self.warnungen
            if fehler:
                t(214, f"{fehler} FEHLGESCHLAGEN", "#ffb000", (FM, 9, "bold"))
                for i, zeile in enumerate(umbrechen(fehltext, 30, 2)):
                    t(230 + 14 * i, zeile, "#ff5c7a", (FM, 9))
            elif warnungen:
                t(214, "⚠ " + (f"{len(warnungen)} WARNUNGEN" if len(warnungen) > 1 else "WARNUNG") + " · BITTE PRÜFEN",
                  "#ffb000", (FM, 9, "bold"))
                for i, zeile in enumerate(umbrechen(warnungen[0], 30, 2)):
                    t(230 + 14 * i, zeile, "#ff9a3c", (FM, 9))
            t(274, "KLICKEN ▸ WEITER" + ("  ·  H ▸ DETAILS" if (fehler or warnungen) else ""), GRAU, (FM, 9))
            t(295, "© BY JOSCHY", akzent, (FM, 9, "bold"))
            if self.erkannt_text:
                c.create_text(M, 42, text=f"[ {self.erkannt_text} ]", fill="#ffb000", font=(FM, 8, "bold"), tags="txt")
            elif warnungen and fehler:
                c.create_text(M, 42, text=f"[ ⚠ {len(warnungen)} WARNUNG(EN) ]", fill="#ffb000",
                              font=(FM, 8, "bold"), tags="txt")

        if hg.SKALA != 1.0:
            c.scale("txt", 0, 0, hg.SKALA, hg.SKALA)
        c.tag_raise("funke")

    # ---- Bausteine für alle Menüs im Stil der Reaktor-Konsole
    def _hud_stil(self):
        td = theme_daten(self.theme_var.get())
        a, bg = hexrgb(td["CYAN"]), hexrgb(td["bg_inner"])
        return {"A": td["CYAN"], "ROT": td["ROT"], "BG": td["bg_dark"], "ROW": hexfarbe(mix(bg, a, 0.07)),
                "ROW_H": hexfarbe(mix(bg, a, 0.22)), "LINIE": td["border"], "TEXT": "#dff7f4",
                "DIM": hexfarbe(mix((128, 150, 156), a, 0.25))}

    def _hud_fenster(self, titel):
        """Rahmenloses Konsolenfenster mit Kopfzeile, Leuchtecken und rotem Akzent."""
        st = self._hud_stil()
        top = tk.Toplevel(self.root)
        top.overrideredirect(True)
        top.attributes("-topmost", True)
        top.configure(bg=st["LINIE"])
        aussen = tk.Frame(top, bg=st["BG"])
        aussen.pack(fill="both", expand=True, padx=1, pady=1)
        kopf = tk.Frame(aussen, bg=st["BG"])
        kopf.pack(fill="x", padx=12, pady=(10, 2))
        tk.Label(kopf, text="◢ " + titel.upper(), fg=st["A"], bg=st["BG"], font=(FM, 9, "bold")).pack(side="left")
        tk.Frame(kopf, bg=st["ROT"], width=18, height=3).pack(side="right")
        tk.Frame(kopf, bg=st["LINIE"], height=1).pack(side="left", fill="x", expand=True, padx=8)
        inhalt = tk.Frame(aussen, bg=st["BG"])
        inhalt.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        for relx, rely, anker in ((0, 0, "nw"), (1, 0, "ne"), (0, 1, "sw"), (1, 1, "se")):
            tk.Frame(top, bg=st["A"], width=18, height=2).place(relx=relx, rely=rely, anchor=anker)
            tk.Frame(top, bg=st["A"], width=2, height=18).place(relx=relx, rely=rely, anchor=anker)
        return top, inhalt, st

    @staticmethod
    def _hud_abschnitt(parent, text, st, kennfarbe=None):
        z = tk.Frame(parent, bg=st["BG"])
        z.pack(fill="x", pady=(9, 3))
        if kennfarbe:
            tk.Frame(z, bg=kennfarbe, width=6, height=6).pack(side="left", padx=(0, 6))
        tk.Label(z, text=text.upper(), fg=st["DIM"], bg=st["BG"], font=(FM, 8, "bold")).pack(side="left")
        tk.Frame(z, bg=st["LINIE"], height=1).pack(side="left", fill="x", expand=True, padx=(6, 0))

    @staticmethod
    def _hud_knopf(parent, text, cmd, st, fg=None, schrift=None, pady=3):
        """Menüzeile mit Leuchtbalken links, der beim Überfahren aufleuchtet."""
        zeile = tk.Frame(parent, bg=st["ROW"])
        balken = tk.Frame(zeile, bg=st["ROW"], width=3)
        balken.pack(side="left", fill="y")
        b = tk.Button(zeile, text=text, fg=fg or st["TEXT"], bg=st["ROW"], activebackground=st["ROW_H"],
                      activeforeground=st["A"], bd=0, relief="flat", highlightthickness=0, font=schrift or (FH, 9),
                      anchor="w", padx=9, pady=pady, cursor="hand2", command=cmd)
        b.pack(side="left", fill="x", expand=True)
        b._fg, b._balken, b._balken_farbe = fg or st["TEXT"], balken, st["ROW"]

        def an(_e=None):
            b.configure(bg=st["ROW_H"], fg=st["A"])
            balken.configure(bg=st["A"])

        def aus(_e=None):
            b.configure(bg=st["ROW"], fg=b._fg)
            balken.configure(bg=b._balken_farbe)
        b.bind("<Enter>", an)
        b.bind("<Leave>", aus)
        zeile.pack(fill="x", pady=1)
        return b

    @staticmethod
    def _hud_markiere(b, aktiv, st, fg_normal=None):
        b._fg = st["A"] if aktiv else (fg_normal or st["TEXT"])
        b._balken_farbe = st["A"] if aktiv else st["ROW"]
        b.configure(fg=b._fg)
        b._balken.configure(bg=b._balken_farbe)

    @staticmethod
    def _hud_scrollleiste(parent, ziel, st):
        """Schmale Leucht-Scrollleiste (statt der grauen Windows-Leiste); mit Mausrad und Ziehen."""
        leiste = tk.Canvas(parent, width=P(4), bg=st["ROW"], highlightthickness=0, bd=0)
        daumen = leiste.create_rectangle(0, 0, P(4), 0, fill=st["A"], outline="")

        def sync(erst, letzt):
            h = max(1, leiste.winfo_height())
            erst, letzt = float(erst), float(letzt)
            leiste.coords(daumen, 0, erst * h, P(4), max(erst * h + P(12), letzt * h))
            leiste.itemconfigure(daumen, state="hidden" if (erst <= 0 and letzt >= 1) else "normal")

        def ziehen(e):
            h = max(1, leiste.winfo_height())
            erst, letzt = ziel.yview()
            ziel.yview_moveto(max(0.0, e.y / h - (letzt - erst) / 2))
        leiste.bind("<B1-Motion>", ziehen)
        leiste.bind("<Button-1>", ziehen)
        leiste.bind("<Configure>", lambda e: sync(*ziel.yview()))
        ziel.configure(yscrollcommand=sync)
        return leiste

    def _modernes_fmt_menu(self, x, y):
        top, inhalt, st = self._hud_fenster("Zielformat")
        breite, hoehe = P(196), P(380)
        sw, sh = top.winfo_screenwidth(), top.winfo_screenheight()
        x = max(0, min(x, sw - breite - 10))
        y = max(0, min(y, sh - hoehe - 10))
        top.geometry(f"{breite}x{hoehe}+{x}+{y}")

        box = tk.Frame(inhalt, bg=st["BG"])
        box.pack(fill="both", expand=True)
        canvas_in = tk.Canvas(box, bg=st["BG"], highlightthickness=0, bd=0)
        liste = tk.Frame(canvas_in, bg=st["BG"])
        fenster = canvas_in.create_window((0, 0), window=liste, anchor="nw")
        liste.bind("<Configure>", lambda e: canvas_in.configure(scrollregion=canvas_in.bbox("all")))
        canvas_in.bind("<Configure>", lambda e: canvas_in.itemconfigure(fenster, width=e.width))
        leiste = self._hud_scrollleiste(box, canvas_in, st)
        leiste.pack(side="right", fill="y", padx=(4, 0))
        canvas_in.pack(side="left", fill="both", expand=True)

        def rad(event):
            canvas_in.yview_scroll(int(-1 * (event.delta / 120)), "units")
        canvas_in.bind_all("<MouseWheel>", rad)

        def schliessen():
            try:
                canvas_in.unbind_all("<MouseWheel>")
                top.destroy()
            except tk.TclError:
                pass
            self._wecken()
            self.root.after(80, self._fokus)

        aktiver = None
        for kat in sichtbare_kategorien():
            formate = KATEGORIEN[kat]
            self._hud_abschnitt(liste, kat, st, kennfarbe=FARBEN[kat][0])
            for fm in formate:
                b = self._hud_knopf(liste, fm.upper(), lambda f=fm: [self._setze_format(f), schliessen()], st,
                                    schrift=(FH, 8), pady=1)
                if fm == self.format:
                    b.configure(text="▶ " + fm.upper())
                    self._hud_markiere(b, True, st)
                    aktiver = b

        def zum_aktiven():                      # aktuelles Format sichtbar machen
            top.update_idletasks()
            if aktiver is not None and liste.winfo_height() > 0:
                canvas_in.yview_moveto(max(0.0, (aktiver.winfo_rooty() - liste.winfo_rooty() - 60) / liste.winfo_height()))
        top.after(30, zum_aktiven)
        top.bind("<FocusOut>", lambda e: e.widget == top and schliessen())
        top.bind("<Escape>", lambda e: schliessen())
        top.focus_set()

    def _historie_dialog(self):
        top, inhalt, st = self._hud_fenster("Protokoll")
        mitte = tk.Frame(inhalt, bg=st["BG"])
        mitte.pack(fill="both", expand=True, pady=(6, 0))
        lb = tk.Listbox(mitte, bg=st["ROW"], fg=st["TEXT"], bd=0, highlightthickness=0, activestyle="none",
                        font=(FM, 9), width=62, height=15, selectbackground=st["ROW_H"], selectforeground=st["A"])
        leiste = self._hud_scrollleiste(mitte, lb, st)
        leiste.pack(side="right", fill="y", padx=(4, 0))
        lb.pack(side="left", fill="both", expand=True)
        eintraege = list(reversed(self.historie[-100:])) or ["Noch keine Einträge."]
        for i, item in enumerate(eintraege):
            lb.insert("end", "  " + item)
            if "[FEHLER]" in item:
                lb.itemconfigure(i, fg=st["ROT"])
            elif "[WARNUNG]" in item:
                lb.itemconfigure(i, fg="#ffb000")
            elif "[OK]" in item:
                lb.itemconfigure(i, fg=st["A"])

        def log_oeffnen():
            pfad = os.path.join(daten_ordner(), "kreis_fehler.log")
            try:
                if hasattr(os, "startfile"):
                    os.startfile(pfad if os.path.exists(pfad) else daten_ordner())
            except Exception:
                pass

        def zu(_e=None):
            top.destroy()
            self.root.after(80, self._fokus)
        tk.Frame(inhalt, bg=st["BG"], height=8).pack()
        self._hud_knopf(inhalt, "▸  Fehler-Log öffnen", log_oeffnen, st)
        self._hud_knopf(inhalt, "✕  Schließen", zu, st, fg=st["ROT"])
        top.update_idletasks()
        top.geometry(f"+{self.root.winfo_x() + P(20)}+{self.root.winfo_y() + P(40)}")
        top.bind("<Escape>", zu)
        top.focus_set()

    def _lizenzen_dialog(self):
        """Zeigt die eigene Lizenz (AGPL-3.0, Link zum Quellcode) und die Lizenzen der mitgelieferten
        Fremdsoftware (Ordner 'lizenzen', in die exe eingebettet)."""
        top, inhalt, st = self._hud_fenster("Lizenzen")
        dateien = lizenz_dateien()
        reiter = tk.Frame(inhalt, bg=st["BG"])
        reiter.pack(fill="x", pady=(6, 4))
        mitte = tk.Frame(inhalt, bg=st["BG"])
        mitte.pack(fill="both", expand=True)
        feld = tk.Text(mitte, bg=st["ROW"], fg=st["TEXT"], bd=0, highlightthickness=0, wrap="word",
                       font=(FM, 9), width=84, height=26, padx=8, pady=6, insertwidth=0,
                       selectbackground=st["ROW_H"], selectforeground=st["A"])
        leiste = self._hud_scrollleiste(mitte, feld, st)
        leiste.pack(side="right", fill="y", padx=(4, 0))
        feld.pack(side="left", fill="both", expand=True)
        knoepfe = []

        def zeige(i):
            if dateien[i][1] is None:
                text = eigener_lizenz_text()
                if len(dateien) == 1:
                    text += ("\n\nDie Lizenztexte der Fremdsoftware wurden nicht gefunden (Ordner 'lizenzen' fehlt).\n"
                             "FFmpeg: GPL-3.0 · 7-Zip: GNU LGPL + unRAR-Einschränkung")
            else:
                try:
                    with open(dateien[i][1], "r", encoding="utf-8-sig", errors="replace") as f:
                        text = f.read()
                except OSError as ex:
                    text = f"Datei konnte nicht gelesen werden: {ex}"
            feld.configure(state="normal")
            feld.delete("1.0", "end")
            feld.insert("1.0", text)
            feld.configure(state="disabled")
            feld.yview_moveto(0)
            for k, kn in enumerate(knoepfe):
                aktiv = k == i
                kn.configure(bg=st["A"] if aktiv else st["ROW"], fg=st["BG"] if aktiv else st["TEXT"])

        for i, (name, _pfad) in enumerate(dateien):
            kn = tk.Button(reiter, text=name, bd=0, relief="flat", highlightthickness=0, font=(FM, 8, "bold"),
                           padx=6, pady=3, cursor="hand2", bg=st["ROW"], fg=st["TEXT"],
                           activebackground=st["ROW_H"], activeforeground=st["A"], command=lambda k=i: zeige(k))
            kn.pack(side="left", padx=(0, 3))
            knoepfe.append(kn)

        def ordner_oeffnen():
            """Kopiert die Texte in den Datenordner (der exe-Entpackordner verschwindet beim Beenden)."""
            ziel = os.path.join(daten_ordner(), "lizenzen")
            try:
                os.makedirs(ziel, exist_ok=True)
                for _name, pfad in dateien:
                    if pfad is None:
                        with open(os.path.join(ziel, "Kreis-Konverter_LIZENZ.txt"), "w", encoding="utf-8") as f:
                            f.write(eigener_lizenz_text())
                    else:
                        shutil.copyfile(pfad, os.path.join(ziel, os.path.basename(pfad)))
                if hasattr(os, "startfile"):
                    os.startfile(ziel)
            except Exception as ex:
                logge(f"Lizenzordner öffnen: {ex}")

        def zu(_e=None):
            try:
                top.destroy()
            except tk.TclError:
                pass
            self.root.after(80, self._fokus)

        def quellcode_oeffnen():
            try:
                import webbrowser
                webbrowser.open(QUELLCODE_URL)
            except Exception as ex:
                logge(f"Quellcode-Link öffnen: {ex}")

        tk.Frame(inhalt, bg=st["BG"], height=8).pack()
        self._hud_knopf(inhalt, "▸  Quellcode öffnen (GitHub)", quellcode_oeffnen, st)
        self._hud_knopf(inhalt, "▸  Lizenztexte als Dateien öffnen", ordner_oeffnen, st)
        self._hud_knopf(inhalt, "✕  Schließen", zu, st, fg=st["ROT"])
        zeige(0)
        top.update_idletasks()
        w, h = top.winfo_reqwidth(), top.winfo_reqheight()
        sw, sh = top.winfo_screenwidth(), top.winfo_screenheight()
        x = max(0, min(self.root.winfo_x() + (self.size - w) // 2, sw - w - 10))
        y = max(0, min(self.root.winfo_y() + (self.size - h) // 2, sh - h - 10))
        top.geometry(f"+{x}+{y}")
        top.bind("<Escape>", zu)
        top.focus_set()
        self._lizenz_fenster = (top, feld, dateien)      # für Tests

    def _modernes_haupt_menu(self, x, y):
        top, inhalt, st = self._hud_fenster("System")

        def schliessen():
            try:
                top.destroy()
            except tk.TclError:
                pass
            self._effekt_knoepfe = None
            self._speichern()
            self._wecken()
            self.root.after(80, self._fokus)

        def aktion(text, cmd, fg=None):
            return self._hud_knopf(inhalt, text, lambda: [schliessen(), cmd()], st, fg)

        def schalter(text, var, danach=None):
            """Ein-/Aus-Schalter, der das Menü offen lässt."""
            b = self._hud_knopf(inhalt, "", None, st)

            def zeige():
                an = var.get()
                b.configure(text=("■  " if an else "□  ") + text)
                self._hud_markiere(b, an, st)

            def umschalten():
                var.set(not var.get())
                zeige()
                if danach:
                    danach()
            b.configure(command=umschalten)
            zeige()

        self._hud_abschnitt(inhalt, "Ziel", st)
        aktion("▸  Zielordner wählen …", self.ordner_waehlen)
        aktion("↺  Zielordner zurücksetzen", self._ordner_reset)

        self._hud_abschnitt(inhalt, "Design  ·  " + self.theme_var.get(), st)
        zeile = tk.Frame(inhalt, bg=st["BG"])
        zeile.pack(fill="x", pady=2)
        for name, td in THEMES.items():
            aktiv = name == self.theme_var.get()
            rahmen = tk.Frame(zeile, bg=st["A"] if aktiv else st["LINIE"])
            rahmen.pack(side="left", padx=(0, 6))
            tk.Button(rahmen, text="", width=2, height=1, bg=td["CYAN"], activebackground=td["CYAN"], bd=0,
                      relief="flat", highlightthickness=0, cursor="hand2",
                      command=lambda n=name: [self._wechsel_theme(n), schliessen()]).pack(padx=2 if aktiv else 1,
                                                                                         pady=2 if aktiv else 1)

        self._hud_abschnitt(inhalt, "Qualität / Bitrate", st)
        zeile = tk.Frame(inhalt, bg=st["BG"])
        zeile.pack(fill="x", pady=2)
        q_knoepfe = {}

        def setze_q(w):
            self.qual.set(w)
            for v, kn in q_knoepfe.items():
                kn.configure(fg=st["BG"] if v == w else st["TEXT"], bg=st["A"] if v == w else st["ROW"])

        for w in (100, 92, 85, 75):
            an = self.qual.get() == w
            kn = tk.Button(zeile, text=f"{w}%", bg=st["A"] if an else st["ROW"], fg=st["BG"] if an else st["TEXT"],
                           activebackground=st["ROW_H"], activeforeground=st["A"], bd=0, relief="flat",
                           highlightthickness=0, font=(FM, 8, "bold"), pady=3, cursor="hand2",
                           command=lambda v=w: setze_q(v))
            kn.pack(side="left", fill="x", expand=True, padx=(0, 3))
            q_knoepfe[w] = kn

        self._hud_abschnitt(inhalt, "Optionen", st)
        zeile = tk.Frame(inhalt, bg=st["BG"])
        zeile.pack(fill="x", pady=2)
        tk.Label(zeile, text="Effekte", fg=st["TEXT"], bg=st["BG"], font=(FM, 8, "bold")).pack(side="left",
                                                                                             padx=(0, 6))
        e_knoepfe = {}
        for w in self.EFFEKT_STUFEN:
            an = self.effekte == w
            kn = tk.Button(zeile, text=w, bg=st["A"] if an else st["ROW"], fg=st["BG"] if an else st["TEXT"],
                           activebackground=st["ROW_H"], activeforeground=st["A"], bd=0, relief="flat",
                           highlightthickness=0, font=(FM, 8, "bold"), pady=3, cursor="hand2",
                           command=lambda v=w: self._setze_effekte(v))
            kn.pack(side="left", fill="x", expand=True, padx=(0, 3))
            e_knoepfe[w] = kn
        self._effekt_knoepfe = (e_knoepfe, st)
        schalter("Ordner danach öffnen", self.open_var)
        schalter("Immer im Vordergrund", self.top_var,
                 lambda: self.root.attributes("-topmost", self.top_var.get()))
        schalter("Audio-Feedback", self.sound_var)
        schalter("JAI-Stimme", self.jai_var, self._jai_umschalten)
        aktion("▶  JAI testen", self._jai_testen)
        schalter("EXIF-Daten entfernen", self.exif_var)
        schalter("Web-Preset (1200 px, ohne EXIF)", self.preset_web_var, self.zeichne_texte)
        if self.watch_var.get():
            b = aktion("■  Überwachung: " + kuerze(os.path.basename(self.watch_ordner.rstrip("/\\")) or self.watch_ordner, 18),
                       self._watch_toggle)
            self._hud_markiere(b, True, st)
        else:
            aktion("□  Ordner überwachen …", self._watch_toggle)
        if TRAY_AVAILABLE:
            schalter("Doppelklick ▸ ins Tray", self.tray_min_var)

        self._hud_abschnitt(inhalt, "System", st)
        aktion("≡  Protokoll / Historie", self._historie_dialog)
        aktion("§  Lizenzen", self._lizenzen_dialog)
        aktion("✕  Beenden", self._beenden, fg=st["ROT"])
        tk.Label(inhalt, text="© JOSCHY – Lizenz: AGPL-3.0", fg=st["DIM"], bg=st["BG"], font=(FM, 7, "bold")).pack(anchor="e", pady=(6, 0))

        top.update_idletasks()
        breite = max(P(270), top.winfo_reqwidth())
        hoehe = top.winfo_reqheight()
        sw, sh = top.winfo_screenwidth(), top.winfo_screenheight()
        x = max(0, min(x, sw - breite - 10))
        y = max(0, min(y, sh - hoehe - 10))
        top.geometry(f"{breite}x{hoehe}+{x}+{y}")
        top.bind("<FocusOut>", lambda e: e.widget == top and schliessen())
        top.bind("<Escape>", lambda e: schliessen())
        top.focus_set()

    def _setze_effekte(self, wert):
        """Effekt-Stufe setzen, speichern, im offenen Menü markieren und das Bild sofort neu zeichnen."""
        if wert not in self.EFFEKT_STUFEN:
            return
        self.effekte = wert
        self._speichern()
        self._wecken()                              # auch im Standbild sofort sichtbar
        if self._effekt_knoepfe:
            knoepfe, st = self._effekt_knoepfe
            for v, kn in knoepfe.items():
                try:
                    kn.configure(fg=st["BG"] if v == wert else st["TEXT"], bg=st["A"] if v == wert else st["ROW"])
                except tk.TclError:
                    self._effekt_knoepfe = None      # Menü schon geschlossen
                    break

    # ---- JAI
    def _ansage(self, cid):
        """Höchstens eine Ansage pro Ereignis; True, wenn JAI spricht (dann kein Piepton)."""
        try:
            return self.jai.sage(cid)
        except Exception:
            return False

    def _jai_umschalten(self):
        self.jai.an = self.jai_var.get()
        if self.jai.an:
            self._ansage("jai_aktiviert")
        self._speichern()

    def _jai_testen(self):
        try:
            self.jai.sage("jai_aktiviert", immer=True)
        except Exception:
            pass

    def _jai_ergebnis(self, ok, fehler, fehltext):
        """Clip-ID für das Ende einer Umwandlung."""
        text = (fehltext or "").lower()
        if fehler and "abgebrochen" in text:
            return "abgebrochen"
        if fehler:
            if "passwort" in text:
                return "passwort_erforderlich"
            if "nicht unterstützt" in text or "kein passendes" in text or "keine passenden" in text:
                return "format_nicht_unterstuetzt"
            return "fehler"
        if ok and self._jai_ziel in PACK_AKTIONEN:
            return "archiv_erstellt"
        if ok > 1:
            return "alle_fertig"
        return "fertig" if ok else None

    def _fokus(self):
        try:
            self.root.focus_force()
            self.c.focus_set()
        except Exception:
            pass

    def _wechsle_art(self, schritt):
        arten = sichtbare_kategorien()
        akt = kat_von_format(self.format)
        i = ((arten.index(akt) if akt in arten else 0) + schritt) % len(arten)
        self._setze_format(KATEGORIEN[arten[i]][0])

    def _taste(self, e):
        self._wecken()
        ks = e.keysym
        kl = ks.lower()
        if self.zustand == "wahl":
            if ks == "Left":
                self._wahl_blaettern(-1)
            elif ks == "Right":
                self._wahl_blaettern(1)
            elif ks in ("Return", "KP_Enter", "space"):
                self._wahl_start()
            elif ks == "Escape":
                self._wahl_abbrechen()
            return
        if self.zustand == "laeuft":
            if ks == "Escape":
                self._abbrechen()
            return
        if self.zustand == "fertig":
            if kl == "h":
                self._historie_dialog()
                return
            if ks in ("Return", "KP_Enter", "space", "Escape"):
                self._zurueck()
                return
            if ks in ("Left", "Right", "Up", "Down"):
                self._zurueck()
            else:
                return
        if ks == "Left":
            self._wechsle_format(-1)
        elif ks == "Right":
            self._wechsle_format(1)
        elif ks == "Up":
            self._wechsle_art(-1)
        elif ks == "Down":
            self._wechsle_art(1)
        elif ks in ("Return", "KP_Enter", "space") or kl == "o":
            if self.format == "qrcode":
                self.qr_dialog_starten()
            else:
                self.dateien_waehlen()
        elif kl == "f":
            self._modernes_fmt_menu(self.root.winfo_rootx() + P(170 - 80),
                        self.root.winfo_rooty() + P(PILL_Y + 25))
        elif kl == "m" or ks in ("Menu", "F10"):
            self._modernes_haupt_menu(self.root.winfo_rootx() + P(170 - BTN_DX),
                                      self.root.winfo_rooty() + P(BTN_Y + 18))
        elif kl == "z":
            self.ordner_waehlen()
        elif ks == "Escape":
            self._beenden()

    def _setze_format(self, fm):
        self.format = fm
        self.zeichne_texte()
        self._speichern()

    def _wechsle_format(self, schritt):
        formate = sichtbare_formate()
        i = ((formate.index(self.format) if self.format in formate else 0) + schritt) % len(formate)
        self._setze_format(formate[i])

    def qr_dialog_starten(self):
        top, inhalt, st = self._hud_fenster("QR-Code Generator")
        tk.Label(inhalt, text="Text oder URL für den QR-Code:", fg=st["DIM"], bg=st["BG"],
                 font=(FH, 9)).pack(anchor="w", pady=(8, 4))
        rahmen = tk.Frame(inhalt, bg=st["A"])
        rahmen.pack(fill="x")
        feld = tk.Entry(rahmen, bg=st["ROW"], fg=st["TEXT"], insertbackground=st["A"], relief="flat", bd=0,
                        highlightthickness=0, font=(FM, 10), width=38)
        feld.pack(fill="x", padx=1, pady=1, ipady=5, ipadx=4)

        def zu(_e=None):
            try:
                top.destroy()
            except tk.TclError:
                pass
            self.root.after(80, self._fokus)

        def ok(_e=None):
            text = feld.get().strip()
            zu()
            if text and self.zustand != "laeuft":
                self.zustand = "laeuft"
                self.fortschritt = self.prog_anzeige = 0.0
                self.info = ("1 / 1", "QR-Code erstellen...", "0 %")
                self.zeichne_texte()
                threading.Thread(target=self._arbeit_qr_direkt, args=(text, self.ordner), daemon=True).start()

        tk.Frame(inhalt, bg=st["BG"], height=8).pack()
        self._hud_knopf(inhalt, "▸  QR-Code erzeugen", ok, st, fg=st["A"])
        self._hud_knopf(inhalt, "✕  Abbrechen", zu, st)
        feld.bind("<Return>", ok)
        top.bind("<Escape>", zu)
        top.update_idletasks()
        w, h = top.winfo_reqwidth(), top.winfo_reqheight()
        x = self.root.winfo_x() + (self.size - w) // 2
        y = self.root.winfo_y() + (self.size - h) // 2
        top.geometry(f"+{max(0, x)}+{max(0, y)}")
        top.after(60, feld.focus_force)

    def _arbeit_qr_direkt(self, text, zielordner):
        try:
            self.q.put((self._schritt, (0.5, "Generiere...", "50 %")))
            img = qr_bild(text)
            ordner = zielordner or standard_ausgabe_ordner()
            os.makedirs(ordner, exist_ok=True)
            safe_name = "".join(c if c.isalnum() else "_" for c in text[:15])
            ziel = freier_pfad(os.path.join(ordner, f"QRCode_{safe_name}.png"))
            img.save(ziel, format="PNG")
            self.erzeugt.add(ziel)
            self.q.put((self._fertig, (1, 0, "", ordner, 0, 0)))
        except Exception as ex:
            self.q.put((self._fertig, (0, 1, verstaendlich(ex), zielordner, 0, 0)))

    def ordner_waehlen(self):
        self.root.attributes("-topmost", False)
        p = filedialog.askdirectory(title="Zielordner wählen")
        self.root.attributes("-topmost", self.top_var.get())
        self.root.after(80, self._fokus)
        if p:
            self.ordner = os.path.normpath(p)
            self.zeichne_texte()
            self._speichern()

    def _ordner_reset(self):
        self.ordner = ""
        self.zeichne_texte()
        self._speichern()

    def _tag_bei(self, x, y):
        x, y = x / hg.SKALA, y / hg.SKALA            # echte Pixel -> logische Koordinaten
        M = 170
        if math.hypot(x - (M - BTN_DX), y - BTN_Y) <= BTN_R + 2:
            return "btn_menu"
        if math.hypot(x - (M + BTN_DX), y - BTN_Y) <= BTN_R + 2:
            return "btn_close"
        if self.zustand in ("bereit", "wahl") and PILL_Y - 23 <= y <= PILL_Y + 23:
            if M - 84 <= x < M - 40:
                return "prev"
            if M + 40 < x <= M + 84:
                return "next"
            if M - 40 <= x <= M + 40:
                return "fmt"
        if self.zustand == "bereit" and math.hypot(x - M, y - 118) <= 35:
            return "datei_auswahl"
        if self.zustand == "wahl":
            if abs(x - M) <= 70 and 275 <= y <= 291:
                return "wahl_abbruch"
            if math.hypot(x - M, y - 140) <= 70:
                return "wahl_start"              # Mitte (Kern + "FORMAT WÄHLEN") startet
        if self.zustand == "fertig" and math.hypot(x - M, y - M) <= 150:
            return "datei_auswahl"          # "KLICKEN ▸ WEITER" funktioniert jetzt wirklich
        if self.zustand == "laeuft" and abs(x - M) <= 70 and 262 <= y <= 286:
            return "abbruch"
        return None

    def _hover_cursor(self, e):
        self._wecken()
        if not self.mini_mode:
            self.c.configure(cursor="hand2" if self._tag_bei(e.x, e.y) else "")

    def _press_ev(self, e):
        self._wecken()
        self._fokus()
        self._press = {"x": e.x_root, "y": e.y_root, "wx": self.root.winfo_x(),
                       "wy": self.root.winfo_y(), "moved": False, "tag": self._tag_bei(e.x, e.y)}

    def _motion_ev(self, e):
        p = self._press
        if not p:
            return
        dx, dy = e.x_root - p["x"], e.y_root - p["y"]
        if abs(dx) + abs(dy) > 4:
            p["moved"] = True
        if p["moved"]:
            self._pos = (p["wx"] + dx, p["wy"] + dy)
            self.root.geometry(f"+{self._pos[0]}+{self._pos[1]}")

    def _release_ev(self, e):
        self._wecken()
        p, self._press = self._press, None
        if p and p["moved"]:
            self._speichern()                 # neue Position merken
            return
        if not p:
            return
        tag = p["tag"]
        if tag == "btn_close":
            self._beenden()
        elif tag == "btn_menu":
            self._modernes_haupt_menu(e.x_root, e.y_root)
        elif tag in ("prev", "next") and self.zustand == "wahl":
            self._wahl_blaettern(-1 if tag == "prev" else 1)
        elif tag == "wahl_start":
            self._wahl_start()
        elif tag == "wahl_abbruch":
            self._wahl_abbrechen()
        elif tag == "fmt" and self.zustand == "wahl":
            pass                                 # im Auswahlmodus nur die passenden Formate (Pfeile)
        elif tag == "prev":
            self._wechsle_format(-1)
        elif tag == "next":
            self._wechsle_format(1)
        elif tag == "fmt":
            self._modernes_fmt_menu(self.root.winfo_rootx() + P(170 - 80),
                        self.root.winfo_rooty() + P(PILL_Y + 25))
        elif tag == "abbruch":
            self._abbrechen()
        elif tag == "datei_auswahl":
            if self.zustand == "fertig":
                self._zurueck()
            elif self.zustand == "bereit":
                if self.format == "qrcode":
                    self.qr_dialog_starten()
                else:
                    self.dateien_waehlen()

    def dateien_waehlen(self):
        muster = " ".join(f"*.{e}" for e in sorted(ALLE_EINGABE))
        self.root.attributes("-topmost", False)
        pfade = filedialog.askopenfilenames(
            title="Dateien auswählen",
            filetypes=[("Unterstützte Dateien", muster), ("Alle Dateien", "*.*")])
        self.root.attributes("-topmost", self.top_var.get())
        self.root.after(80, self._fokus)
        if pfade:
            self.starten(list(pfade))

    def _on_enter(self, event):
        self._wecken()
        if self.zustand == "bereit":
            self.hover = True
            self.zeichne_texte()
        return event.action

    def _on_leave(self, event):
        self.hover = False
        self.zeichne_texte()
        return event.action

    def _on_drop(self, event):
        self._wecken()
        self.hover = False
        if self.sound_var.get():
            spiele_sound("drop")
        pfade = list(self.root.tk.splitlist(event.data))
        if self.zustand in ("laeuft", "wahl"):
            self.wartend.extend(pfade)       # in die Warteschlange statt verwerfen
            self.zeichne_texte()
            return event.action
        self.starten(pfade)
        return event.action

    def _sammle(self, pfade):
        ergebnis = []
        for p in pfade:
            p = os.path.normpath(p)
            if os.path.isdir(p):
                for root, _, files in os.walk(p):
                    for file in files:
                        ergebnis.append(os.path.join(root, file))
            else:
                ergebnis.append(p)
        return [p for p in ergebnis
                if os.path.isfile(p) and ext_von(p) in ALLE_EINGABE and ".~part" not in p]

    def _effektiv(self, fmt):
        """Tatsächliches Ziel und EXIF-Option (Web-Preset macht aus jedem Bildformat 'resize 1200px')."""
        if self.preset_web_var.get() and kat_von_format(fmt) == "Bild":
            return "resize 1200px", True
        return fmt, self.exif_var.get()

    @staticmethod
    def _passt(pfad, ziel_ext):
        """Passt die Datei zum Ziel? Eine von einer Erweiterung erkannte Datei passt immer (Erkennung hat Vorrang)."""
        if ziel_passt(ext_von(pfad), ziel_ext):
            return True
        return erkanntes_format(pfad, ziel_ext) is not None

    def starten(self, pfade, format_wahl=None, aus_watch=False):
        """Passende Dateien sofort umwandeln. Unpassende kommen (gruppiert) in den Modus FORMAT WÄHLEN –
        außer bei der Ordner-Überwachung, dort werden sie still übersprungen.
        format_wahl: im Modus gewähltes Format (ändert das gespeicherte Standardformat nicht)."""
        ziel_ext, strip_exif = self._effektiv(format_wahl or self.format)

        if ziel_ext in PACK_AKTIONEN:            # packen: jede Datei und ganze Ordner
            dateien, rest = [os.path.normpath(p) for p in pfade if os.path.exists(p)], []
        else:
            dateien, rest = [], []
            for d in self._sammle(pfade):
                if ist_folgeteil(d):
                    continue                     # .part2.rar ... kommt beim Entpacken von .part1.rar mit
                (dateien if self._passt(d, ziel_ext) else rest).append(d)

        if self.reset_id:
            self.root.after_cancel(self.reset_id)
            self.reset_id = None
        self.warnungen = []
        self.erkannt_text = ""
        if rest:
            if aus_watch or format_wahl:
                for d in rest:
                    logge(f"Übersprungen (passt nicht zu {ziel_ext}): {d}")
            else:
                self._wahl_einreihen(rest)
        if not dateien:
            if self.wahl_queue and self.zustand in ("bereit", "fertig"):
                self.zustand = "bereit"
                self._naechste_wahl()
                return
            self.ergebnis = (0, 1, "Keine passenden Dateien")
            self._zeige_fertig()
            self._ansage("format_nicht_unterstuetzt")
            return

        self.zustand = "laeuft"
        self.fortschritt = self.prog_anzeige = 0.0
        self.info = (f"0 / {len(dateien)}", "", "0 %")
        self.erparnis_text = ""
        self.abbruch = threading.Event()
        self.zeichne_texte()
        self._jai_ziel = ziel_ext
        self._ansage("datei_erkannt" if len(dateien) == 1 else "konvertierung_laeuft")
        args = (dateien, ziel_ext, self.ordner, self.qual.get(), strip_exif, self.abbruch)
        threading.Thread(target=self._arbeit, args=args, daemon=True).start()

    # ---- Modus FORMAT WÄHLEN
    def _wahl_einreihen(self, dateien):
        """Gruppiert unpassende Dateien nach ihren möglichen Zielformaten (gleiche Liste = eine Gruppe)."""
        gruppen = {}
        for d in dateien:
            formate = tuple(passende_formate(ext_von(d), lambda f: self._effektiv(f)[0]))
            if not formate:
                self.historie.append(f"{os.path.basename(d)} [FEHLER] Kein passendes Zielformat")
                continue
            gruppen.setdefault(formate, []).append(d)
        for formate, liste in gruppen.items():
            self.wahl_queue.append((liste, list(formate)))

    def _naechste_wahl(self):
        if not self.wahl_queue:
            return False
        dateien, formate = self.wahl_queue.pop(0)
        self._wahl_nr += 1
        self.wahl = {"dateien": dateien, "formate": formate, "i": 0,
                     "nr": self._wahl_nr, "gesamt": self._wahl_nr + len(self.wahl_queue)}
        if self.mini_mode:
            self._toggle_mini()              # die Auswahl braucht den großen Kreis
        self.hover = False
        self.zustand = "wahl"
        if not self._ansage("format_waehlen") and self.sound_var.get():
            spiele_sound("drop")
        self.zeichne_texte()
        self._fokus()
        return True

    def _wahl_blaettern(self, schritt):
        if self.zustand != "wahl" or not self.wahl:
            return
        w = self.wahl
        w["i"] = (w["i"] + schritt) % len(w["formate"])
        self.zeichne_texte()

    def _wahl_ende(self):
        self.wahl = None
        self.zustand = "bereit"
        if not self.wahl_queue:
            self._wahl_nr = 0

    def _wahl_start(self):
        if self.zustand != "wahl" or not self.wahl:
            return
        w = self.wahl
        fm = w["formate"][w["i"]]
        self._wahl_ende()
        self._wahl_start_t = time.time()
        self.starten(w["dateien"], format_wahl=fm)

    def _wahl_abbrechen(self):
        """ESC: nichts aus dieser und den noch offenen Gruppen wird umgewandelt."""
        if self.zustand != "wahl":
            return
        n = len(self.wahl["dateien"]) + sum(len(g[0]) for g in self.wahl_queue) if self.wahl else 0
        self.historie.append(f"{n} Datei(en) nicht umgewandelt (Format-Auswahl abgebrochen)")
        self.wahl_queue.clear()
        self._wahl_ende()
        self.zeichne_texte()
        self._ansage("abgebrochen")
        if self.wartend:
            pfade, self.wartend = self.wartend, []
            self.starten(pfade)

    def _abbrechen(self):
        if self.zustand != "laeuft":
            return
        self.abbruch.set()
        self.wartend.clear()
        stoppe_prozesse()
        self.info = (self.info[0], "Breche ab …", self.info[2])
        self.zeichne_texte()

    def _arbeit(self, dateien, ziel_ext, zielordner, qualitaet, strip_exif, abbruch):
        try:
            self._arbeit_lauf(dateien, ziel_ext, zielordner, qualitaet, strip_exif, abbruch)
        except Exception as ex:
            logge("Unerwarteter Fehler im Arbeits-Thread:\n" + traceback.format_exc())
            self.q.put((self._fertig, (0, 1, f"{type(ex).__name__}: {ex}", zielordner, 0, 0)))

    def _arbeit_lauf(self, dateien, ziel_ext, zielordner, qualitaet, strip_exif, abbruch):
        n = len(dateien)
        state = {"ok": 0, "fehler": 0, "erster_fehler": "", "letzter_ordner": zielordner,
                 "completed": 0, "uebersprungen": 0, "warnungen": []}
        groesse = {"orig": 0, "neu": 0}
        teil = {}
        sperre = threading.Lock()
        letzte_meldung = [0.0]

        def melde(name, force=False):
            jetzt = time.time()
            with sperre:
                if not force and jetzt - letzte_meldung[0] < 0.12:
                    return
                letzte_meldung[0] = jetzt
                frac = min(1.0, (state["completed"] + sum(teil.values())) / n)
                curr = state["completed"]
            self.q.put((self._schritt, (frac, f"{curr} / {n}", name)))

        def zu_ende(ergebnis_zeile=None):
            if abbruch.is_set():
                state["fehler"] += state["uebersprungen"]
                if not state["erster_fehler"]:
                    state["erster_fehler"] = "Vom Benutzer abgebrochen"
            self.q.put((self._fertig, (state["ok"], state["fehler"], state["erster_fehler"],
                                       state["letzter_ordner"], groesse["orig"], groesse["neu"],
                                       state["warnungen"])))

        # ---- Sammelaktionen: eine Ausgabe für mehrere Eingaben
        if ziel_ext in ("pdf zusammenführen", "zip (alle)", "7z (alle)"):
            try:
                self.q.put((self._schritt, (0.0, "1 / 1", "Verarbeite …")))
                if ziel_ext == "pdf zusammenführen":
                    pdfs = [d for d in dateien if ext_von(d) == "pdf"]
                    if not pdfs:
                        raise RuntimeError("Keine PDF-Dateien zum Zusammenführen ausgewählt")
                    ordner = zielordner or os.path.dirname(pdfs[0])
                    state["letzter_ordner"] = ordner
                    ziel = konvertiere_pdf_zusammenfuehren(pdfs, ordner)
                    state["ok"] = len(pdfs)
                    groesse["orig"] = sum(groesse_von(p) for p in pdfs)
                else:
                    ordner = zielordner or os.path.dirname(dateien[0])
                    state["letzter_ordner"] = ordner
                    os.makedirs(ordner, exist_ok=True)
                    name = (os.path.splitext(os.path.basename(dateien[0]))[0]
                            if len(dateien) == 1 and os.path.isfile(dateien[0]) else
                            os.path.basename(dateien[0].rstrip("/\\")) if len(dateien) == 1 else "Archiv")
                    endung = "7z" if ziel_ext == "7z (alle)" else "zip"
                    ziel = freier_pfad(os.path.join(ordner, f"{name}.{endung}"))
                    tmp = os.path.splitext(ziel)[0] + ".~part." + endung
                    try:
                        if endung == "7z":
                            siebenzip_erstellen(dateien, tmp, abbruch,
                                                lambda f: self.q.put((self._schritt, (f, "1 / 1", f"{name}.7z"))),
                                                ausschluss=[ziel])
                        else:
                            zip_erstellen(dateien, tmp, {tmp, ziel}, abbruch,
                                          lambda f, nm: self.q.put((self._schritt, (f, "1 / 1", nm))))
                        os.replace(tmp, ziel)
                    except BaseException:
                        raeume_auf(tmp)
                        raeume_auf(ziel)
                        raise
                    state["ok"] = len(dateien)
                    groesse["orig"] = sum(groesse_von(p) for p in dateien)
                groesse["neu"] = groesse_von(ziel)
                self.erzeugt.add(ziel)
                self.historie.append(f"{len(dateien)} Datei(en) -> {os.path.basename(ziel)} [OK]")
            except Abgebrochen:
                state["uebersprungen"] = 1
            except Exception as ex:
                state["fehler"] = 1
                state["erster_fehler"] = verstaendlich(ex)
                self.historie.append(f"{ziel_ext} [FEHLER] {verstaendlich(ex)}")
                logge(f"{ziel_ext}: {traceback.format_exc()}")
            zu_ende()
            return

        # ---- Einzelverarbeitung (parallel)
        def verarbeite_eine_datei(quelle):
            name = os.path.basename(quelle.rstrip("/\\")) or quelle
            erfolg, fehltext = False, ""
            erkannt, warnungen, zf = False, [], ziel_ext
            try:
                if abbruch.is_set():
                    raise Abgebrochen()
                quell_ext = ext_von(quelle)
                # Von einer Erweiterung erkannte Datei immer mit deren Format umwandeln (gewähltes Format bleibt)
                treffer = erkanntes_format(quelle, zf) if ERWEITERUNGEN else None
                if treffer:
                    zf = treffer[0]
                    erkannt = True
                    if treffer[1].get("erkannt_text"):
                        self.q.put((self._erkannt, (str(treffer[1]["erkannt_text"]),)))
                if zf.startswith("resize "):
                    target_dir = zielordner or os.path.join(os.path.dirname(quelle), "Verkleinert")
                else:
                    target_dir = zielordner or os.path.dirname(quelle)
                if zf.startswith("resize ") and quell_ext not in BILD_EINGABE:
                    raise RuntimeError("Verkleinern nur für Bilder (.jpg, .png, .webp, .heic)")
                with sperre:
                    state["letzter_ordner"] = target_dir
                os.makedirs(target_dir, exist_ok=True)
                basis = os.path.basename(quelle.rstrip("/\\"))
                if os.path.isfile(quelle):
                    basis = os.path.splitext(basis)[0]
                erzeugt = None

                def melde_teil(frac):
                    with sperre:
                        teil[quelle] = frac
                    melde(name)

                if zf.startswith("resize "):
                    if quell_ext not in BILD_EINGABE:
                        raise RuntimeError("Verkleinern nur für Bilder (.jpg, .png, .webp, .heic)")
                    erzeugt = konvertiere_bild_verkleinern(quelle, target_dir, zf, qualitaet, strip_exif)
                elif zf == "pdf teilen":
                    if quell_ext != "pdf":
                        raise RuntimeError("PDF teilen ist nur für PDF-Dateien möglich")
                    erzeugt = konvertiere_pdf_teilen(quelle, target_dir)
                    with sperre:
                        state["letzter_ordner"] = erzeugt
                elif zf == "pdf zu bild":
                    if quell_ext != "pdf":
                        raise RuntimeError("Nur für PDF-Dateien möglich")
                    erzeugt = konvertiere_pdf_zu_bild(quelle, target_dir)
                    with sperre:
                        state["letzter_ordner"] = erzeugt
                elif zf == "entpacken":
                    erzeugt = archiv_entpacken(quelle, target_dir, abbruch, melde_teil)
                    with sperre:
                        state["letzter_ordner"] = erzeugt
                else:
                    # Alles mit einer einzelnen Ausgabedatei: erst Ausgabe-Namen reservieren, in eine
                    # Hilfsdatei schreiben, am Ende atomar umbenennen -> nie halbfertige Dateien.
                    if zf == "txt (ocr)":
                        endung = "txt"
                        suffix = "_extracted"
                        if quell_ext == "pdf":
                            def arbeit(t):
                                with BILD_SEM:
                                    konvertiere_pdf_ocr(quelle, t, abbruch, melde_teil)
                        elif quell_ext in BILD_EINGABE:
                            def arbeit(t):
                                with BILD_SEM:
                                    konvertiere_bild_ocr(quelle, t)
                        else:
                            raise RuntimeError("Texterkennung geht nur bei PDF-Dateien und Bildern")
                    elif zf == "qrcode":
                        endung, suffix = "png", "_qrcode"
                        arbeit = lambda t: konvertiere_qrcode(quelle, t)
                    elif zf in ("zip", "7z"):
                        endung, suffix = zf, ""
                        arbeit = None   # unten
                    else:
                        konv, art = waehle_konverter(quell_ext, zf)
                        endung, suffix = ausgabe_ext(zf), ""
                        auf = Auftrag(zf, qualitaet, strip_exif, self.ffmpeg, abbruch, melde_teil)
                        warnungen = auf.warnungen
                        arbeit = lambda t: konv(quelle, t, auf)
                    final = freier_pfad(os.path.join(target_dir, f"{basis}{suffix}.{endung}"))
                    tmp = os.path.splitext(final)[0] + ".~part." + endung
                    try:
                        if zf == "zip":
                            zip_erstellen([quelle], tmp, {tmp, final}, abbruch)
                        elif zf == "7z":
                            siebenzip_erstellen([quelle], tmp, abbruch, melde_teil, ausschluss=[final])
                        else:
                            arbeit(tmp)
                        if not os.path.exists(tmp) or os.path.getsize(tmp) == 0:
                            raise RuntimeError("Es wurde keine Ausgabedatei erzeugt")
                        os.replace(tmp, final)
                    except BaseException:
                        raeume_auf(tmp)
                        raeume_auf(final)
                        raise
                    erzeugt = final

                if erzeugt:
                    self.erzeugt.add(erzeugt)
                    with sperre:
                        groesse["orig"] += groesse_von(quelle)
                        groesse["neu"] += groesse_von(erzeugt)
                erfolg = True
            except Abgebrochen:
                with sperre:
                    state["uebersprungen"] += 1
            except BaseException as ex:       # niemals den Thread-Pool abstürzen lassen
                fehltext = verstaendlich(ex)
                if isinstance(ex, RuntimeError):
                    logge(f"{quelle} -> {zf}: {fehltext}")
                else:
                    logge(f"{quelle} -> {zf}:\n{traceback.format_exc()}")

            with sperre:
                state["completed"] += 1
                teil.pop(quelle, None)
                if erfolg:
                    state["ok"] += 1
                    zusatz = " (erkannt)" if erkannt else ""
                    if warnungen:
                        warntext = " · ".join(dict.fromkeys(warnungen))
                        state["warnungen"].append(warntext)
                        self.historie.append(f"{name} -> .{zf}{zusatz} [WARNUNG] {warntext}")
                        logge(f"WARNUNG {quelle} -> {zf}: {warntext}")
                    else:
                        self.historie.append(f"{name} -> .{zf}{zusatz} [OK]")
                elif fehltext:
                    state["fehler"] += 1
                    if not state["erster_fehler"]:
                        state["erster_fehler"] = fehltext
                    self.historie.append(f"{name} [FEHLER] {fehltext}")
            melde(name, force=True)

        max_threads = min(8, max(2, os.cpu_count() or 4))
        with ThreadPoolExecutor(max_workers=max_threads) as executor:
            list(executor.map(verarbeite_eine_datei, dateien))
        zu_ende()

    def _schritt(self, frac, zaehler, name):
        if self.zustand != "laeuft":
            return
        self.fortschritt = frac
        self.info = (zaehler, name, f"{int(frac * 100)} %")
        self.zeichne_texte()

    def _erkannt(self, text):
        self.erkannt_text = text
        self.zeichne_texte()

    def _fertig(self, ok, fehler, fehltext, ordner, sz_orig, sz_new, warnungen=None):
        if sz_orig > 0 and sz_new > 0:
            diff = sz_orig - sz_new
            prozent = (diff / sz_orig) * 100
            if diff > 0:
                self.erparnis_text = f"📉 {prozent:.1f}% SPEICHER GESPART"
            else:
                self.erparnis_text = "✨ ERFOLGREICH KONVERTIERT"
        else:
            self.erparnis_text = ""

        self.ergebnis = (ok, fehler, fehltext)
        self.warnungen = list(warnungen or [])
        del self.historie[:-200]
        self._zeige_fertig()
        cid = self._jai_ergebnis(ok, fehler, fehltext)
        if not (cid and self._ansage(cid)) and self.sound_var.get():
            spiele_sound("ok" if ok else "err")
        if ok and self.open_var.get() and ordner and not self.wartend and hasattr(os, "startfile"):
            try:
                os.startfile(ordner)
            except Exception:
                pass

    def _zeige_fertig(self):
        self.zustand = "fertig"
        self.t_fertig = time.time()
        self.zeichne_texte()
        if self.ergebnis[0]:
            self._funken_start()
        if self.reset_id:
            self.root.after_cancel(self.reset_id)
            self.reset_id = None
        if self.wahl_queue and not (self.ergebnis[1] or self.warnungen):
            verzoegerung = 1500            # danach geht es mit FORMAT WÄHLEN für den Rest weiter
        elif self.wartend:
            verzoegerung = 1500
        elif self.ergebnis[1] or self.warnungen:
            verzoegerung = None            # Fehler-/Warnansicht bleibt, bis geklickt wird
        else:
            verzoegerung = 7000
        if verzoegerung:
            self.reset_id = self.root.after(verzoegerung, self._zurueck)

    def _zurueck(self, event=None):
        if self.reset_id:
            self.root.after_cancel(self.reset_id)
            self.reset_id = None
        self.zustand = "bereit"
        self.fortschritt = self.prog_anzeige = 0.0
        self.erparnis_text = ""
        self.warnungen = []
        self.erkannt_text = ""
        self.zeichne_texte()
        if self.wartend:
            pfade, self.wartend = self.wartend, []
            self.starten(pfade)
        elif self.wahl_queue:
            self._naechste_wahl()

    def _poll(self):
        try:
            while True:
                funktion, args = self.q.get_nowait()
                funktion(*args)
        except queue.Empty:
            pass
        self.root.after(60, self._poll)

    # ---- Einstellungen, Beenden, Tray, Überwachung
    def _speichern(self):
        try:
            d = lade_einstellungen()
            d.update({
                "theme": self.theme_var.get(), "design": 2, "qual": self.qual.get(),
                "open": self.open_var.get(), "top": self.top_var.get(), "sound": self.sound_var.get(),
                "exif": self.exif_var.get(), "preset_web": self.preset_web_var.get(),
                "tray_min": self.tray_min_var.get(), "ordner": self.ordner, "format": self.format,
                "watch_ordner": self.watch_ordner, "office_bevorzugen": OFFICE_BEVORZUGEN[0],
                "fenster": {"x": self._pos[0] + self.size // 2, "y": self._pos[1] + self.size // 2},
                "mini": self.mini_mode, "jai": self.jai_var.get(), "effekte": self.effekte,
            })
            speichere_einstellungen(d)
        except Exception:
            pass

    def _beenden(self):
        try:
            self.abbruch.set()
            stoppe_prozesse()
            self._speichern()
            if self._tray is not None:
                self._tray.stop()
        finally:
            self.root.destroy()

    def _tray_verstecken(self):
        if not TRAY_AVAILABLE:
            return
        if self._tray is None:
            farbe = hexrgb(theme_daten(self.theme_var.get())["CYAN"])
            bild = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
            d = ImageDraw.Draw(bild)
            d.ellipse([3, 3, 61, 61], fill=(5, 9, 24, 255), outline=farbe + (255,), width=5)
            d.ellipse([22, 22, 42, 42], fill=farbe + (255,))
            menue = pystray.Menu(
                pystray.MenuItem("Anzeigen", lambda i, it: self.q.put((self._tray_zeigen, ())), default=True),
                pystray.MenuItem("Beenden", lambda i, it: self.q.put((self._beenden, ()))))
            self._tray = pystray.Icon("KreisKonverter", bild, "Kreis-Konverter", menue)
            self._tray.run_detached()
        self.root.withdraw()

    def _tray_zeigen(self):
        self.root.deiconify()
        self.root.lift()
        self.root.attributes("-topmost", self.top_var.get())
        if self._tray is not None:
            self._tray.stop()
            self._tray = None
        self.root.after(80, self._fokus)

    def _watch_toggle(self):
        if self.watch_var.get():
            self.watch_var.set(False)
            self.zeichne_texte()
            return
        self.root.attributes("-topmost", False)
        p = filedialog.askdirectory(title="Ordner zum Überwachen wählen",
                                    initialdir=self.watch_ordner or self.ordner or os.path.expanduser("~"))
        self.root.attributes("-topmost", self.top_var.get())
        self.root.after(80, self._fokus)
        if not p:
            return
        self.watch_ordner = os.path.normpath(p)
        # Bereits vorhandene Dateien gelten als "gesehen" und werden nicht konvertiert
        self._watch_gesehen = {os.path.join(self.watch_ordner, f) for f in os.listdir(self.watch_ordner)}
        self._watch_groessen = {}
        self.watch_var.set(True)
        self._speichern()
        self.zeichne_texte()

    def _watch_moeglich(self, pfad):
        if ist_folgeteil(pfad):
            return False
        return self._passt(pfad, self._effektiv(self.format)[0])

    def _watch_loop(self):
        try:
            if (self.watch_var.get() and self.watch_ordner and os.path.isdir(self.watch_ordner)
                    and self.zustand == "bereit"):
                neu = []
                for f in os.listdir(self.watch_ordner):
                    p = os.path.join(self.watch_ordner, f)
                    if (p in self._watch_gesehen or p in self.erzeugt or ".~part" in f
                            or ext_von(f) not in ALLE_EINGABE or not os.path.isfile(p)):
                        continue
                    try:
                        sz = os.path.getsize(p)
                    except OSError:
                        continue
                    # Erst konvertieren, wenn die Größe zwei Prüfungen lang gleich blieb (Kopieren fertig)
                    if sz > 0 and self._watch_groessen.get(p) == sz:
                        neu.append(p)
                    else:
                        self._watch_groessen[p] = sz
                if neu:
                    self._watch_gesehen.update(neu)
                    for p in neu:
                        self._watch_groessen.pop(p, None)
                    neu = [p for p in neu if self._watch_moeglich(p)]   # Unpassendes still überspringen
                    if neu:
                        self.starten(neu, aus_watch=True)
        except Exception:
            logge("Watch-Folder: " + traceback.format_exc())
        self.root.after(2000, self._watch_loop)

    def start(self):
        self.root.mainloop()


if __name__ == "__main__":
    app = KreisApp()
    # Dateien, die per "Öffnen mit" / "Senden an" übergeben wurden, sofort verarbeiten
    uebergeben = [a for a in sys.argv[1:] if os.path.exists(a)]
    if uebergeben:
        app.root.after(400, lambda: app.starten(uebergeben))
    app.start()
